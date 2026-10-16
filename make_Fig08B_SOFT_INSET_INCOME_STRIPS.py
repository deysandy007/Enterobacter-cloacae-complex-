
#!/usr/bin/env python3

# ============================================================
# FIGURE 8B — COMBINED MLST TREEMAP
# SOFT, INSET COUNTRY-INCOME STRIPS
#
# Every valid ST gets its own rectangle.
# Each rectangle receives a faded income-composition strip
# INSIDE the lower portion, separated from its boundaries.
#
# One landscape JPEG at 600 DPI + complete Excel.
# ============================================================

from pathlib import Path
import re
import gc

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from matplotlib.patches import Rectangle, Patch
from matplotlib.colors import to_rgb

import squarify

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

INCOME_DB = (
    ROOT / "Fig_01A_Global_Geographical_Distribution"
    / "Fig_01A_Country_Income_Database.xlsx"
)

OUT = (
    ROOT / "Fig_08B_Combined_MLST_Treemap"
    / "SOFT_INSET_INCOME_STRIPS"
)
OUT.mkdir(parents=True, exist_ok=True)

JPG = OUT / "Fig08B_Combined_All_ST_Soft_Inset_Income_600DPI.jpg"
XLSX = OUT / "Fig08B_Combined_All_ST_Soft_Inset_Data.xlsx"

DPI = 600
FIG_W, FIG_H = 24, 16.7
W, H = 1000.0, 720.0

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis"
]

SPECIES_COLORS = dict(zip(SPECIES, [
    "#0072B2", "#D55E00", "#009E73",
    "#CC79A7", "#7851A9", "#A6761D", "#56B4E9"
]))

ROWS = [
    [SPECIES[0]],
    [SPECIES[1], SPECIES[2], SPECIES[3]],
    [SPECIES[4], SPECIES[5], SPECIES[6]]
]

GROUPS = ["H", "M", "L+LM", "Unclassified"]

INCOME_COLORS = {
    "H": "#C98A00",
    "M": "#673AB7",
    "L+LM": "#C94F82",
    "Unclassified": "#929DA8"
}

# ============================================================
# AESTHETIC CONTROLS
# ============================================================

# 0.00 = original income color; 1.00 = pure white.
# 0.37 provides a softer, faded income strip.
STRIP_FADE = 0.37

# Strip occupies 16% of ST tile height.
STRIP_HEIGHT_FRACTION = 0.16

# Margin between strip and left/right tile edges.
STRIP_LEFT_RIGHT_MARGIN = 0.045

# Margin between strip and bottom tile edge.
STRIP_BOTTOM_MARGIN = 0.055

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "figure.facecolor": "white"
})


# ============================================================
# FUNCTIONS
# ============================================================

def clean_st(x):
    if pd.isna(x):
        return ""

    t = str(x).strip().strip("'\"").strip()

    if t.upper().startswith("ST"):
        t = t[2:]

    if re.fullmatch(r"\d+\.0+", t):
        t = t.split(".")[0]

    if re.fullmatch(r"\d+", t) and int(t) > 0:
        return "ST" + str(int(t))

    return ""


def clean_country(x):
    if pd.isna(x):
        return ""

    return re.sub(
        r"\s+", " ", str(x).strip()
    ).casefold()


def clean_income(x):
    if pd.isna(x):
        return "Unclassified"

    t = re.sub(
        r"[\s_\-]+", "",
        str(x).upper().strip()
    )

    if t in {
        "H", "HIGH", "HIGHINCOME",
        "HIGHINCOMECOUNTRIES"
    }:
        return "H"

    if t in {
        "M", "UM", "UMIC", "UPPERMIDDLE",
        "UPPERMIDDLEINCOME",
        "UPPERMIDDLEINCOMECOUNTRIES"
    }:
        return "M"

    if t in {
        "L+LM", "LM+L", "LLM", "LMICLIC",
        "LOWERMIDDLE+LOW", "LOW+LOWERMIDDLE",
        "LOWERMIDDLEANDLOW",
        "LOWERMIDDLEANDLOWINCOME",
        "LOWERMIDDLE", "LOWERMIDDLEINCOME",
        "LOW", "LOWINCOME"
    }:
        return "L+LM"

    return "Unclassified"


def blend(color, white):
    rgb = np.asarray(to_rgb(color))

    return tuple(
        np.clip(
            rgb * (1 - white) + white,
            0, 1
        )
    )


def ink(face):
    r, g, b = to_rgb(face)

    lum = 0.2126*r + 0.7152*g + 0.0722*b

    return "white" if lum < .46 else "#173148"


def workbook_style(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for c in ws[1]:
            c.fill = PatternFill(
                "solid", fgColor="173352"
            )
            c.font = Font(
                color="FFFFFF", bold=True
            )

        for j, c in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(j)
            ].width = min(
                42, max(15, len(str(c.value)) + 3)
            )


# ============================================================
# LOAD MASTER
# ============================================================

print("=" * 90)
print("FIGURE 8B — SOFT INSET INCOME STRIPS")
print("=" * 90)

cols = [
    "Assembly_Accession",
    "Normalized_Species",
    "MLST_ST",
    "MLST_Scheme",
    "Country_Harmonized"
]

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=cols,
    dtype=str,
    engine="openpyxl"
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9157 genomes: {len(df)}"
    )

if (
    df.Assembly_Accession.isna().any()
    or df.Assembly_Accession.duplicated().any()
):
    raise ValueError("Missing/duplicate accessions")

df["Species"] = (
    df.Normalized_Species.fillna("").str.strip()
)

df["ST"] = df.MLST_ST.map(clean_st)

df["Scheme"] = (
    df.MLST_Scheme
    .fillna("")
    .str.strip()
    .str.lower()
)

df["Country_Key"] = (
    df.Country_Harmonized.map(clean_country)
)

df["Resolved"] = df.Species.isin(SPECIES)

df["Eligible"] = (
    df.Resolved
    & df.Scheme.eq("ecloacae")
    & df.ST.ne("")
)

# ============================================================
# FROZEN COUNTRY-INCOME DATABASE
# ============================================================

inc = pd.read_excel(
    INCOME_DB,
    sheet_name="COUNTRY_INCOME_COUNTS",
    engine="openpyxl",
    dtype=str
)

for col in ("Country_Harmonized", "Income_Group"):
    if col not in inc:
        raise ValueError(
            "Missing income mapping column: " + col
        )

inc["Country_Key"] = (
    inc.Country_Harmonized.map(clean_country)
)

inc["Income"] = (
    inc.Income_Group.map(clean_income)
)

inc = inc[inc.Country_Key.ne("")].copy()

conf = inc.groupby("Country_Key").Income.nunique()

if (conf > 1).any():
    raise ValueError(
        "Conflicting income entries: "
        + str(conf[conf > 1].index.tolist()[:10])
    )

lookup = (
    inc.drop_duplicates("Country_Key")
    .set_index("Country_Key")
    .Income
)

df["Income"] = (
    df.Country_Key
    .map(lookup)
    .fillna("Unclassified")
)

resolved = df[df.Resolved].copy()
typed = df[df.Eligible].copy()

if len(resolved) != 8882:
    raise ValueError(
        f"Expected 8882 resolved genomes: {len(resolved)}"
    )

if typed.empty:
    raise ValueError("No ST-assigned genomes")

# ============================================================
# CALCULATE COMPLETE ST AND INCOME ABUNDANCE
# ============================================================

spall = (
    resolved.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

sptyped = (
    typed.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

counts = (
    typed.groupby(["Species", "ST"])
    .size()
    .rename("Genome_Count")
    .reset_index()
)

inc_counts = (
    typed.groupby(["Species", "ST", "Income"])
    .size()
    .unstack("Income", fill_value=0)
    .reindex(columns=GROUPS, fill_value=0)
    .reset_index()
)

inc_counts = inc_counts.rename(columns={
    "H": "N_H",
    "M": "N_M",
    "L+LM": "N_LLM",
    "Unclassified": "N_Unclassified"
})

st = counts.merge(
    inc_counts,
    on=["Species", "ST"],
    validate="one_to_one"
)

st["Species_Total"] = st.Species.map(spall)
st["Typed_Total"] = st.Species.map(sptyped)

st["Within_Species_Percent"] = (
    100 * st.Genome_Count / st.Species_Total
)

st["Within_Typed_Percent"] = (
    100 * st.Genome_Count / st.Typed_Total
)

st["ST_Numeric"] = st.ST.str[2:].astype(int)

st["Income_Breadth"] = (
    (st[["N_H", "N_M", "N_LLM"]] > 0)
    .sum(axis=1)
)

st = st.sort_values(
    ["Species", "Genome_Count", "ST_Numeric"],
    ascending=[True, False, True]
).reset_index(drop=True)

for group, col in zip(
    GROUPS,
    ["N_H", "N_M", "N_LLM", "N_Unclassified"]
):
    st["Pct_" + group] = (
        100 * st[col] / st.Genome_Count
    )

income_count_cols = [
    "N_H", "N_M", "N_LLM", "N_Unclassified"
]

if not (
    st[income_count_cols].sum(axis=1)
    == st.Genome_Count
).all():
    raise ValueError("Income counts mismatch")

# ============================================================
# CANVAS
# ============================================================

fig = plt.figure(
    figsize=(FIG_W, FIG_H),
    facecolor="white"
)

left, bottom, aw, ah = .026, .138, .948, .778

ax = fig.add_axes([left, bottom, aw, ah])

ax.set_xlim(0, W)
ax.set_ylim(0, H)
ax.invert_yaxis()
ax.axis("off")

ax_w_in = FIG_W * aw
ax_h_in = FIG_H * ah

records = []
blocks = []

yy = 0.0
total = len(typed)

# ============================================================
# DRAW ALL SPECIES AND ST TILES
# ============================================================

for line in ROWS:

    row_n = sum(int(sptyped[s]) for s in line)
    row_h = H * row_n / total

    xx = 0.0

    for species in line:

        n = int(sptyped[species])
        all_n = int(spall[species])

        bw = W * n / row_n
        base = SPECIES_COLORS[species]

        ax.add_patch(Rectangle(
            (xx + 1, yy + 1),
            bw - 2,
            row_h - 2,
            facecolor=base,
            edgecolor="white",
            linewidth=2,
            zorder=1
        ))

        head = min(30, row_h * .22)

        label = (
            species if bw >= 340
            else species.replace("Enterobacter ", "E. ")
        )

        title = (
            f"{label} "
            f"(n={n:,}; {100*n/total:.1f}%)"
        )

        ax.text(
            xx + 8,
            yy + head*.54,
            title,
            color="white",
            fontsize=13 if bw >= 340 else 9,
            weight="bold",
            style="italic",
            va="center",
            zorder=7
        )

        px = xx + 6
        py = yy + head + 2
        pw = max(.1, bw - 12)
        ph = max(.1, row_h - head - 8)

        ss = (
            st[st.Species.eq(species)]
            .sort_values(
                ["Genome_Count", "ST_Numeric"],
                ascending=[False, True]
            )
            .reset_index(drop=True)
        )

        rects = squarify.squarify(
            squarify.normalize_sizes(
                ss.Genome_Count.astype(float).tolist(),
                pw, ph
            ),
            px, py, pw, ph
        )

        if len(rects) != len(ss):
            raise ValueError(
                "Tile mismatch for " + species
            )

        for i, (geom, r) in enumerate(
            zip(rects, ss.itertuples(index=False))
        ):

            x, y, w, h = (
                float(geom[k])
                for k in ["x", "y", "dx", "dy"]
            )

            # Soft species background.
            face = blend(
                base,
                .14 + .64*i/max(1, len(ss)-1)
            )

            outline = (
                "#263B4B"
                if r.Income_Breadth >= 2
                else "white"
            )

            patch = Rectangle(
                (x, y),
                w, h,
                facecolor=face,
                edgecolor=outline,
                linewidth=.10 if min(w, h) < 1 else .40,
                zorder=2
            )
            ax.add_patch(patch)

            # ============================================
            # INSET INCOME STRIP
            # ============================================

            margin_x = (
                w * STRIP_LEFT_RIGHT_MARGIN
            )

            margin_bottom = (
                h * STRIP_BOTTOM_MARGIN
            )

            sw = w - 2*margin_x
            sh = h * STRIP_HEIGHT_FRACTION

            sx = x + margin_x
            sy = y + h - margin_bottom - sh

            cursor = sx

            for group, key in zip(
                GROUPS,
                income_count_cols
            ):

                count = int(getattr(r, key))

                if count <= 0:
                    continue

                seg = (
                    sw * count / int(r.Genome_Count)
                )

                piece = Rectangle(
                    (cursor, sy),
                    seg, sh,
                    facecolor=blend(
                        INCOME_COLORS[group],
                        STRIP_FADE
                    ),
                    edgecolor="none",
                    linewidth=0,
                    zorder=5,
                    antialiased=False
                )

                piece.set_clip_path(patch)
                ax.add_patch(piece)

                cursor += seg

            # Subtle outline around inset strip.
            if (
                sw/W*ax_w_in > .12
                and sh/H*ax_h_in > .035
            ):

                frame = Rectangle(
                    (sx, sy),
                    sw, sh,
                    fill=False,
                    edgecolor="white",
                    linewidth=.32,
                    zorder=6
                )

                frame.set_clip_path(patch)
                ax.add_patch(frame)

            # ============================================
            # ST LABEL ABOVE THE STRIP
            # ============================================

            width_inches = (
                w/W*ax_w_in
            )

            text_inches = max(
                0,
                (sy-y)/H*ax_h_in
            )

            label_status = "Too small to label"

            if (
                width_inches >= .88
                and text_inches >= .55
            ):

                txt = (
                    f"{r.ST}\n"
                    f"n={int(r.Genome_Count):,}\n"
                    f"{r.Within_Species_Percent:.1f}%"
                )

                fs = min(
                    12.5,
                    width_inches*9.7,
                    text_inches*13
                )

            elif (
                width_inches >= .56
                and text_inches >= .32
            ):

                txt = (
                    f"{r.ST}\n"
                    f"n={int(r.Genome_Count):,}"
                )

                fs = min(
                    9.5,
                    width_inches*10.5,
                    text_inches*16
                )

            elif (
                width_inches >= .36
                and text_inches >= .17
            ):

                txt = r.ST

                fs = min(
                    7.5,
                    width_inches*10.8,
                    text_inches*22
                )

            else:
                txt = None
                fs = 0

            if txt and fs >= 5.5:

                artist = ax.text(
                    x+w/2,
                    y+(sy-y)/2,
                    txt,
                    ha="center",
                    va="center",
                    fontsize=fs,
                    color=ink(face),
                    weight="bold",
                    linespacing=1.10,
                    zorder=8
                )

                artist.set_clip_path(patch)
                label_status = "Labeled"

            records.append({
                "Species": species,
                "ST": r.ST,
                "Genome_Count": int(r.Genome_Count),
                "Within_Species_Percent": float(
                    r.Within_Species_Percent
                ),
                "N_H": int(r.N_H),
                "N_M": int(r.N_M),
                "N_LLM": int(r.N_LLM),
                "N_Unclassified": int(r.N_Unclassified),
                "Income_Breadth": int(
                    r.Income_Breadth
                ),
                "Strip_Drawn": True,
                "Label_Status": label_status,
                "X": x,
                "Y": y,
                "Width": w,
                "Height": h
            })

        blocks.append({
            "Species": species,
            "Total_Genomes": all_n,
            "ST_Assigned": n,
            "Unassigned_ST": all_n-n,
            "Unique_ST": len(ss),
            "Species_Block_Proportion": n/total
        })

        xx += bw

    yy += row_h

# ============================================================
# FIGURE TEXT AND INCOME LEGEND
# ============================================================

fig.text(
    .5, .973,
    "Combined MLST abundance and country-income composition",
    ha="center",
    fontsize=21,
    weight="bold",
    color="#173352"
)

fig.text(
    .5, .943,
    (
        "Enterobacter cloacae complex | "
        f"{total:,} ST-assigned genomes"
    ),
    ha="center",
    fontsize=12,
    color="#516779"
)

handles = [
    Patch(
        facecolor=blend(
            INCOME_COLORS[group],
            STRIP_FADE
        ),
        label=label
    )
    for group, label in [
        ("H", "High income"),
        ("M", "Upper-middle income"),
        ("L+LM", "Lower-middle + low income"),
        ("Unclassified", "Unclassified")
    ]
]

fig.legend(
    handles=handles,
    loc="lower center",
    bbox_to_anchor=(.5, .098),
    ncol=4,
    frameon=False,
    fontsize=10
)

fig.text(
    .5, .072,
    "Each rectangle represents one ST. "
    "Soft inset strip shows income composition within that ST.",
    ha="center",
    fontsize=10,
    color="#394F63"
)

fig.text(
    .5, .050,
    "Rectangle areas represent ST-assigned genome counts. "
    "No STs are pooled into an Other category.",
    ha="center",
    fontsize=9.4,
    color="#51677A"
)

fig.text(
    .5, .030,
    "Very small ST tiles retain their inset income strips; "
    "individual segments may be smaller than one pixel.",
    ha="center",
    fontsize=8.8,
    color="#65798A"
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 94,
        "subsampling": 1
    }
)

plt.close(fig)
gc.collect()

# ============================================================
# EXPORT AND VERIFY
# ============================================================

tiles = pd.DataFrame(records)

assert len(tiles) == len(st)
assert int(tiles.Genome_Count.sum()) == total
assert tiles.Strip_Drawn.all()

qc = pd.DataFrame([
    ("Master genomes", len(df)),
    ("Resolved genomes", len(resolved)),
    ("Assigned ST genomes", len(typed)),
    (
        "Unassigned ST genomes",
        len(resolved)-len(typed)
    ),
    ("Individual ST tiles", len(tiles)),
    ("Other ST grouping", 0),
    ("Inset strips", len(tiles)),
    ("Strip fade toward white", STRIP_FADE),
    ("Income database", str(INCOME_DB)),
    (
        "Small segments",
        "Subpixel segments may be visually indistinguishable"
    )
], columns=["Parameter", "Value"])

with pd.ExcelWriter(
    XLSX,
    engine="openpyxl"
) as writer:

    df[[
        "Assembly_Accession",
        "Species",
        "ST",
        "Country_Harmonized",
        "Income",
        "Resolved",
        "Eligible"
    ]].to_excel(
        writer,
        sheet_name="ALL_9157_GENOMES",
        index=False
    )

    st.to_excel(
        writer,
        sheet_name="ALL_ST_INCOME_COUNTS",
        index=False
    )

    tiles.to_excel(
        writer,
        sheet_name="EVERY_ST_TILE",
        index=False
    )

    pd.DataFrame(blocks).to_excel(
        writer,
        sheet_name="SPECIES_SUMMARY",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    workbook_style(writer)

print()
print("=" * 90)
print("FIGURE 8B COMPLETED")
print("=" * 90)
print("ST tiles:", len(tiles))
print("Income strips:", len(tiles))
print("JPEG:", JPG)
print("Excel:", XLSX)
print("=" * 90)
