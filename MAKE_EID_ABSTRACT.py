
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_LINE_SPACING

OUT = Path(
    r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES\EID_Abstract_200_Words.docx"
)

TITLE = (
    "Global Genomic Surveillance Reveals Species-Specific, "
    "Income-Stratified, and One Health Heterogeneity in "
    "Antimicrobial Resistance across the Enterobacter cloacae Complex"
)

ABSTRACT = """The Enterobacter cloacae complex (ECC) comprises emerging opportunistic pathogens increasingly associated with multidrug-resistant infections. However, global differences in its antimicrobial resistome across socioeconomic settings and One Health reservoirs remain poorly characterized. We analyzed 9,157 quality-filtered, nonredundant ECC genomes collected between 2000 and 2026, integrating geographical origin, World Bank income classification, isolation source, multilocus sequence typing, and antimicrobial resistance gene profiles. Genomes originated from 97 countries, with substantial geographical inequalities and predominance of human-associated isolates (79.5%). Seven resolved ECC species exhibited extensive sequence type diversity and significant isolation-source-associated differences in species composition. Although genome-level species distributions varied across income groups, country-level rarefaction and sensitivity analyses revealed no consistent income-associated differences in taxonomic diversity. Antimicrobial resistance gene burden varied across species, ecological compartments, income strata, and collection periods. Clinically important determinants, including blaCTX-M-15, mcr-10.1, and tet(X4), displayed heterogeneous distributions, with selected resistance-gene combinations concentrated within specific countries and sequencing projects. These findings reveal substantial species-specific and ecological heterogeneity in the sampled global ECC resistome, emphasizing the importance of geographically balanced genomic surveillance and integrated One Health monitoring to inform antimicrobial resistance surveillance priorities and strengthen the interpretation of emerging resistance threats worldwide."""

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.9)
section.bottom_margin = Inches(0.9)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE)
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(14)
r.font.italic = False

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ABSTRACT")
r.bold = True
r.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing = 1.5

# Italicize the organism name wherever it occurs.
organism = "Enterobacter cloacae"
parts = ABSTRACT.split(organism)

for i, part in enumerate(parts):
    p.add_run(part)
    if i < len(parts) - 1:
        p.add_run(organism).italic = True

word_count = len(ABSTRACT.split())

doc.add_paragraph()
note = doc.add_paragraph(
    f"Word count: {word_count} (excluding title)"
)
note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
note.runs[0].italic = True
note.runs[0].font.size = Pt(10)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

print("=" * 65)
print("EID ABSTRACT WORD DOCUMENT CREATED")
print("=" * 65)
print("Word count:", word_count)
print("Saved to:", OUT)
