
#!/usr/bin/env python3

# ============================================================
# FIGURE 8B — COMBINED HIERARCHICAL MLST TREEMAP
# ENTEROBACTER CLOACAE COMPLEX
#
# 7 species in ONE LANDSCAPE figure.
#
# SPECIES BLOCK AREA = total species genome count
# ST TILE AREA = genomes assigned to each ST
# LABEL = ST + genome count + within-species percentage
#
# Actual data from 9,157-genome final master.
# Unresolved species excluded from the seven-species figure.
# Unassigned STs retained as a separate category.
#
# Outputs:
#   1 combined JPEG, 600 DPI
#   1 complete Excel workbook
# ============================================================

from pathlib import Path
import re
import gc
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.colors import to_rgb
import squarify

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT
    / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

OUT = ROOT / "Fig_08B_Combined_MLST_Treemap"
OUT.mkdir(parents=True, exist_ok=True)

JPG = OUT / "Fig08B_Combined_Seven_Species_MLST_Treemap_600DPI.jpg"

EXCEL = OUT / "Fig08B_Combined_MLST_Treemap_Data.xlsx"

# ============================================================
# SETTINGS
# ============================================================

DPI = 600

FIG_WIDTH = 18
FIG_HEIGHT = 13.5

# Individual major STs displayed.
# Others are combined into "Other STs" in the figure.
TOP_ST_HORMAECHEI = 10
TOP_ST_OTHER_SPECIES = 5

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

# Same order as the approved composite figure:
# first row: largest species
# second row: next three species
# third row: final three species
SPECIES_ROWS = [
    ["Enterobacter hormaechei"],
    [
        "Enterobacter cloacae",
        "Enterobacter roggenkampii",
        "Enterobacter asburiae"
    ],
    [
        "Enterobacter kobei",
        "Enterobacter ludwigii",
        "Enterobacter bugandensis"
    ]
]

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 11,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

# ============================================================
# FUNCTIONS
# ============================================================

def clean_st(value):
    if pd.isna(value):
        return ""

    text = str(value).strip().strip("'\"").strip()

    if text.upper().startswith("ST"):
        text = text[2:].strip()

    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]

    if not re.fullmatch(r"\d+", text):
        return ""

    number = int(text)

    return f"ST{number}" if number > 0 else ""


def blend(base, amount):
    """Lighten a species color without changing its identity."""
    rgb = np.array(to_rgb(base))
    return tuple(
        np.clip(rgb * (1.0 - amount) + amount, 0, 1)
    )


def contrast_color(rgb):
    r, g, b = to_rgb(rgb)
    brightness = 0.2126*r + 0.7152*g + 0.0722*b
    return "white" if brightness < 0.54 else "#142B3D"


def fmt(value):
    return f"{int(value):,}"


def draw_fitted_label(
    ax, patch, x, y, w, h,
    title, count, percent, face,
    font_scale=1.0
):
    """
    Labels adapt to available tile dimensions.
    All text is clipped to its own rectangle.
    """

    if w < 24 or h < 15:
        return

    color = contrast_color(face)

    # Estimate available label width in points.
    # Axis data width maps to the physical figure width.
    width_in = w / 1000 * FIG_WIDTH * 0.93
    height_in = h / 720 * FIG_HEIGHT * 0.80

    max_font_by_width = max(
        5.5,
        width_in * 72 / max(5, len(title) * 0.57)
    )

    if width_in >= 1.65 and height_in >= 0.75:
        lines = [
            title,
            f"n={fmt(count)}",
            f"{percent:.1f}%"
        ]
        fontsize = min(
            14.5, max_font_by_width, height_in * 19
        )

    elif width_in >= 0.85 and height_in >= 0.47:
        lines = [
            title,
            f"n={fmt(count)}"
        ]
        fontsize = min(
            11.2, max_font_by_width, height_in * 22
        )

    elif width_in >= 0.55 and height_in >= 0.22:
        lines = [title]
        fontsize = min(
            9.0, max_font_by_width, height_in * 25
        )

    else:
        return

    fontsize = max(5.5, fontsize * font_scale)

    text = ax.text(
        x + w/2,
        y + h/2,
        "\n".join(lines),
        ha="center",
        va="center",
        fontsize=fontsize,
        fontweight="bold",
        linespacing=1.15,
        color=color,
        zorder=6
    )

    text.set_clip_path(patch)


def species_short(species):
    return species.replace("Enterobacter ", "E. ")


def style_excel(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="173352"
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )

        for i, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(i)
            ].width = min(
                46,
                max(16, len(str(cell.value)) + 3)
            )


# ============================================================
# READ FINAL MASTER
# ============================================================

print("=" * 95)
print("FIGURE 8B — COMBINED HIERARCHICAL MLST TREEMAP")
print("=" * 95)

if not MASTER.is_file():
    raise FileNotFoundError(MASTER)

columns = [
    "Assembly_Accession",
    "Normalized_Species",
    "MLST_Scheme",
    "MLST_ST",
    "MLST_Status"
]

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=columns,
    engine="openpyxl",
    dtype=str
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9,157 genomes, found {len(df)}"
    )

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicated genome accessions."
    )

df["Species"] = (
    df["Normalized_Species"]
    .fillna("")
    .astype(str)
    .str.strip()
)

df["Scheme"] = (
    df["MLST_Scheme"]
    .fillna("")
    .astype(str)
    .str.lower()
    .str.strip()
)

df["ST"] = df["MLST_ST"].map(clean_st)

df["Resolved"] = df["Species"].isin(SPECIES)
df["Valid_ST"] = df["ST"].ne("")
df["Valid_Scheme"] = df["Scheme"].eq("ecloacae")

df["Typed"] = (
    df["Resolved"]
    & df["Valid_ST"]
    & df["Valid_Scheme"]
)

resolved = df[df["Resolved"]].copy()
typed = df[df["Typed"]].copy()

species_totals = (
    resolved.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

if int(species_totals.sum()) != 8882:
    raise ValueError(
        "Expected 8,882 resolved-species genomes, "
        f"found {int(species_totals.sum())}."
    )

typed_totals = (
    typed.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

abundance = (
    typed.groupby(["Species", "ST"])
    .size()
    .reset_index(name="Genome_Count")
)

abundance["ST_Number"] = (
    abundance["ST"].str[2:].astype(int)
)

abundance = abundance.sort_values(
    ["Species", "Genome_Count", "ST_Number"],
    ascending=[True, False, True]
).reset_index(drop=True)

abundance["Rank"] = (
    abundance.groupby("Species").cumcount() + 1
)

abundance["Species_Total"] = (
    abundance["Species"].map(species_totals)
)

abundance["Typed_Total"] = (
    abundance["Species"].map(typed_totals)
)

abundance["Prevalence_Within_Species_Percent"] = (
    100.0 * abundance["Genome_Count"]
    / abundance["Species_Total"]
)

abundance["Prevalence_Among_Typed_Percent"] = (
    100.0 * abundance["Genome_Count"]
    / abundance["Typed_Total"]
)

# ============================================================
# BUILD PLOTTED ST CATEGORIES
# ============================================================

plot_rows = []
summary_rows = []

for species in SPECIES:

    sp = abundance[
        abundance["Species"].eq(species)
    ].copy()

    total = int(species_totals[species])
    typed_n = int(typed_totals[species])
    unassigned = total - typed_n

    limit = (
        TOP_ST_HORMAECHEI
        if species == "Enterobacter hormaechei"
        else TOP_ST_OTHER_SPECIES
    )

    major = sp.head(limit)
    other = sp.iloc[limit:]

    for row in major.itertuples(index=False):
        plot_rows.append({
            "Species": species,
            "ST": row.ST,
            "Genome_Count": int(row.Genome_Count),
            "Percentage_Within_Species": (
                100.0 * row.Genome_Count / total
            ),
            "Category": "Individual ST",
            "Rank": int(row.Rank)
        })

    other_n = int(other["Genome_Count"].sum())

    if other_n > 0:
        plot_rows.append({
            "Species": species,
            "ST": "Other STs",
            "Genome_Count": other_n,
            "Percentage_Within_Species": (
                100.0 * other_n / total
            ),
            "Category": "Other STs",
            "Rank": limit + 1
        })

    if unassigned > 0:
        plot_rows.append({
            "Species": species,
            "ST": "Unassigned ST",
            "Genome_Count": unassigned,
            "Percentage_Within_Species": (
                100.0 * unassigned / total
            ),
            "Category": "Unassigned ST",
            "Rank": limit + 2
        })

    summary_rows.append({
        "Species": species,
        "Total_Genomes": total,
        "Species_Percent_of_Resolved": (
            100.0 * total / len(resolved)
        ),
        "Typed_Genomes": typed_n,
        "Unassigned_ST_Genomes": unassigned,
        "Unique_STs": len(sp),
        "Individual_STs_Shown": len(major),
        "Other_STs_Count": len(other),
        "Other_STs_Genomes": other_n,
        "Displayed_Major_ST_Genomes": (
            int(major["Genome_Count"].sum())
        )
    })

plot_df = pd.DataFrame(plot_rows)
summary = pd.DataFrame(summary_rows)

# Audit the arithmetic before drawing.
for species in SPECIES:
    observed = int(
        plot_df.loc[
            plot_df["Species"].eq(species),
            "Genome_Count"
        ].sum()
    )

    expected = int(species_totals[species])

    if observed != expected:
        raise ValueError(
            f"Treemap categories do not sum for {species}: "
            f"{observed} vs {expected}"
        )

print("\nSPECIES SUMMARY")
print(summary.to_string(index=False))

# ============================================================
# CREATE ONE LANDSCAPE COMBINED TREEMAP
# ============================================================

fig = plt.figure(
    figsize=(FIG_WIDTH, FIG_HEIGHT),
    facecolor="white"
)

ax = fig.add_axes([
    0.035, 0.105, 0.93, 0.80
])

ax.set_xlim(0, 1000)
ax.set_ylim(0, 720)
ax.invert_yaxis()
ax.set_aspect("auto")
ax.axis("off")

TOTAL = int(species_totals.sum())

X0 = 0.0
Y0 = 0.0
FULL_WIDTH = 1000.0
FULL_HEIGHT = 720.0

species_rectangles = []

# Species blocks have exactly proportional area.
current_y = Y0

for group in SPECIES_ROWS:

    row_total = sum(
        int(species_totals[s])
        for s in group
    )

    row_height = (
        FULL_HEIGHT * row_total / TOTAL
    )

    current_x = X0

    for species in group:

        count = int(species_totals[species])

        block_width = (
            FULL_WIDTH * count / row_total
        )

        x = current_x
        y = current_y
        w = block_width
        h = row_height

        base = SPECIES_COLORS[species]

        # Species border.
        ax.add_patch(
            Rectangle(
                (x+2, y+2),
                w-4, h-4,
                facecolor=base,
                edgecolor="white",
                linewidth=2.3,
                joinstyle="round",
                zorder=1
            )
        )

        # Proportional header band.
        header_h = min(31.0, h*0.22)

        species_percentage = (
            100.0 * count / TOTAL
        )

        # Full species names when space permits.
        if w >= 300:
            heading = (
                f"{species}  "
                f"(n={fmt(count)}; "
                f"{species_percentage:.1f}%)"
            )
            heading_size = 14.0
        else:
            heading = (
                f"{species_short(species)} "
                f"(n={fmt(count)})"
            )
            heading_size = 10.0 if w > 180 else 8.2

        ax.text(
            x + 11,
            y + header_h*0.53,
            heading,
            ha="left",
            va="center",
            fontsize=heading_size,
            fontstyle="italic",
            fontweight="bold",
            color="white",
            zorder=6
        )

        # Nested ST treemap region.
        px = x + 8
        py = y + header_h + 2
        pw = max(1.0, w - 16)
        ph = max(1.0, h - header_h - 10)

        part = plot_df[
            plot_df["Species"].eq(species)
        ].copy()

        # Sort for proper squarified rectangle placement.
        part = part.sort_values(
            ["Genome_Count", "Rank"],
            ascending=[False, True]
        ).reset_index(drop=True)

        sizes = part["Genome_Count"].tolist()

        normalized_sizes = squarify.normalize_sizes(
            sizes,
            pw,
            ph
        )

        rects = squarify.squarify(
            normalized_sizes,
            px,
            py,
            pw,
            ph
        )

        for i, (rectangle, row) in enumerate(
            zip(rects, part.itertuples(index=False))
        ):

            rx = rectangle["x"]
            ry = rectangle["y"]
            rw = rectangle["dx"]
            rh = rectangle["dy"]

            if row.Category == "Unassigned ST":
                face = "#D8DEE5"
            elif row.Category == "Other STs":
                face = blend(base, 0.72)
            else:
                amount = min(0.73, i*0.065)
                face = blend(base, amount)

            patch = Rectangle(
                (rx+1, ry+1),
                max(0, rw-2),
                max(0, rh-2),
                facecolor=face,
                edgecolor="white",
                linewidth=1.45,
                zorder=3
            )
            ax.add_patch(patch)

            draw_fitted_label(
                ax,
                patch,
                rx, ry, rw, rh,
                row.ST,
                row.Genome_Count,
                row.Percentage_Within_Species,
                face
            )

        species_rectangles.append({
            "Species": species,
            "X": x,
            "Y": y,
            "Width": w,
            "Height": h,
            "Total_Genomes": count,
            "Percentage_of_Resolved": species_percentage
        })

        current_x += block_width

    current_y += row_height

# ============================================================
# HEADER, LEGEND AND FOOTER
# ============================================================

fig.text(
    0.5,
    0.967,
    "Combined hierarchical MLST abundance treemap",
    ha="center",
    va="center",
    fontsize=22,
    fontweight="bold",
    color="#173352"
)

fig.text(
    0.5,
    0.938,
    (
        "Enterobacter cloacae complex"
        f"  |  {TOTAL:,} resolved-species genomes"
    ),
    ha="center",
    fontsize=12.5,
    color="#53687C"
)

fig.text(
    0.5,
    0.073,
    "Species block area = species genome count. "
    "Nested ST tile area = number of genomes assigned "
    "to each sequence type.",
    ha="center",
    fontsize=10.0,
    color="#344B60"
)

fig.text(
    0.5,
    0.053,
    "Percentages are calculated within each species. "
    "Rare STs are grouped as Other STs; unassigned "
    "STs are retained separately.",
    ha="center",
    fontsize=9.3,
    color="#566A7D"
)

fig.text(
    0.5,
    0.034,
    "Seven resolved ECC species; unresolved ECC genomes "
    "are not included in this treemap.",
    ha="center",
    fontsize=8.8,
    color="#66798B"
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 94,
        "subsampling": 1
    }
)

plt.close(fig)
gc.collect()

print("\n[SAVED FIGURE]", JPG)

# ============================================================
# EXPORT COMPLETE EXCEL
# ============================================================

qc = pd.DataFrame([
    ["Total genomes in master", len(df)],
    ["Resolved seven-species genomes", len(resolved)],
    ["Unresolved/other genomes", len(df)-len(resolved)],
    ["Valid numeric-ST genomes", len(typed)],
    ["MLST scheme", "ecloacae"],
    ["Figure layout", "One combined landscape treemap"],
    ["Species block area", "Total species genome count"],
    ["ST tile area", "Genomes assigned to that ST"],
    ["ST percentage denominator", "All genomes of that species"],
    ["Other STs", "Rare assigned STs pooled for visualization"],
    ["Unassigned STs", "Separate tile, not included in Other STs"],
    ["Top STs, hormaechei", TOP_ST_HORMAECHEI],
    ["Top STs, other species", TOP_ST_OTHER_SPECIES],
    ["Interpretation", "Sequence type abundance, not serotype"],
    ["ST caveat", "Same ST number across species is not proof of identical lineage"],
], columns=["Parameter", "Description"])

with pd.ExcelWriter(
    EXCEL,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="ALL_9157_MLST_QC",
        index=False
    )

    abundance.to_excel(
        writer,
        sheet_name="ALL_SPECIES_ST_ABUNDANCE",
        index=False
    )

    plot_df.to_excel(
        writer,
        sheet_name="PLOTTED_ST_TILES",
        index=False
    )

    summary.to_excel(
        writer,
        sheet_name="SEVEN_SPECIES_SUMMARY",
        index=False
    )

    pd.DataFrame(
        species_rectangles
    ).to_excel(
        writer,
        sheet_name="SPECIES_BLOCK_LAYOUT",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    style_excel(writer)

print("\n" + "=" * 95)
print("FIGURE 8B COMPLETED")
print("=" * 95)
print("Master genomes:", len(df))
print("Resolved species genomes:", len(resolved))
print("Valid assigned ST genomes:", len(typed))
print("Species represented:", len(SPECIES))
print("JPEG:", JPG)
print("Excel:", EXCEL)
print("=" * 95)
