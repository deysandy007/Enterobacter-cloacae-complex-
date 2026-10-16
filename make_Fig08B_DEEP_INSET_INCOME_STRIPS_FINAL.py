
#!/usr/bin/env python3

# ================================================================
# FIGURE 8B — COMBINED ECC MLST ABUNDANCE TREEMAP
# FINAL DESIGN: DEEP INSET COUNTRY-INCOME COLOR STRIPS
#
# DATA:
# 9,157 ECC genomes from the final master database
# Seven resolved Enterobacter cloacae complex species
#
# FIGURE:
# - ONE combined landscape treemap
# - Every valid ST receives an individual rectangle
# - No "Other STs" grouping
# - Species-specific background colors remain unchanged
# - DEEP income strips INSIDE the individual ST boxes
# - Income segment width proportional to genome counts
# - ST labels positioned above the income strip
# - Thin outlines for cross-income STs
#
# OUTPUT:
# 1 publication-quality 600-DPI JPEG
# 1 complete Excel workbook
#
# ================================================================

from pathlib import Path
import re
import gc

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt

from matplotlib.patches import Rectangle, Patch
from matplotlib.colors import to_rgb

import squarify

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# ================================================================
# 1. INPUT AND OUTPUT PATHS
# ================================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT
    / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

INCOME_DB = (
    ROOT
    / "Fig_01A_Global_Geographical_Distribution"
    / "Fig_01A_Country_Income_Database.xlsx"
)

OUT = (
    ROOT
    / "Fig_08B_Combined_MLST_Treemap"
    / "DEEP_INSET_INCOME_STRIPS_FINAL"
)

OUT.mkdir(parents=True, exist_ok=True)

JPG = (
    OUT
    / "Fig08B_Combined_All_ST_Deep_Income_Strips_600DPI.jpg"
)

XLSX = (
    OUT
    / "Fig08B_Combined_All_ST_Deep_Income_Complete_Data.xlsx"
)


# ================================================================
# 2. FIGURE SETTINGS
# ================================================================

DPI = 600

FIG_WIDTH = 24.0
FIG_HEIGHT = 16.7

CANVAS_WIDTH = 1000.0
CANVAS_HEIGHT = 720.0

# INSET STRIP CONTROLS
# The strip appears inside the ST rectangle,
# near the bottom, with a visible margin.

STRIP_HEIGHT_FRACTION = 0.17
STRIP_LEFT_RIGHT_MARGIN = 0.045
STRIP_BOTTOM_MARGIN = 0.055

# IMPORTANT:
# 0.00 means completely saturated income colors.
# Only income strips use this setting.
# Species backgrounds will not change.

STRIP_FADE = 0.00


# ================================================================
# 3. ECC SPECIES AND THEIR ORIGINAL COLORS
# ================================================================

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

# Maintain the previous combined treemap arrangement.

SPECIES_ROWS = [
    [
        "Enterobacter hormaechei"
    ],
    [
        "Enterobacter cloacae",
        "Enterobacter roggenkampii",
        "Enterobacter asburiae"
    ],
    [
        "Enterobacter kobei",
        "Enterobacter ludwigii",
        "Enterobacter bugandensis"
    ],
]


# ================================================================
# 4. DEEP COUNTRY-INCOME COLORS
# ================================================================

INCOME_GROUPS = [
    "H",
    "M",
    "L+LM",
    "Unclassified"
]

INCOME_COLORS = {
    "H": "#9B6100",
    "M": "#461080",
    "L+LM": "#A51F55",
    "Unclassified": "#58636F"
}

INCOME_LABELS = {
    "H": "High income",
    "M": "Upper-middle income",
    "L+LM": "Lower-middle + low income",
    "Unclassified": "Income unclassified"
}

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})


# ================================================================
# 5. HELPER FUNCTIONS
# ================================================================

def clean_st(value):

    if pd.isna(value):
        return ""

    text = (
        str(value)
        .strip()
        .strip("'\"")
        .strip()
    )

    if text.upper().startswith("ST"):
        text = text[2:]

    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]

    if re.fullmatch(r"\d+", text):
        number = int(text)

        if number > 0:
            return f"ST{number}"

    return ""


def clean_country(value):

    if pd.isna(value):
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip()
    ).casefold()


def clean_income(value):

    if pd.isna(value):
        return "Unclassified"

    text = re.sub(
        r"[\s_\-]+",
        "",
        str(value).upper().strip()
    )

    mapping = {
        "H": "H",
        "HIGH": "H",
        "HIGHINCOME": "H",
        "HIGHINCOMECOUNTRIES": "H",

        "M": "M",
        "UM": "M",
        "UMIC": "M",
        "UPPERMIDDLE": "M",
        "UPPERMIDDLEINCOME": "M",
        "UPPERMIDDLEINCOMECOUNTRIES": "M",

        "L+LM": "L+LM",
        "LM+L": "L+LM",
        "LLM": "L+LM",
        "LMICLIC": "L+LM",

        "LOWERMIDDLE+LOW": "L+LM",
        "LOW+LOWERMIDDLE": "L+LM",
        "LOWERMIDDLEANDLOW": "L+LM",
        "LOWERMIDDLEANDLOWINCOME": "L+LM",

        "LOWERMIDDLE": "L+LM",
        "LOWERMIDDLEINCOME": "L+LM",
        "LOW": "L+LM",
        "LOWINCOME": "L+LM",
    }

    return mapping.get(
        text,
        "Unclassified"
    )


def blend_color(color, fade):

    rgb = np.asarray(
        to_rgb(color),
        dtype=float
    )

    return tuple(
        np.clip(
            rgb * (1.0 - fade) + fade,
            0.0,
            1.0
        )
    )


def text_color(face):

    red, green, blue = to_rgb(face)

    luminance = (
        0.2126 * red
        + 0.7152 * green
        + 0.0722 * blue
    )

    if luminance < 0.46:
        return "white"

    return "#173148"


def format_integer(value):

    return f"{int(value):,}"


def style_excel(writer):

    for worksheet in writer.book.worksheets:

        worksheet.freeze_panes = "A2"

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )

        for cell in worksheet[1]:

            cell.fill = PatternFill(
                "solid",
                fgColor="173352"
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

        for column_number, cell in enumerate(
            worksheet[1], start=1
        ):

            worksheet.column_dimensions[
                get_column_letter(column_number)
            ].width = min(
                46,
                max(
                    17,
                    len(str(cell.value)) + 4
                )
            )


# ================================================================
# 6. LOAD THE 9,157-GENOME MASTER DATABASE
# ================================================================

print("=" * 100)
print("FIGURE 8B — DEEP INSET INCOME-STRIP TREEMAP")
print("=" * 100)

for required_file in [MASTER, INCOME_DB]:

    if not required_file.is_file():

        raise FileNotFoundError(
            required_file
        )

columns = [
    "Assembly_Accession",
    "Normalized_Species",
    "MLST_ST",
    "MLST_Scheme",
    "Country_Harmonized"
]

data = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=columns,
    dtype=str,
    engine="openpyxl"
)

if len(data) != 9157:

    raise ValueError(
        "Expected 9157 genomes; "
        f"found {len(data)}"
    )

if (
    data["Assembly_Accession"].isna().any()
    or data["Assembly_Accession"].duplicated().any()
):

    raise ValueError(
        "Missing or duplicated assembly accessions."
    )

data["Species"] = (
    data["Normalized_Species"]
    .fillna("")
    .str.strip()
)

data["ST"] = (
    data["MLST_ST"]
    .map(clean_st)
)

data["Scheme"] = (
    data["MLST_Scheme"]
    .fillna("")
    .str.strip()
    .str.lower()
)

data["Country_Key"] = (
    data["Country_Harmonized"]
    .map(clean_country)
)

data["Resolved_Species"] = (
    data["Species"].isin(SPECIES)
)

data["Valid_ST"] = (
    data["ST"].ne("")
    & data["Scheme"].eq("ecloacae")
)

data["Eligible"] = (
    data["Resolved_Species"]
    & data["Valid_ST"]
)


# ================================================================
# 7. LOAD THE FROZEN COUNTRY-INCOME DATABASE
# ================================================================

income_data = pd.read_excel(
    INCOME_DB,
    sheet_name="COUNTRY_INCOME_COUNTS",
    engine="openpyxl",
    dtype=str
)

for required_column in [
    "Country_Harmonized",
    "Income_Group"
]:

    if required_column not in income_data.columns:

        raise ValueError(
            "Missing income column: "
            + required_column
        )

income_data["Country_Key"] = (
    income_data["Country_Harmonized"]
    .map(clean_country)
)

income_data["Income"] = (
    income_data["Income_Group"]
    .map(clean_income)
)

income_data = income_data[
    income_data["Country_Key"].ne("")
].copy()

conflicts = (
    income_data
    .groupby("Country_Key")["Income"]
    .nunique()
)

if (conflicts > 1).any():

    conflicting_countries = (
        conflicts[conflicts > 1]
        .index.tolist()
    )

    raise ValueError(
        "Conflicting income classifications: "
        + str(conflicting_countries[:15])
    )

income_lookup = (
    income_data
    .drop_duplicates("Country_Key")
    .set_index("Country_Key")["Income"]
)

data["Income"] = (
    data["Country_Key"]
    .map(income_lookup)
    .fillna("Unclassified")
)

resolved = data[
    data["Resolved_Species"]
].copy()

typed = data[
    data["Eligible"]
].copy()

if len(resolved) != 8882:

    raise ValueError(
        "Expected 8882 resolved genomes; "
        f"found {len(resolved)}"
    )

if typed.empty:

    raise ValueError(
        "No valid ST-assigned genomes."
    )

print("Total master genomes:", len(data))
print("Resolved genomes:", len(resolved))
print("ST-assigned genomes:", len(typed))

print("\nIncome distribution:")

print(
    typed["Income"]
    .value_counts()
    .to_string()
)


# ================================================================
# 8. COUNT EVERY ST AND ITS INCOME COMPOSITION
# ================================================================

species_all_counts = (
    resolved
    .groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

species_typed_counts = (
    typed
    .groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

st_counts = (
    typed
    .groupby(["Species", "ST"])
    .size()
    .rename("Genome_Count")
    .reset_index()
)

income_counts = (
    typed
    .groupby([
        "Species",
        "ST",
        "Income"
    ])
    .size()
    .unstack(
        "Income",
        fill_value=0
    )
    .reindex(
        columns=INCOME_GROUPS,
        fill_value=0
    )
    .reset_index()
)

income_counts = income_counts.rename(
    columns={
        "H": "N_H",
        "M": "N_M",
        "L+LM": "N_LLM",
        "Unclassified": "N_Unclassified"
    }
)

st_table = st_counts.merge(
    income_counts,
    on=["Species", "ST"],
    validate="one_to_one"
)

st_table["Species_Total"] = (
    st_table["Species"]
    .map(species_all_counts)
)

st_table["Species_Typed_Total"] = (
    st_table["Species"]
    .map(species_typed_counts)
)

st_table["Prevalence_Within_Species"] = (
    100.0
    * st_table["Genome_Count"]
    / st_table["Species_Total"]
)

st_table["Prevalence_Within_Typed"] = (
    100.0
    * st_table["Genome_Count"]
    / st_table["Species_Typed_Total"]
)

st_table["ST_Number"] = (
    st_table["ST"]
    .str[2:]
    .astype(int)
)

income_count_columns = [
    "N_H",
    "N_M",
    "N_LLM",
    "N_Unclassified"
]

for group, column in zip(
    INCOME_GROUPS,
    income_count_columns
):

    st_table[
        "Percent_" + group
    ] = (
        100.0
        * st_table[column]
        / st_table["Genome_Count"]
    )

st_table["Income_Breadth"] = (
    (
        st_table[
            ["N_H", "N_M", "N_LLM"]
        ] > 0
    ).sum(axis=1)
)

st_table["Cross_Income"] = (
    st_table["Income_Breadth"] >= 2
)

st_table = (
    st_table
    .sort_values(
        [
            "Species",
            "Genome_Count",
            "ST_Number"
        ],
        ascending=[
            True,
            False,
            True
        ]
    )
    .reset_index(drop=True)
)

st_table["Abundance_Rank"] = (
    st_table
    .groupby("Species")
    .cumcount()
    + 1
)

# Validate actual income counts.

if not (
    st_table[income_count_columns].sum(axis=1)
    == st_table["Genome_Count"]
).all():

    raise ValueError(
        "Income count totals do not match ST counts."
    )

if int(st_table["Genome_Count"].sum()) != len(typed):

    raise ValueError(
        "ST genome counts do not match typed genomes."
    )


# ================================================================
# 9. CREATE THE COMBINED TREEMAP CANVAS
# ================================================================

fig = plt.figure(
    figsize=(
        FIG_WIDTH,
        FIG_HEIGHT
    ),
    facecolor="white"
)

AX_LEFT = 0.026
AX_BOTTOM = 0.138
AX_WIDTH = 0.948
AX_HEIGHT = 0.778

ax = fig.add_axes([
    AX_LEFT,
    AX_BOTTOM,
    AX_WIDTH,
    AX_HEIGHT
])

ax_width_inches = (
    FIG_WIDTH * AX_WIDTH
)

ax_height_inches = (
    FIG_HEIGHT * AX_HEIGHT
)

ax.set_xlim(
    0,
    CANVAS_WIDTH
)

ax.set_ylim(
    0,
    CANVAS_HEIGHT
)

ax.invert_yaxis()
ax.set_aspect("auto")
ax.axis("off")

tile_records = []
species_records = []

TOTAL_TYPED = len(typed)


# ================================================================
# 10. DRAW DEEP INSET INCOME STRIPS
# ================================================================

def draw_income_strip(x, y, w, h, row, tile_patch):

    # Small inset margins preserve the visual appearance
    # of the species-colored ST rectangle.

    margin_x = (
        w * STRIP_LEFT_RIGHT_MARGIN
    )

    margin_bottom = (
        h * STRIP_BOTTOM_MARGIN
    )

    strip_width = (
        w - 2.0 * margin_x
    )

    strip_height = (
        h * STRIP_HEIGHT_FRACTION
    )

    strip_x = x + margin_x

    strip_y = (
        y + h
        - margin_bottom
        - strip_height
    )

    n_by_income = [
        int(row.N_H),
        int(row.N_M),
        int(row.N_LLM),
        int(row.N_Unclassified)
    ]

    total = int(row.Genome_Count)

    cursor = strip_x

    segments_drawn = 0

    for income, count in zip(
        INCOME_GROUPS,
        n_by_income
    ):

        if count == 0:
            continue

        segment_width = (
            strip_width
            * count
            / total
        )

        # Deep income colors, without fading.
        segment_face = blend_color(
            INCOME_COLORS[income],
            STRIP_FADE
        )

        piece = Rectangle(
            (cursor, strip_y),
            segment_width,
            strip_height,
            facecolor=segment_face,
            edgecolor="none",
            linewidth=0,
            antialiased=False,
            zorder=6
        )

        piece.set_clip_path(
            tile_patch
        )

        ax.add_patch(piece)

        cursor += segment_width
        segments_drawn += 1

    # Thin, subtle inset-strip outline.
    # Applied only where physically visible.

    physical_width = (
        strip_width
        / CANVAS_WIDTH
        * ax_width_inches
    )

    physical_height = (
        strip_height
        / CANVAS_HEIGHT
        * ax_height_inches
    )

    if (
        physical_width >= 0.14
        and physical_height >= 0.035
    ):

        outline = Rectangle(
            (
                strip_x,
                strip_y
            ),
            strip_width,
            strip_height,
            fill=False,
            edgecolor="#F0F2F4",
            linewidth=0.35,
            zorder=7
        )

        outline.set_clip_path(
            tile_patch
        )

        ax.add_patch(outline)

    return strip_y, segments_drawn


# ================================================================
# 11. DRAW ST LABELS ABOVE INSET STRIPS
# ================================================================

def draw_st_label(
    tile_patch,
    x, y, w, h,
    strip_y,
    row,
    face
):

    physical_width = (
        w / CANVAS_WIDTH
        * ax_width_inches
    )

    usable_height = max(
        0.0,
        strip_y - y
    )

    physical_height = (
        usable_height
        / CANVAS_HEIGHT
        * ax_height_inches
    )

    label = None
    fontsize = 0

    if (
        physical_width >= 0.90
        and physical_height >= 0.56
    ):

        label = (
            f"{row.ST}\n"
            f"n={int(row.Genome_Count):,}\n"
            f"{row.Prevalence_Within_Species:.1f}%"
        )

        fontsize = min(
            12.3,
            physical_width * 10,
            physical_height * 13
        )

    elif (
        physical_width >= 0.57
        and physical_height >= 0.32
    ):

        label = (
            f"{row.ST}\n"
            f"n={int(row.Genome_Count):,}"
        )

        fontsize = min(
            9.7,
            physical_width * 10.5,
            physical_height * 16
        )

    elif (
        physical_width >= 0.36
        and physical_height >= 0.17
    ):

        label = row.ST

        fontsize = min(
            7.7,
            physical_width * 10.8,
            physical_height * 21
        )

    if (
        label is None
        or fontsize < 5.5
    ):

        return "Too small to label"

    artist = ax.text(
        x + w/2,
        y + usable_height/2,
        label,
        ha="center",
        va="center",
        fontsize=fontsize,
        fontweight="bold",
        color=text_color(face),
        linespacing=1.12,
        zorder=9
    )

    artist.set_clip_path(
        tile_patch
    )

    return "Labeled"


# ================================================================
# 12. DRAW ALL SEVEN SPECIES AND EVERY INDIVIDUAL ST
# ================================================================

current_y = 0.0

for species_row in SPECIES_ROWS:

    row_genomes = sum(
        int(species_typed_counts[s])
        for s in species_row
    )

    if row_genomes == 0:
        continue

    row_height = (
        CANVAS_HEIGHT
        * row_genomes
        / TOTAL_TYPED
    )

    current_x = 0.0

    for species in species_row:

        typed_genomes = int(
            species_typed_counts[species]
        )

        all_genomes = int(
            species_all_counts[species]
        )

        if typed_genomes == 0:
            continue

        block_width = (
            CANVAS_WIDTH
            * typed_genomes
            / row_genomes
        )

        block_height = row_height

        x = current_x
        y = current_y

        species_color = (
            SPECIES_COLORS[species]
        )

        # Original species colors are preserved.
        ax.add_patch(
            Rectangle(
                (
                    x+1,
                    y+1
                ),
                block_width-2,
                block_height-2,
                facecolor=species_color,
                edgecolor="white",
                linewidth=2.0,
                zorder=1
            )
        )

        header_height = min(
            30.0,
            block_height*0.22
        )

        species_prevalence = (
            100.0
            * typed_genomes
            / TOTAL_TYPED
        )

        if block_width >= 340:

            species_label = (
                f"{species} "
                f"(n={typed_genomes:,}; "
                f"{species_prevalence:.1f}%)"
            )

            species_font = 13

        else:

            species_label = (
                species.replace(
                    "Enterobacter ",
                    "E. "
                )
                + f" (n={typed_genomes:,})"
            )

            species_font = (
                10 if block_width > 180
                else 8
            )

        ax.text(
            x+8,
            y+header_height*0.54,
            species_label,
            fontsize=species_font,
            fontweight="bold",
            fontstyle="italic",
            color="white",
            va="center",
            ha="left",
            zorder=12
        )

        # Area available for individual ST tiles.
        tile_x = x+6
        tile_y = y+header_height+2

        tile_width = max(
            0.1,
            block_width-12
        )

        tile_height = max(
            0.1,
            block_height-header_height-8
        )

        species_st = (
            st_table[
                st_table["Species"].eq(species)
            ]
            .sort_values(
                [
                    "Genome_Count",
                    "ST_Number"
                ],
                ascending=[
                    False,
                    True
                ]
            )
            .reset_index(drop=True)
        )

        sizes = (
            species_st["Genome_Count"]
            .astype(float)
            .tolist()
        )

        rectangles = squarify.squarify(
            squarify.normalize_sizes(
                sizes,
                tile_width,
                tile_height
            ),
            tile_x,
            tile_y,
            tile_width,
            tile_height
        )

        if len(rectangles) != len(species_st):

            raise ValueError(
                "Not every ST was assigned a rectangle "
                f"for {species}"
            )

        for i, (geometry, row) in enumerate(
            zip(
                rectangles,
                species_st.itertuples(index=False)
            )
        ):

            rx = float(geometry["x"])
            ry = float(geometry["y"])

            rw = float(geometry["dx"])
            rh = float(geometry["dy"])

            # Preserve the existing lighter and darker
            # shades of the species background.

            shade = (
                0.14
                + 0.64 * i
                / max(
                    1,
                    len(species_st)-1
                )
            )

            face = blend_color(
                species_color,
                shade
            )

            if row.Income_Breadth >= 2:

                border_color = "#263B4B"

            else:

                border_color = "white"

            line_width = (
                0.10
                if min(rw, rh) < 1
                else 0.40
            )

            tile_patch = Rectangle(
                (rx, ry),
                rw,
                rh,
                facecolor=face,
                edgecolor=border_color,
                linewidth=line_width,
                zorder=2
            )

            ax.add_patch(
                tile_patch
            )

            # Apply the inset income strip.
            strip_y, n_segments = draw_income_strip(
                rx,
                ry,
                rw,
                rh,
                row,
                tile_patch
            )

            # ST label remains above the colored strip.
            label_status = draw_st_label(
                tile_patch,
                rx,
                ry,
                rw,
                rh,
                strip_y,
                row,
                face
            )

            tile_records.append({
                "Species": species,
                "ST": row.ST,
                "Genome_Count": int(
                    row.Genome_Count
                ),
                "Prevalence_Within_Species": float(
                    row.Prevalence_Within_Species
                ),
                "N_H": int(row.N_H),
                "N_M": int(row.N_M),
                "N_LLM": int(row.N_LLM),
                "N_Unclassified": int(
                    row.N_Unclassified
                ),
                "Income_Breadth": int(
                    row.Income_Breadth
                ),
                "Cross_Income": bool(
                    row.Cross_Income
                ),
                "Income_Segments_Drawn": n_segments,
                "Inset_Strip": True,
                "Label_Status": label_status,
                "Tile_X": rx,
                "Tile_Y": ry,
                "Tile_Width": rw,
                "Tile_Height": rh
            })

        species_records.append({
            "Species": species,
            "All_Genomes": all_genomes,
            "ST_Assigned_Genomes": typed_genomes,
            "Unassigned_ST_Genomes": (
                all_genomes - typed_genomes
            ),
            "Unique_STs": len(
                species_st
            ),
            "Species_Proportion_Percent": (
                species_prevalence
            ),
            "Block_X": x,
            "Block_Y": y,
            "Block_Width": block_width,
            "Block_Height": block_height
        })

        current_x += block_width

    current_y += row_height


# ================================================================
# 13. FIGURE TITLE AND LEGEND
# ================================================================

fig.text(
    0.5,
    0.973,
    "Combined MLST abundance and country-income composition",
    ha="center",
    fontsize=21,
    fontweight="bold",
    color="#173352"
)

fig.text(
    0.5,
    0.943,
    (
        "Enterobacter cloacae complex | "
        f"{TOTAL_TYPED:,} ST-assigned genomes"
    ),
    ha="center",
    fontsize=12,
    color="#516779"
)

legend_handles = [
    Patch(
        facecolor=blend_color(
            INCOME_COLORS[group],
            STRIP_FADE
        ),
        edgecolor="none",
        label=INCOME_LABELS[group]
    )
    for group in INCOME_GROUPS
]

fig.legend(
    handles=legend_handles,
    loc="lower center",
    bbox_to_anchor=(
        0.5,
        0.098
    ),
    ncol=4,
    frameon=False,
    fontsize=10,
    columnspacing=2,
    handlelength=1.8
)

fig.text(
    0.5,
    0.072,
    "Each rectangle represents one ST; deep inset "
    "color strips show the income composition of that ST.",
    ha="center",
    fontsize=10,
    color="#394F63"
)

fig.text(
    0.5,
    0.050,
    "Species colors remain unchanged. "
    "Rectangle areas are proportional to genome counts.",
    ha="center",
    fontsize=9.4,
    color="#51677A"
)

fig.text(
    0.5,
    0.030,
    "Every valid ST is retained separately. "
    "For very small rectangles, exact income counts "
    "are available in the Excel workbook.",
    ha="center",
    fontsize=8.8,
    color="#65798A"
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 94,
        "subsampling": 1
    }
)

plt.close(fig)
gc.collect()

print("[SAVED FIGURE]", JPG)


# ================================================================
# 14. VALIDATION
# ================================================================

tiles = pd.DataFrame(
    tile_records
)

species_summary = pd.DataFrame(
    species_records
)

if len(tiles) != len(st_table):

    raise ValueError(
        "Some STs were omitted."
    )

if tiles.duplicated([
    "Species",
    "ST"
]).any():

    raise ValueError(
        "Duplicate ST rectangles detected."
    )

if int(tiles["Genome_Count"].sum()) != TOTAL_TYPED:

    raise ValueError(
        "Treemap counts do not match assigned genomes."
    )

if not tiles["Inset_Strip"].all():

    raise ValueError(
        "At least one ST did not receive an inset strip."
    )


# ================================================================
# 15. COMPLETE EXCEL OUTPUT
# ================================================================

genome_export = data[[
    "Assembly_Accession",
    "Species",
    "ST",
    "Scheme",
    "Country_Harmonized",
    "Income",
    "Resolved_Species",
    "Valid_ST",
    "Eligible"
]].copy()

species_income = (
    typed
    .groupby([
        "Species",
        "Income"
    ])
    .size()
    .unstack(
        "Income",
        fill_value=0
    )
    .reindex(
        index=SPECIES,
        columns=INCOME_GROUPS,
        fill_value=0
    )
    .reset_index()
)

cross_income = st_table[
    st_table["Cross_Income"]
].copy()

qc = pd.DataFrame([
    [
        "Master genomes",
        len(data)
    ],
    [
        "Seven resolved-species genomes",
        len(resolved)
    ],
    [
        "ST-assigned genomes",
        TOTAL_TYPED
    ],
    [
        "Unassigned ST genomes",
        len(resolved)-TOTAL_TYPED
    ],
    [
        "Distinct species-ST combinations",
        len(st_table)
    ],
    [
        "Individual ST rectangles",
        len(tiles)
    ],
    [
        "Inset income strips",
        int(tiles["Inset_Strip"].sum())
    ],
    [
        "STs grouped into Other",
        0
    ],
    [
        "Cross-income STs",
        len(cross_income)
    ],
    [
        "Income-strip fade",
        STRIP_FADE
    ],
    [
        "Income-strip colors",
        str(INCOME_COLORS)
    ],
    [
        "Species colors",
        "Original species palette preserved"
    ],
    [
        "Tile area",
        "ST-assigned genome count"
    ],
    [
        "Income-strip width",
        "Within-ST genome proportion"
    ],
    [
        "Income categories",
        "H, M, L+LM, Unclassified"
    ],
    [
        "MLST scheme",
        "ecloacae"
    ],
    [
        "Figure",
        "Single landscape JPEG, 600 DPI"
    ],
    [
        "Interpretation",
        "Cross-income ST occurrence does not prove transmission"
    ]
], columns=[
    "Parameter",
    "Value"
])

with pd.ExcelWriter(
    XLSX,
    engine="openpyxl"
) as writer:

    genome_export.to_excel(
        writer,
        sheet_name="ALL_9157_GENOMES",
        index=False
    )

    st_table.to_excel(
        writer,
        sheet_name="ALL_ST_INCOME_COUNTS",
        index=False
    )

    tiles.to_excel(
        writer,
        sheet_name="EVERY_ST_TILE",
        index=False
    )

    cross_income.to_excel(
        writer,
        sheet_name="CROSS_INCOME_STS",
        index=False
    )

    species_income.to_excel(
        writer,
        sheet_name="SPECIES_INCOME",
        index=False
    )

    species_summary.to_excel(
        writer,
        sheet_name="SPECIES_SUMMARY",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    style_excel(writer)


# ================================================================
# 16. FINAL LOG
# ================================================================

print()
print("=" * 100)
print("FIGURE 8B — COMPLETED")
print("=" * 100)

print("Total master genomes:", len(data))
print("Resolved genomes:", len(resolved))
print("ST-assigned genomes:", TOTAL_TYPED)

print("Individual ST rectangles:", len(tiles))
print("Inset income strips:", len(tiles))
print("Cross-income STs:", len(cross_income))

print("STs grouped:", 0)
print("Strip fade:", STRIP_FADE)

print()
print("JPEG:", JPG)
print("Excel:", XLSX)

print()
print("SPECIES SUMMARY:")
print(
    species_summary[
        [
            "Species",
            "All_Genomes",
            "ST_Assigned_Genomes",
            "Unassigned_ST_Genomes",
            "Unique_STs"
        ]
    ].to_string(index=False)
)

print("=" * 100)
