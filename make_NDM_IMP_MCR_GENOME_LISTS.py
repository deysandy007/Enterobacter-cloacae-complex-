
#!/usr/bin/env python3

# ============================================================
# ECC NDM / IMP / MCR GENOME AND VARIANT EXTRACTION
#
# INPUT:
# ECC_FINAL_2000_2026_MASTER.xlsx (9,157 genomes)
#
# OUTPUT:
# - Master XLSX of gene-positive genomes
# - All variants and country-level counts
# - Unique NCBI assembly accession lists
# - Country-specific accession lists
# - Source- and year-stratified summaries
# - Multi-family co-carriage list
# - FASTA path verification
#
# No fixed gene variants or top-N filtering.
# ============================================================

from pathlib import Path
from collections import Counter
import re
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

OUT = ROOT / "Fig_10_NDM_IMP_MCR_VARIATION"
OUT.mkdir(parents=True, exist_ok=True)

LIST_DIR = OUT / "NCBI_ACCESSION_LISTS"
LIST_DIR.mkdir(exist_ok=True)

COUNTRY_DIR = OUT / "COUNTRY_SPECIFIC_LISTS"
COUNTRY_DIR.mkdir(exist_ok=True)

EXCEL = OUT / "ECC_NDM_IMP_MCR_GENOME_MASTER.xlsx"

FAMILIES = ["NDM", "IMP", "MCR"]

# Match whole gene names instead of substring fragments.
NDM = re.compile(
    r"^bla[\s_-]*NDM[\s_-]*(\d+[A-Za-z]?)$",
    re.I
)

IMP = re.compile(
    r"^bla[\s_-]*IMP[\s_-]*(\d+[A-Za-z]?)$",
    re.I
)

MCR = re.compile(
    r"^mcr[\s_-]*(10|[1-9])((?:[._-]\d+)*)$",
    re.I
)


def normalize_gene(gene):
    gene = str(gene).strip().strip("'\"")

    match = NDM.fullmatch(gene)
    if match:
        return "NDM", "blaNDM-" + match.group(1).upper()

    match = IMP.fullmatch(gene)
    if match:
        return "IMP", "blaIMP-" + match.group(1).upper()

    match = MCR.fullmatch(gene)
    if match:
        suffix = match.group(2)
        suffix = suffix.replace("_", ".").replace("-", ".")
        return "MCR", "mcr-" + match.group(1) + suffix

    return None, None


def identify_genes(value):
    result = {f: set() for f in FAMILIES}
    review = set()

    text = "" if pd.isna(value) else str(value)

    for raw in re.split(r"[;,]", text):
        gene = raw.strip()
        if not gene:
            continue

        family, allele = normalize_gene(gene)

        if family is not None:
            result[family].add(allele)

        elif re.match(
            r"^(?:bla[\s_-]*(?:NDM|IMP)|mcr[\s_-]*\d)",
            gene, re.I
        ):
            review.add(gene)

    return result, sorted(review)


def filename_safe(value):
    name = re.sub(
        r"[^A-Za-z0-9._-]+",
        "_",
        str(value).strip()
    )
    return name.strip("._")[:100] or "UNKNOWN"


def save_list(path, values):
    unique = sorted({
        str(x).strip()
        for x in values
        if pd.notna(x) and str(x).strip()
    })

    path.write_text(
        "\n".join(unique) + ("\n" if unique else ""),
        encoding="utf-8"
    )

    return len(unique)


def summarize(df, columns):
    return (
        df.groupby(columns, dropna=False)
        .agg(
            Genome_Count=(
                "Assembly_Accession", "nunique"
            )
        )
        .reset_index()
        .sort_values(
            "Genome_Count", ascending=False
        )
    )


def style_excel(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid", fgColor="173352"
            )
            cell.font = Font(
                bold=True, color="FFFFFF"
            )

        for j, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(j)
            ].width = min(
                44, max(16, len(str(cell.value)) + 3)
            )


def main():
    print("=" * 100)
    print("ECC NDM / IMP / MCR GENOME EXTRACTION")
    print("=" * 100)

    if not MASTER.is_file():
        raise FileNotFoundError(MASTER)

    required = [
        "Assembly_Accession",
        "AMR_Genes",
        "Normalized_Species",
        "Country_Harmonized",
        "MLST_ST",
        "Final_Source_Category",
        "Collection_Year_Numeric",
    ]

    optional = [
        "Original_NCBI_Accession",
        "Genome_File",
        "BioSample",
        "BioProject",
        "AMRFinder_Status",
        "MLST_Scheme",
        "Raw_Isolation_Source",
        "Raw_Host",
        "Isolation_Source",
        "Host",
    ]

    header = pd.read_excel(
        MASTER,
        sheet_name="FINAL_MASTER_2000_2026",
        nrows=0
    ).columns.tolist()

    missing = [c for c in required if c not in header]
    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))

    data = pd.read_excel(
        MASTER,
        sheet_name="FINAL_MASTER_2000_2026",
        usecols=required + [
            c for c in optional if c in header
        ],
        dtype=str,
        engine="openpyxl"
    )

    if len(data) != 9157:
        raise ValueError(
            f"Expected 9157 genomes; found {len(data)}"
        )

    if (
        data["Assembly_Accession"].isna().any()
        or data["Assembly_Accession"].duplicated().any()
    ):
        raise ValueError("Duplicate or missing accessions")

    for col in optional:
        if col not in data:
            data[col] = ""

    calls = data["AMR_Genes"].map(identify_genes)

    for family in FAMILIES:
        data[family + "_Variants"] = calls.map(
            lambda item, f=family:
            "; ".join(sorted(item[0][f]))
        )

        data[family + "_Positive"] = (
            data[family + "_Variants"].ne("")
        )

    data["Genes_For_Review"] = calls.map(
        lambda item: "; ".join(item[1])
    )

    positive_flags = [
        f + "_Positive" for f in FAMILIES
    ]

    data["Positive_Family_Count"] = (
        data[positive_flags].sum(axis=1)
    )

    data["Any_Target_Positive"] = (
        data["Positive_Family_Count"].gt(0)
    )

    positive = data[
        data["Any_Target_Positive"]
    ].copy()

    positive["Country"] = (
        positive["Country_Harmonized"]
        .fillna("Unclassified")
        .str.strip()
        .replace("", "Unclassified")
    )

    positive["Year"] = pd.to_numeric(
        positive["Collection_Year_Numeric"],
        errors="coerce"
    )

    positive["FASTA_Path_Exists"] = (
        positive["Genome_File"].fillna("")
        .map(
            lambda x:
            bool(str(x).strip()) and Path(str(x)).is_file()
        )
    )

    positive["FASTA_Path_Status"] = positive[
        "FASTA_Path_Exists"
    ].map({
        True: "AVAILABLE",
        False: "DOWNLOAD_OR_LOCATE"
    })

    positive = positive.sort_values(
        [
            "Positive_Family_Count",
            "Normalized_Species",
            "Country",
            "Assembly_Accession",
        ],
        ascending=[False, True, True, True]
    )

    # Convert each genome's variants into a long format.
    long_rows = []

    for row in positive.itertuples(index=False):
        for family in FAMILIES:
            value = getattr(row, family + "_Variants")

            for allele in value.split("; "):
                if not allele:
                    continue

                long_rows.append({
                    "Assembly_Accession":
                        row.Assembly_Accession,
                    "Species":
                        row.Normalized_Species,
                    "ST":
                        row.MLST_ST,
                    "Country":
                        row.Country,
                    "Year":
                        row.Year,
                    "Source":
                        row.Final_Source_Category,
                    "Family":
                        family,
                    "Allele":
                        allele,
                })

    long = pd.DataFrame(long_rows)

    allele_summary = summarize(
        long, ["Family", "Allele"]
    )

    allele_country = summarize(
        long, ["Family", "Allele", "Country"]
    )

    allele_species = summarize(
        long, [
            "Family", "Allele", "Species"
        ]
    )

    allele_st = summarize(
        long, [
            "Family", "Allele", "Species", "ST"
        ]
    )

    allele_year = summarize(
        long, [
            "Family", "Allele", "Year"
        ]
    )

    allele_source = summarize(
        long, [
            "Family", "Allele", "Source"
        ]
    )

    # Country-year denominators must come from all
    # genomes, not only target-positive genomes.
    denominator = data.copy()

    denominator["Country"] = (
        denominator["Country_Harmonized"]
        .fillna("Unclassified")
        .str.strip()
        .replace("", "Unclassified")
    )

    denominator["Year"] = pd.to_numeric(
        denominator["Collection_Year_Numeric"],
        errors="coerce"
    )

    country_year_total = summarize(
        denominator,
        ["Country", "Year"]
    ).rename(columns={
        "Genome_Count": "All_ECC_Genomes"
    })

    family_country_year = []

    for family in FAMILIES:
        flagged = denominator[
            denominator[family + "_Positive"]
        ]

        tab = summarize(
            flagged, ["Country", "Year"]
        )

        tab["Family"] = family

        family_country_year.append(tab)

    family_country_year = pd.concat(
        family_country_year,
        ignore_index=True
    )

    family_country_year = family_country_year.merge(
        country_year_total,
        on=["Country", "Year"],
        how="left",
        validate="many_to_one"
    )

    family_country_year["Detection_Percent"] = (
        100.0
        * family_country_year["Genome_Count"]
        / family_country_year["All_ECC_Genomes"]
    )

    # Exact accession files.
    for family in FAMILIES:
        subset = positive[
            positive[family + "_Positive"]
        ]

        number = save_list(
            LIST_DIR / f"{family}_ALL_ACCESSIONS.txt",
            subset["Assembly_Accession"]
        )

        print(f"{family:5s}: {number:,} genomes")

    save_list(
        LIST_DIR / "NDM_IMP_MCR_UNION_ACCESSIONS.txt",
        positive["Assembly_Accession"]
    )

    multiple = positive[
        positive["Positive_Family_Count"].ge(2)
    ]

    save_list(
        LIST_DIR / "MULTIPLE_FAMILY_ACCESSIONS.txt",
        multiple["Assembly_Accession"]
    )

    # Also create each country × family list.
    country_index = []

    for (country, family), group in long.groupby(
        ["Country", "Family"]
    ):
        accessions = group["Assembly_Accession"].unique()

        path = (
            COUNTRY_DIR
            / filename_safe(country)
            / f"{family}_ACCESSIONS.txt"
        )

        path.parent.mkdir(
            parents=True, exist_ok=True
        )

        n = save_list(path, accessions)

        country_index.append({
            "Country": country,
            "Family": family,
            "Genome_Count": n,
            "Accession_List": str(path)
        })

    country_index = pd.DataFrame(country_index)

    # Write a clear audit table.
    qc = pd.DataFrame([
        ["Total master genomes", len(data)],
        ["Any target gene", len(positive)],
        ["NDM genomes",
         int(data["NDM_Positive"].sum())],
        ["IMP genomes",
         int(data["IMP_Positive"].sum())],
        ["mcr genomes",
         int(data["MCR_Positive"].sum())],
        ["Two or more families", len(multiple)],
        ["Unique allele annotations", len(allele_summary)],
        ["FASTAs located",
         int(positive["FASTA_Path_Exists"].sum())],
        ["FASTAs not yet located",
         int((~positive["FASTA_Path_Exists"]).sum())],
        ["Unresolved target-like gene calls",
         int(data["Genes_For_Review"].ne("").sum())],
        ["Interpretation",
         "Gene carriage is not a confirmed phenotype"],
    ], columns=["Metric", "Value"])

    with pd.ExcelWriter(
        EXCEL,
        engine="openpyxl"
    ) as writer:

        qc.to_excel(
            writer, sheet_name="QC", index=False
        )

        positive.to_excel(
            writer,
            sheet_name="ALL_POSITIVE_GENOMES",
            index=False
        )

        for family in FAMILIES:
            positive[
                positive[family + "_Positive"]
            ].to_excel(
                writer,
                sheet_name=f"{family}_GENOMES",
                index=False
            )

        multiple.to_excel(
            writer,
            sheet_name="MULTIPLE_FAMILIES",
            index=False
        )

        long.to_excel(
            writer,
            sheet_name="GENOME_ALLELE_LONG",
            index=False
        )

        allele_summary.to_excel(
            writer,
            sheet_name="ALLELE_FREQUENCIES",
            index=False
        )

        allele_country.to_excel(
            writer,
            sheet_name="ALLELE_BY_COUNTRY",
            index=False
        )

        allele_species.to_excel(
            writer,
            sheet_name="ALLELE_BY_SPECIES",
            index=False
        )

        allele_st.to_excel(
            writer,
            sheet_name="ALLELE_BY_ST",
            index=False
        )

        allele_year.to_excel(
            writer,
            sheet_name="ALLELE_BY_YEAR",
            index=False
        )

        allele_source.to_excel(
            writer,
            sheet_name="ALLELE_BY_SOURCE",
            index=False
        )

        family_country_year.to_excel(
            writer,
            sheet_name="COUNTRY_YEAR_PREVALENCE",
            index=False
        )

        country_index.to_excel(
            writer,
            sheet_name="COUNTRY_LIST_INDEX",
            index=False
        )

        data.loc[
            data["Genes_For_Review"].ne(""),
            [
                "Assembly_Accession",
                "AMR_Genes",
                "Genes_For_Review"
            ]
        ].to_excel(
            writer,
            sheet_name="GENE_REVIEW",
            index=False
        )

        style_excel(writer)

    print()
    print("EXCEL:", EXCEL)
    print("FAMILY ACCESSION LISTS:", LIST_DIR)
    print("COUNTRY ACCESSION LISTS:", COUNTRY_DIR)
    print("UNIQUE POSITIVE GENOMES:", len(positive))
    print("=" * 100)


if __name__ == "__main__":
    main()
