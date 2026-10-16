
#!/usr/bin/env python3

# =============================================================
# FIGURE 6 — ECC TEMPORAL AMR GENE BURDEN, 2000–2026
#
# 6A: 7-species temporal trajectories
# 6B: 7-species income-stratified temporal trajectories
# 6C: Sampling coverage heatmap
#
# All figures: 600-DPI JPEG
# Excel: genome-level data, summaries, CI, QC and methods
#
# Burden = number of distinct AMR gene names per genome.
# No MDR classification and no invented gene-class mapping.
#
# IMPORTANT:
# Set YEAR_COLUMN explicitly if you know your preferred
# collection/isolation year column. Otherwise AUTO selects
# an identifiable collection year first, and reports any
# submission/release-year fallback transparently.
# =============================================================

from pathlib import Path
from datetime import datetime
import re
import warnings

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import LogNorm
import matplotlib.patheffects as pe

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings(
    "ignore",
    message="Glyph .* missing from font"
)

# =============================================================
# PATHS
# =============================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

INCOME_FILE = (
    ROOT / "Fig_01A_Global_Geographical_Distribution"
    / "Fig_01A_Country_Income_Database.xlsx"
)

OUT = ROOT / "Fig_06_ECC_Temporal_AMR_Burden"
OUT.mkdir(parents=True, exist_ok=True)

FIG6A = OUT / "Fig06A_Seven_Species_Temporal_AMR_Burden_600DPI.jpg"
FIG6B = OUT / "Fig06B_Income_Stratified_Temporal_AMR_Burden_600DPI.jpg"
FIG6C = OUT / "Fig06C_Temporal_Genome_Sampling_Coverage_600DPI.jpg"

EXCEL = OUT / "Fig06_Temporal_AMR_Burden_ALL_SOURCE_DATA.xlsx"

DPI = 600

# Change to actual collection/isolation year column if known.
# Example: YEAR_COLUMN = "Collection_Year"
YEAR_COLUMN = None

MIN_N_PLOT = 5
BOOTSTRAPS = 1000
RANDOM_SEED = 2026

# =============================================================
# SPECIES
# =============================================================

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_SHORT = {
    "Enterobacter hormaechei": "E. hormaechei",
    "Enterobacter cloacae": "E. cloacae",
    "Enterobacter roggenkampii": "E. roggenkampii",
    "Enterobacter asburiae": "E. asburiae",
    "Enterobacter kobei": "E. kobei",
    "Enterobacter ludwigii": "E. ludwigii",
    "Enterobacter bugandensis": "E. bugandensis",
}

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

INCOME_ORDER = ["H", "M", "L+LM"]

INCOME_COLORS = {
    "H": "#C98A00",
    "M": "#673AB7",
    "L+LM": "#C94F82",
}

INCOME_NAMES = {
    "H": "High income",
    "M": "Upper-middle income",
    "L+LM": "Lower-middle + low income",
}

PERIODS = [
    ("2000–2004", 2000, 2004),
    ("2005–2009", 2005, 2009),
    ("2010–2014", 2010, 2014),
    ("2015–2019", 2015, 2019),
    ("2020–2024", 2020, 2024),
    ("2025–2026", 2025, 2026),
]

PERIOD_ORDER = [x[0] for x in PERIODS]
PERIOD_MIDPOINTS = [
    (x[1] + x[2]) / 2 for x in PERIODS
]

rng = np.random.default_rng(RANDOM_SEED)

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 11,
    "axes.linewidth": 0.9,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

# =============================================================
# HELPERS
# =============================================================

def require_file(path):
    if not path.is_file():
        raise FileNotFoundError(str(path))


def clean_string(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def standardize_income(value):
    s = clean_string(value).upper()
    s = re.sub(r"[\s_\-]+", "", s)

    mapping = {
        "H": "H",
        "HIGH": "H",
        "HIGHINCOME": "H",

        "M": "M",
        "UM": "M",
        "UPPERMIDDLE": "M",
        "UPPERMIDDLEINCOME": "M",

        "L+LM": "L+LM",
        "LM+L": "L+LM",
        "LOWERMIDDLE+LOW": "L+LM",
        "LOW+LOWERMIDDLE": "L+LM",
        "LOWERMIDDLEANDLOW": "L+LM",
    }

    return mapping.get(s, "Unclassified")


def parse_amr_genes(value):
    if pd.isna(value):
        return set()

    s = str(value).strip()

    if s.lower() in {
        "", "nan", "none", "no hits",
        "no hit", "not detected", "-", "na", "n/a"
    }:
        return set()

    return {
        gene.strip()
        for gene in s.split(";")
        if gene.strip()
    }


def get_year(value):
    if pd.isna(value):
        return np.nan

    if isinstance(value, (pd.Timestamp, datetime)):
        return int(value.year)

    s = str(value).strip()

    if re.fullmatch(r"\d{4}(?:\.0+)?", s):
        return int(float(s))

    # Recognizes YYYY-MM-DD or YYYY/MM/DD.
    match = re.search(r"(?<!\d)(19\d{2}|20\d{2})(?!\d)", s)

    if match:
        return int(match.group(1))

    return np.nan


def identify_year_column(frame):
    if YEAR_COLUMN is not None:
        if YEAR_COLUMN not in frame.columns:
            raise ValueError(
                "Specified YEAR_COLUMN not in master: "
                + str(YEAR_COLUMN)
            )
        return YEAR_COLUMN, "USER-SPECIFIED"

    candidates = [
        ("Isolation_Year", "ISOLATION"),
        ("Isolate_Collection_Year", "COLLECTION"),
        ("Collection_Year", "COLLECTION"),
        ("Year_of_Collection", "COLLECTION"),
        ("Collection_Date", "COLLECTION"),
        ("Isolation_Date", "ISOLATION"),
        ("Isolate_Collection_Date", "COLLECTION"),
        ("Sample_Collection_Date", "COLLECTION"),
        ("Date_Collected", "COLLECTION"),
        ("Year", "UNSPECIFIED_YEAR"),
        ("Release_Year", "RELEASE"),
        ("Assembly_Release_Year", "RELEASE"),
        ("Assembly_Release_Date", "RELEASE"),
        ("Submit_Date", "SUBMISSION"),
        ("Submission_Date", "SUBMISSION"),
        ("Submission_Year", "SUBMISSION"),
        ("Release_Date", "RELEASE"),
    ]

    lower_lookup = {
        c.lower(): c for c in frame.columns
    }

    for candidate, kind in candidates:
        key = candidate.lower()

        if key not in lower_lookup:
            continue

        col = lower_lookup[key]
        parsed = frame[col].map(get_year)

        eligible = parsed.between(
            2000, 2026, inclusive="both"
        ).sum()

        if eligible > 0:
            return col, kind

    raise ValueError(
        "No usable year column was identified. "
        "Available columns:\n"
        + "\n".join(map(str, frame.columns))
        + "\n\nSet YEAR_COLUMN near the top of the script."
    )


def period_from_year(year):
    if pd.isna(year):
        return None

    year = int(year)

    for name, first, last in PERIODS:
        if first <= year <= last:
            return name

    return None


def ci_mean(values):
    a = np.asarray(values, dtype=float)
    a = a[np.isfinite(a)]

    n = len(a)

    if n == 0:
        return {
            "N": 0,
            "Mean": np.nan,
            "Median": np.nan,
            "SD": np.nan,
            "Q1": np.nan,
            "Q3": np.nan,
            "CI_Lower": np.nan,
            "CI_Upper": np.nan,
        }

    mean = float(a.mean())
    median = float(np.median(a))

    result = {
        "N": n,
        "Mean": mean,
        "Median": median,
        "SD": float(a.std(ddof=1)) if n > 1 else np.nan,
        "Q1": float(np.percentile(a, 25)),
        "Q3": float(np.percentile(a, 75)),
        "CI_Lower": np.nan,
        "CI_Upper": np.nan,
    }

    if n >= MIN_N_PLOT:
        # Bootstrap sampling of the mean.
        # Fixed seed for reproducibility.
        boot_means = np.empty(
            BOOTSTRAPS, dtype=float
        )

        for b in range(BOOTSTRAPS):
            sample = rng.choice(a, size=n, replace=True)
            boot_means[b] = sample.mean()

        result["CI_Lower"] = float(
            np.percentile(boot_means, 2.5)
        )
        result["CI_Upper"] = float(
            np.percentile(boot_means, 97.5)
        )

    return result


def summary_by(frame, group_columns):
    records = []

    for key, part in frame.groupby(
        group_columns, dropna=False, observed=True
    ):
        if not isinstance(key, tuple):
            key = (key,)

        record = dict(zip(group_columns, key))

        stats = ci_mean(
            part["AMR_Gene_Burden"].to_numpy()
        )

        record.update(stats)
        records.append(record)

    return pd.DataFrame(records)


def format_axis(ax, ymax):
    ax.set_ylim(0, ymax)

    ax.grid(
        axis="y",
        linestyle=(0, (2, 3)),
        linewidth=0.55,
        alpha=0.26
    )

    ax.set_axisbelow(True)

    ax.tick_params(
        axis="both",
        labelsize=9,
        length=3
    )

    ax.set_xticks(range(len(PERIOD_ORDER)))
    ax.set_xticklabels(
        ["00–04", "05–09", "10–14",
         "15–19", "20–24", "25–26"],
        fontsize=8.8
    )

    for spine in ["left", "bottom"]:
        ax.spines[spine].set_color("#AEB9C4")


def save_jpg(fig, path):
    fig.savefig(
        path,
        dpi=DPI,
        format="jpeg",
        facecolor="white",
        pil_kwargs={
            "quality": 94,
            "subsampling": 1
        }
    )
    plt.close(fig)

    if not path.is_file():
        raise RuntimeError(
            "JPEG not saved: " + str(path)
        )

    print(
        "[SAVED]",
        path.name,
        f"({path.stat().st_size / (1024*1024):.1f} MB)"
    )


def style_excel(writer):
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="173352"
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for i, column in enumerate(
            sheet.columns, start=1
        ):
            first = next(iter(column)).value
            width = min(
                38,
                max(14, len(str(first or "")) + 3)
            )
            sheet.column_dimensions[
                get_column_letter(i)
            ].width = width


# =============================================================
# LOAD MASTER AND COUNTRY INCOME MAPPING
# =============================================================

print("=" * 95)
print("FIGURE 6 — TEMPORAL AMR GENE BURDEN")
print("=" * 95)

require_file(MASTER)
require_file(INCOME_FILE)

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    engine="openpyxl",
    dtype={"Assembly_Accession": str}
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9157 genomes; found {len(df)}"
    )

needed = [
    "Assembly_Accession",
    "Normalized_Species",
    "Country_Harmonized",
    "AMR_Genes",
]

for col in needed:
    if col not in df.columns:
        raise ValueError(
            "Missing column: " + col
        )

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicate accession(s) in master."
    )

income = pd.read_excel(
    INCOME_FILE,
    sheet_name="COUNTRY_INCOME_COUNTS",
    engine="openpyxl"
)

for col in ["Country_Harmonized", "Income_Group"]:
    if col not in income.columns:
        raise ValueError(
            "Missing income database column: " + col
        )

df["Country_Key"] = (
    df["Country_Harmonized"]
    .map(clean_string)
)

income["Country_Key"] = (
    income["Country_Harmonized"]
    .map(clean_string)
)

income["Income_Group"] = (
    income["Income_Group"]
    .map(standardize_income)
)

if income["Country_Key"].duplicated().any():
    raise ValueError(
        "Duplicate country names in income mapping."
    )

income_map = income.set_index(
    "Country_Key"
)["Income_Group"]

df["Income_Group_Fig06"] = (
    df["Country_Key"]
    .map(income_map)
    .fillna("Unclassified")
)

df["Species_Fig06"] = (
    df["Normalized_Species"]
    .map(clean_string)
)

# =============================================================
# YEAR SELECTION
# =============================================================

year_col, year_kind = identify_year_column(df)

df["Year_Fig06"] = (
    df[year_col].map(get_year)
)

df["Period_Fig06"] = (
    df["Year_Fig06"].map(period_from_year)
)

print()
print("Selected year column:", year_col)
print("Year type:", year_kind)

if year_kind in {
    "RELEASE",
    "SUBMISSION",
    "UNSPECIFIED_YEAR"
}:
    print(
        "WARNING: This is not confirmed to be "
        "an isolate collection year."
    )
    print(
        "Temporal interpretation must reflect "
        "the actual year-field meaning."
    )

# =============================================================
# AMR GENE BURDEN
# =============================================================

df["AMR_Gene_Burden"] = (
    df["AMR_Genes"]
    .map(parse_amr_genes)
    .map(len)
)

# Flag obvious analysis failures rather than reporting
# unprocessed genomes as zero AMR genes.
if "AMRFinder_Status" in df.columns:
    status = (
        df["AMRFinder_Status"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.lower()
    )

    failed = status.str.contains(
        r"failed|failure|error|not.run|not processed",
        regex=True
    )

    df["AMR_Annotation_Failed"] = failed
else:
    df["AMR_Annotation_Failed"] = False

df["Eligible_Species"] = (
    df["Species_Fig06"].isin(SPECIES)
)

df["Eligible_Year"] = (
    df["Period_Fig06"].notna()
)

df["Eligible_AMR"] = (
    ~df["AMR_Annotation_Failed"]
)

df["Eligible_Fig06A"] = (
    df["Eligible_Species"]
    & df["Eligible_Year"]
    & df["Eligible_AMR"]
)

df["Eligible_Fig06B"] = (
    df["Eligible_Fig06A"]
    & df["Income_Group_Fig06"].isin(INCOME_ORDER)
)

analysis_a = df[
    df["Eligible_Fig06A"]
].copy()

analysis_b = df[
    df["Eligible_Fig06B"]
].copy()

if analysis_a.empty:
    raise ValueError(
        "No genomes eligible for temporal analysis. "
        "Check the year column and species names."
    )

print()
print("Total master genomes:", len(df))
print("Resolved seven-species genomes:",
      int(df["Eligible_Species"].sum()))
print("Figure 6A genomes:", len(analysis_a))
print("Figure 6B genomes:", len(analysis_b))
print("Unclassified income among 6A:",
      int((~analysis_a["Income_Group_Fig06"]
           .isin(INCOME_ORDER)).sum()))

print()
print("Year range used:",
      int(analysis_a["Year_Fig06"].min()),
      "to",
      int(analysis_a["Year_Fig06"].max()))

# =============================================================
# GROUP STATISTICS
# =============================================================

s6a = summary_by(
    analysis_a,
    ["Species_Fig06", "Period_Fig06"]
)

s6b = summary_by(
    analysis_b,
    [
        "Species_Fig06",
        "Income_Group_Fig06",
        "Period_Fig06"
    ]
)

s6a["Period_Fig06"] = pd.Categorical(
    s6a["Period_Fig06"],
    categories=PERIOD_ORDER,
    ordered=True
)

s6b["Period_Fig06"] = pd.Categorical(
    s6b["Period_Fig06"],
    categories=PERIOD_ORDER,
    ordered=True
)

s6a = s6a.sort_values(
    ["Species_Fig06", "Period_Fig06"]
)

s6b = s6b.sort_values(
    [
        "Species_Fig06",
        "Income_Group_Fig06",
        "Period_Fig06"
    ]
)

max_observed_burden = max(
    float(s6a["Mean"].max()),
    float(s6b["Mean"].max())
    if not s6b.empty else 0
)

max_ci = max(
    float(s6a["CI_Upper"].max())
    if s6a["CI_Upper"].notna().any() else 0,
    float(s6b["CI_Upper"].max())
    if (
        not s6b.empty
        and s6b["CI_Upper"].notna().any()
    ) else 0
)

y_max = max(
    5,
    math_ceil := int(
        np.ceil(max(
            max_observed_burden,
            max_ci
        ) * 1.18)
    )
)

# =============================================================
# FIGURE 6A — SEVEN SPECIES
# =============================================================

print("\n[FIGURE 6A] Building species trajectories...")

fig, axes = plt.subplots(
    4, 2,
    figsize=(16.0, 15.6),
    sharex=True,
    sharey=True
)

axes = axes.flatten()

for idx, sp in enumerate(SPECIES):
    ax = axes[idx]
    color = SPECIES_COLORS[sp]

    sub = s6a[
        s6a["Species_Fig06"].eq(sp)
    ].copy()

    x = [
        PERIOD_ORDER.index(p)
        for p in sub["Period_Fig06"]
    ]

    values = sub["Mean"].to_numpy(float)
    ns = sub["N"].to_numpy(int)

    lower = sub["CI_Lower"].to_numpy(float)
    upper = sub["CI_Upper"].to_numpy(float)

    x = np.asarray(x)
    valid = (ns >= MIN_N_PLOT)

    # Lines break at periods with fewer than MIN_N_PLOT
    # genomes instead of implying reliable continuity.
    yline = np.where(valid, values, np.nan)

    ax.plot(
        x, yline,
        color=color,
        linewidth=2.35,
        marker="o",
        markersize=6.3,
        markeredgecolor="white",
        markeredgewidth=0.9,
        zorder=5
    )

    valid_ci = (
        valid
        & np.isfinite(lower)
        & np.isfinite(upper)
    )

    ax.fill_between(
        x,
        np.where(valid_ci, lower, np.nan),
        np.where(valid_ci, upper, np.nan),
        color=color,
        alpha=0.16,
        linewidth=0
    )

    # Mark sparse periods without joining their points.
    sparse = (ns > 0) & (~valid)

    if sparse.any():
        ax.scatter(
            x[sparse],
            values[sparse],
            marker="x",
            s=43,
            color=color,
            linewidth=1.6,
            zorder=7
        )

    for xx, yy, nn in zip(
        x, values, ns
    ):
        ax.annotate(
            f"n={nn}",
            (xx, yy),
            xytext=(0, 9),
            textcoords="offset points",
            ha="center",
            fontsize=8,
            color="#566B7F",
            zorder=8
        )

    ax.set_title(
        SPECIES_SHORT[sp],
        loc="left",
        fontsize=12.5,
        fontweight="bold",
        fontstyle="italic",
        color=color,
        pad=13
    )

    format_axis(ax, y_max)

    if idx % 2 == 0:
        ax.set_ylabel(
            "AMR genes per genome",
            fontsize=10
        )

# Final empty panel: definitions and legend
ax = axes[7]
ax.axis("off")

ax.text(
    0.06, 0.82,
    "Figure interpretation",
    fontsize=13,
    fontweight="bold",
    color="#173352",
    transform=ax.transAxes
)

ax.text(
    0.06, 0.66,
    "Point: mean unique AMR genes/genome\n"
    "Shading: bootstrap 95% CI\n"
    "n: number of genomes in the interval\n"
    "×: fewer than five genomes\n"
    "Last interval: 2025–2026 only",
    fontsize=10.5,
    color="#4F6377",
    va="top",
    linespacing=1.8,
    transform=ax.transAxes
)

fig.suptitle(
    "Temporal dynamics of AMR gene burden "
    "across seven ECC species",
    fontsize=20,
    fontweight="bold",
    color="#173352",
    y=0.986
)

fig.text(
    0.5, 0.958,
    "Five-year intervals, 2000–2024; "
    "final partial interval, 2025–2026",
    ha="center",
    color="#586D80",
    fontsize=11
)

fig.text(
    0.5, 0.020,
    f"Time variable: {year_col} ({year_kind}). "
    "Distinct annotated AMR genes; "
    "no acquired-only filtering.",
    ha="center",
    fontsize=8.5,
    color="#5C6E7D"
)

fig.subplots_adjust(
    left=0.083, right=0.985,
    top=0.915, bottom=0.070,
    hspace=0.50, wspace=0.27
)

save_jpg(fig, FIG6A)

# =============================================================
# FIGURE 6B — SPECIES × INCOME × TIME
# =============================================================

print("\n[FIGURE 6B] Building income trajectories...")

fig, axes = plt.subplots(
    4, 2,
    figsize=(17.0, 16.2),
    sharex=True,
    sharey=True
)

axes = axes.flatten()

for idx, sp in enumerate(SPECIES):
    ax = axes[idx]

    for inc in INCOME_ORDER:

        sub = s6b[
            s6b["Species_Fig06"].eq(sp)
            & s6b["Income_Group_Fig06"].eq(inc)
        ].copy()

        if sub.empty:
            continue

        x = np.array([
            PERIOD_ORDER.index(p)
            for p in sub["Period_Fig06"]
        ])

        y = sub["Mean"].to_numpy(float)
        n = sub["N"].to_numpy(int)

        lo = sub["CI_Lower"].to_numpy(float)
        hi = sub["CI_Upper"].to_numpy(float)

        valid = n >= MIN_N_PLOT
        sparse = (n > 0) & (~valid)

        color = INCOME_COLORS[inc]

        ax.plot(
            x,
            np.where(valid, y, np.nan),
            color=color,
            linewidth=2.2,
            marker="o",
            markersize=5.0,
            markeredgecolor="white",
            markeredgewidth=0.7,
            zorder=5
        )

        ci_valid = (
            valid
            & np.isfinite(lo)
            & np.isfinite(hi)
        )

        ax.fill_between(
            x,
            np.where(ci_valid, lo, np.nan),
            np.where(ci_valid, hi, np.nan),
            color=color,
            alpha=0.12,
            linewidth=0
        )

        if sparse.any():
            ax.scatter(
                x[sparse],
                y[sparse],
                marker="x",
                s=42,
                linewidth=1.6,
                color=color,
                zorder=7
            )

    n_sp = int(
        analysis_b["Species_Fig06"].eq(sp).sum()
    )

    ax.set_title(
        f"{SPECIES_SHORT[sp]}  (n={n_sp:,})",
        fontsize=12.0,
        fontweight="bold",
        fontstyle="italic",
        loc="left",
        color=SPECIES_COLORS[sp],
        pad=12
    )

    format_axis(ax, y_max)

    if idx % 2 == 0:
        ax.set_ylabel(
            "AMR genes per genome",
            fontsize=10
        )

axes[7].axis("off")

legend_handles = [
    Line2D(
        [0], [0],
        color=INCOME_COLORS[inc],
        marker="o",
        markersize=6.5,
        linewidth=2.5,
        label=INCOME_NAMES[inc]
    )
    for inc in INCOME_ORDER
]

axes[7].legend(
    handles=legend_handles,
    loc="upper left",
    bbox_to_anchor=(0.04, 0.92),
    frameon=False,
    fontsize=11,
    title="Country-income group",
    title_fontsize=12
)

axes[7].text(
    0.05, 0.32,
    "Points: mean AMR genes/genome\n"
    "Shading: bootstrap 95% CI\n"
    "×: n < 5 (sparse group)\n"
    "Unclassified-income genomes excluded",
    transform=axes[7].transAxes,
    fontsize=10,
    color="#586D80",
    linespacing=1.7
)

fig.suptitle(
    "Income-stratified temporal changes "
    "in ECC AMR gene burden",
    fontsize=20,
    fontweight="bold",
    color="#173352",
    y=0.985
)

fig.text(
    .5, .957,
    "Seven species analyzed separately "
    "using identical burden and temporal scales",
    ha="center",
    fontsize=11,
    color="#536B7E"
)

fig.text(
    .5, .020,
    "H = high income; M = upper-middle income; "
    "L+LM = lower-middle and low income. "
    f"Year field: {year_col}.",
    ha="center",
    fontsize=8.5,
    color="#607385"
)

fig.subplots_adjust(
    left=.080, right=.984,
    top=.913, bottom=.070,
    hspace=.49, wspace=.26
)

save_jpg(fig, FIG6B)

# =============================================================
# FIGURE 6C — SAMPLE COUNT COVERAGE
# =============================================================

print("\n[FIGURE 6C] Building coverage heatmap...")

# 7 species × 3 income groups, 6 time intervals.
row_order = [
    (sp, inc)
    for sp in SPECIES
    for inc in INCOME_ORDER
]

coverage = np.zeros(
    (len(row_order), len(PERIOD_ORDER)),
    dtype=int
)

count_lookup = (
    analysis_b.groupby(
        [
            "Species_Fig06",
            "Income_Group_Fig06",
            "Period_Fig06"
        ]
    )
    .size()
    .to_dict()
)

for i, (sp, inc) in enumerate(row_order):
    for j, period in enumerate(PERIOD_ORDER):
        coverage[i, j] = int(
            count_lookup.get(
                (sp, inc, period), 0
            )
        )

masked = np.ma.masked_where(
    coverage == 0, coverage
)

fig, ax = plt.subplots(
    figsize=(13.2, 11.8)
)

cmap = plt.get_cmap("YlGnBu").copy()
cmap.set_bad("#F0F2F5")

positive = coverage[coverage > 0]
vmax = max(
    2,
    int(positive.max())
    if len(positive) else 2
)

im = ax.imshow(
    masked,
    interpolation="nearest",
    aspect="auto",
    cmap=cmap,
    norm=LogNorm(vmin=1, vmax=vmax)
)

ax.set_xticks(
    np.arange(len(PERIOD_ORDER))
)

ax.set_xticklabels(
    PERIOD_ORDER,
    fontsize=10,
    rotation=25,
    ha="right"
)

ax.set_yticks(
    np.arange(len(row_order))
)

ax.set_yticklabels(
    [
        f"{SPECIES_SHORT[sp]}  |  {inc}"
        for sp, inc in row_order
    ],
    fontsize=10
)

for i in range(coverage.shape[0]):
    for j in range(coverage.shape[1]):

        value = coverage[i, j]

        if value == 0:
            text_color = "#697782"
        else:
            norm_val = im.norm(value)
            text_color = (
                "white" if norm_val > 0.53
                else "#142A3E"
            )

        ax.text(
            j, i,
            f"{value:,}",
            ha="center",
            va="center",
            fontsize=9.5,
            fontweight="bold",
            color=text_color
        )

# Separate the seven species visually.
for i in range(3, len(row_order), 3):
    ax.axhline(
        i - 0.5,
        color="white",
        linewidth=3
    )

ax.set_xticks(
    np.arange(-.5, len(PERIOD_ORDER), 1),
    minor=True
)

ax.set_yticks(
    np.arange(-.5, len(row_order), 1),
    minor=True
)

ax.grid(
    which="minor",
    color="white",
    linewidth=0.9
)

ax.tick_params(
    which="minor",
    bottom=False,
    left=False
)

cbar = fig.colorbar(
    im,
    ax=ax,
    shrink=0.75,
    pad=0.03
)

cbar.set_label(
    "Genomes per species–income–period group (log scale)",
    fontsize=10
)

fig.suptitle(
    "Temporal sampling coverage across "
    "ECC species and income groups",
    fontsize=18,
    fontweight="bold",
    color="#173352",
    y=.982
)

fig.text(
    .5, .950,
    "Cells show genome counts, "
    "not AMR prevalence",
    ha="center",
    fontsize=10.5,
    color="#586D80"
)

fig.text(
    .5, .017,
    "H: high; M: upper-middle; "
    "L+LM: lower-middle and low. "
    "Gray cells: zero genomes.",
    ha="center",
    fontsize=9,
    color="#586D80"
)

fig.subplots_adjust(
    left=.235, right=.935,
    top=.915, bottom=.10
)

save_jpg(fig, FIG6C)

# =============================================================
# COVERAGE AND AUDIT TABLES
# =============================================================

coverage_rows = []

for i, (sp, inc) in enumerate(row_order):
    for j, period in enumerate(PERIOD_ORDER):
        coverage_rows.append({
            "Species": sp,
            "Income_Group": inc,
            "Period": period,
            "Genome_Count": int(coverage[i, j]),
            "Sufficient_For_Trajectory": (
                coverage[i, j] >= MIN_N_PLOT
            )
        })

coverage_df = pd.DataFrame(coverage_rows)

qc = pd.DataFrame([
    ["Master genome count", len(df)],
    [
        "Seven resolved species",
        int(df["Eligible_Species"].sum())
    ],
    [
        "Valid year 2000–2026",
        int(df["Eligible_Year"].sum())
    ],
    [
        "Failed AMR annotation",
        int(df["AMR_Annotation_Failed"].sum())
    ],
    [
        "Eligible for Figure 6A",
        len(analysis_a)
    ],
    [
        "Eligible for Figure 6B",
        len(analysis_b)
    ],
    [
        "Excluded from 6B for income not classified",
        len(analysis_a) - len(analysis_b)
    ],
    ["Year column used", year_col],
    ["Year field type", year_kind],
    [
        "Last interval",
        "2025–2026 (2 years)"
    ],
    [
        "Minimum n for line point",
        MIN_N_PLOT
    ],
    [
        "Bootstrap iterations",
        BOOTSTRAPS
    ],
    [
        "Distinct gene counting",
        "Unique semicolon-delimited AMR gene names"
    ],
    [
        "AMR type",
        "All annotated AMR genes; intrinsic not excluded"
    ],
    [
        "Blank AMR_Genes cells",
        "Counted as zero unless AMRFinder_Status indicates failure; review QC"
    ],
    [
        "Interpretation",
        "Temporal patterns in sampled genomes, not population-level AMR incidence"
    ],
], columns=["QC_Parameter", "Value"])

methods = pd.DataFrame([
    [
        "Primary data",
        "Final 9,157-genome ECC master"
    ],
    [
        "Species",
        "Seven resolved ECC species"
    ],
    [
        "AMR gene burden",
        "Number of distinct AMR gene annotations per genome"
    ],
    [
        "Periods",
        "2000–2004; 2005–2009; 2010–2014; "
        "2015–2019; 2020–2024; 2025–2026"
    ],
    [
        "Income classification",
        "Frozen country-income mapping from Figure 1A"
    ],
    [
        "Figure 6A",
        "Mean annotated AMR genes/genome by species and period"
    ],
    [
        "Figure 6B",
        "Mean annotated AMR genes/genome by species, income and period"
    ],
    [
        "Confidence intervals",
        f"Percentile bootstrap 95% CI, {BOOTSTRAPS} resamples"
    ],
    [
        "Sparse groups",
        f"Groups with n < {MIN_N_PLOT} shown as unconnected x markers"
    ],
    [
        "Figure 6C",
        "Genome counts supporting each species-income-time group"
    ],
    [
        "Temporal interpretation",
        f"Source field: {year_col}; type: {year_kind}"
    ],
    [
        "Caveat",
        "NCBI/public genomes are not population-representative; "
        "sampling and annotation differences can affect comparisons"
    ],
], columns=["Parameter", "Definition"])

# =============================================================
# EXPORT EXCEL
# =============================================================

print("\n[EXCEL] Exporting data...")

genome_columns = [
    "Assembly_Accession",
    "Species_Fig06",
    "Country_Harmonized",
    "Income_Group_Fig06",
    "Year_Fig06",
    "Period_Fig06",
    "AMR_Gene_Burden",
    "AMR_Genes",
    "AMR_Annotation_Failed",
    "Eligible_Fig06A",
    "Eligible_Fig06B",
]

genome_columns = [
    c for c in genome_columns
    if c in df.columns
]

with pd.ExcelWriter(
    EXCEL,
    engine="openpyxl"
) as writer:

    df[genome_columns].to_excel(
        writer,
        sheet_name="ALL_9157_GENOMES",
        index=False
    )

    s6a.to_excel(
        writer,
        sheet_name="FIG6A_SPECIES_TEMPORAL",
        index=False
    )

    s6b.to_excel(
        writer,
        sheet_name="FIG6B_SPECIES_INCOME_TIME",
        index=False
    )

    coverage_df.to_excel(
        writer,
        sheet_name="FIG6C_SAMPLE_COVERAGE",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC",
        index=False
    )

    methods.to_excel(
        writer,
        sheet_name="METHODS",
        index=False
    )

    style_excel(writer)

print()
print("=" * 95)
print("FIGURE 6 COMPLETED")
print("=" * 95)
print("Master genomes:", len(df))
print("Figure 6A genomes:", len(analysis_a))
print("Figure 6B genomes:", len(analysis_b))
print("Year field:", year_col)
print("Year type:", year_kind)
print()
print("6A:", FIG6A)
print("6B:", FIG6B)
print("6C:", FIG6C)
print("Excel:", EXCEL)
print("=" * 95)
