
#!/usr/bin/env python3

# ============================================================
# FIGURE 6 — TEMPORAL AMR BURDEN IN SEVEN ECC SPECIES
# Figure 3A-inspired layout, Nature Communications reference
#
# 7 species-specific panels, full five-year labels
# Mean ± SEM, Kruskal-Wallis and pairwise significance
# 600-DPI JPEG + statistical Excel workbook
#
# Data: 9,157-genome final ECC master
# ============================================================

from pathlib import Path
import re
import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from matplotlib.ticker import MaxNLocator
from scipy.stats import kruskal, mannwhitneyu
from statsmodels.stats.multitest import multipletests
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

OUT = (
    ROOT / "Fig_06_ECC_Temporal_AMR_Burden"
    / "FIG3A_STYLE_FINAL"
)

OUT.mkdir(parents=True, exist_ok=True)

JPG = OUT / "Fig06_Seven_ECC_Species_ARG_Burden_Fig3A_Style_600DPI.jpg"
XLS = OUT / "Fig06_Seven_ECC_Species_ARG_Burden_Statistics.xlsx"

DPI = 600

# Set explicitly if you know the collection-year field.
YEAR_COLUMN = None

YEAR_CANDIDATES = [
    "Isolation_Year",
    "Collection_Year",
    "Isolate_Collection_Year",
    "Year_of_Collection",
    "Collection_Date",
    "Isolation_Date",
    "Isolate_Collection_Date",
    "Sample_Collection_Date",
    "Year",
    "Release_Year",
    "Assembly_Release_Year",
    "Assembly_Release_Date",
    "Submission_Year",
    "Submission_Date",
    "Release_Date",
]

MIN_N = 5
MAX_BRACKETS = 3

SPECIES = [
    ("Enterobacter hormaechei", "#0072B2"),
    ("Enterobacter cloacae", "#D55E00"),
    ("Enterobacter roggenkampii", "#009E73"),
    ("Enterobacter asburiae", "#CC79A7"),
    ("Enterobacter kobei", "#7851A9"),
    ("Enterobacter ludwigii", "#A6761D"),
    ("Enterobacter bugandensis", "#56B4E9"),
]

PERIODS = [
    (2000, 2004),
    (2005, 2009),
    (2010, 2014),
    (2015, 2019),
    (2020, 2024),
    (2025, 2026),
]

LABELS = [f"{a}–{b}" for a, b in PERIODS]


# ============================================================
# FUNCTIONS
# ============================================================

def get_year(value):
    if pd.isna(value):
        return np.nan

    if isinstance(value, pd.Timestamp):
        return value.year

    text = str(value).strip()

    if re.fullmatch(r"\d{4}(?:\.0+)?", text):
        return int(float(text))

    match = re.search(
        r"(?<!\d)(19\d{2}|20\d{2})(?!\d)",
        text
    )

    return int(match.group(1)) if match else np.nan


def get_genes(value):
    if pd.isna(value):
        return set()

    text = str(value).strip()

    if text.lower() in {
        "", "nan", "na", "n/a", "none",
        "not detected", "no hits", "-"
    }:
        return set()

    return {
        gene.strip()
        for gene in text.split(";")
        if gene.strip()
    }


def get_period(year):
    if pd.isna(year):
        return None

    for i, (start, end) in enumerate(PERIODS):
        if start <= int(year) <= end:
            return i

    return None


def significance(p):
    if p < 0.0001:
        return "****"
    if p < 0.001:
        return "***"
    if p < 0.01:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


# ============================================================
# LOAD MASTER
# ============================================================

print("=" * 100)
print("FIGURE 6 — ECC SPECIES TEMPORAL AMR BURDEN")
print("=" * 100)

if not MASTER.is_file():
    raise FileNotFoundError(MASTER)

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    engine="openpyxl",
    dtype={"Assembly_Accession": str}
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9157 genomes, found {len(df)}"
    )

required = [
    "Assembly_Accession",
    "Normalized_Species",
    "AMR_Genes",
]

missing = [c for c in required if c not in df.columns]

if missing:
    raise ValueError(f"Missing master columns: {missing}")

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicate assembly accessions"
    )

# ============================================================
# IDENTIFY YEAR FIELD
# ============================================================

if YEAR_COLUMN is not None:
    if YEAR_COLUMN not in df.columns:
        raise ValueError(
            f"YEAR_COLUMN not found: {YEAR_COLUMN}"
        )
    year_col = YEAR_COLUMN
else:
    year_col = next(
        (
            c for c in YEAR_CANDIDATES
            if c in df.columns
            and df[c].map(get_year).between(2000, 2026).any()
        ),
        None
    )

if year_col is None:
    raise ValueError(
        "No usable year column found. Available columns:\n"
        + "\n".join(df.columns)
    )

year_lower = year_col.lower()

if "collection" in year_lower or "isolation" in year_lower:
    year_kind = "Collection/isolation"
elif "release" in year_lower or "submi" in year_lower:
    year_kind = "Submission/release"
else:
    year_kind = "Unspecified"

print("YEAR COLUMN:", year_col)
print("YEAR TYPE:", year_kind)

if year_kind != "Collection/isolation":
    print(
        "WARNING: Year field is not confirmed as collection year."
    )

# ============================================================
# GENOME-LEVEL AMR BURDEN
# ============================================================

df["Year_Used"] = df[year_col].map(get_year)
df["Period_Index"] = df["Year_Used"].map(get_period)

df["Species_Used"] = (
    df["Normalized_Species"].astype(str).str.strip()
)

df["AMR_Genes_Per_Genome"] = (
    df["AMR_Genes"].map(
        lambda value: len(get_genes(value))
    )
)

if "AMRFinder_Status" in df.columns:
    status = (
        df["AMRFinder_Status"]
        .fillna("")
        .astype(str)
        .str.lower()
    )

    df["Annotation_Failed"] = status.str.contains(
        r"fail|error|not.run|not processed",
        regex=True
    )
else:
    df["Annotation_Failed"] = False

valid_species = {name for name, _ in SPECIES}

df["Figure_Eligible"] = (
    df["Species_Used"].isin(valid_species)
    & df["Period_Index"].notna()
    & ~df["Annotation_Failed"]
)

work = df[df["Figure_Eligible"]].copy()

if work.empty:
    raise ValueError("No eligible genomes found")

work["Period_Index"] = (
    work["Period_Index"].astype(int)
)

print("MASTER GENOMES:", len(df))
print("ELIGIBLE GENOMES:", len(work))


# ============================================================
# SUMMARY STATISTICS
# ============================================================

summary_rows = []
omnibus_rows = []
pair_data = {}

for species, color in SPECIES:

    sub = work[work["Species_Used"].eq(species)]
    groups = {}

    for i, label in enumerate(LABELS):

        values = sub.loc[
            sub["Period_Index"].eq(i),
            "AMR_Genes_Per_Genome"
        ].to_numpy(float)

        groups[i] = values

        n = len(values)

        summary_rows.append({
            "Species": species,
            "Period": label,
            "Period_Index": i,
            "N": n,
            "Mean_ARGs_Per_Genome": (
                float(np.mean(values)) if n else np.nan
            ),
            "SEM": (
                float(np.std(values, ddof=1) / np.sqrt(n))
                if n > 1 else np.nan
            ),
            "Median": (
                float(np.median(values)) if n else np.nan
            ),
            "Q1": (
                float(np.quantile(values, 0.25))
                if n else np.nan
            ),
            "Q3": (
                float(np.quantile(values, 0.75))
                if n else np.nan
            ),
            "Sufficient_N": n >= MIN_N,
        })

    adequate = [
        (i, values)
        for i, values in groups.items()
        if len(values) >= MIN_N
    ]

    if len(adequate) >= 2:
        try:
            h, p = kruskal(
                *(values for _, values in adequate)
            )
        except ValueError:
            h, p = 0.0, 1.0
    else:
        h, p = np.nan, np.nan

    omnibus_rows.append({
        "Species": species,
        "Eligible_Genomes": len(sub),
        "Testable_Periods": len(adequate),
        "Kruskal_Wallis_H": h,
        "P_raw": p,
    })

    # Predefined comparisons:
    # adjacent periods and first-versus-last.
    comparisons = [
        (0, 1), (1, 2), (2, 3),
        (3, 4), (4, 5), (0, 5)
    ]

    results = []

    for i, j in comparisons:
        a = groups[i]
        b = groups[j]

        if len(a) < MIN_N or len(b) < MIN_N:
            continue

        p_pair = mannwhitneyu(
            a, b,
            alternative="two-sided"
        ).pvalue

        median_difference = (
            float(np.median(b) - np.median(a))
        )

        results.append((
            i, j, p_pair, median_difference
        ))

    pair_data[species] = results


# ============================================================
# CORRECT STATISTICAL TESTS
# ============================================================

summary = pd.DataFrame(summary_rows)
omnibus = pd.DataFrame(omnibus_rows)

omnibus["Q_BH_7Species"] = np.nan

valid = omnibus["P_raw"].notna()

if valid.any():
    omnibus.loc[
        valid,
        "Q_BH_7Species"
    ] = multipletests(
        omnibus.loc[valid, "P_raw"],
        method="fdr_bh"
    )[1]

omnibus_q = dict(
    zip(
        omnibus["Species"],
        omnibus["Q_BH_7Species"]
    )
)

pair_rows = []

for species, _ in SPECIES:

    tests = pair_data[species]

    adjusted = (
        multipletests(
            [r[2] for r in tests],
            method="holm"
        )[1]
        if tests else []
    )

    for (i, j, p, diff), q in zip(tests, adjusted):

        overall_q = omnibus_q.get(species, np.nan)

        significant = bool(
            q < 0.05
            and np.isfinite(overall_q)
            and overall_q < 0.05
        )

        pair_rows.append({
            "Species": species,
            "Period_1": LABELS[i],
            "Period_2": LABELS[j],
            "Index_1": i,
            "Index_2": j,
            "Median_Difference_Later_Minus_Earlier": diff,
            "Mann_Whitney_P_raw": p,
            "Holm_P_adjusted": q,
            "Omnibus_BH_Q": overall_q,
            "Significant_After_Correction": significant,
            "Stars": significance(q) if significant else "ns"
        })

pairwise = pd.DataFrame(pair_rows)

# ============================================================
# FIGURE — SEVEN SPECIES, FULL YEAR LABELS
# ============================================================

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "figure.facecolor": "white",
    "axes.spines.top": False,
    "axes.spines.right": False,
})

fig, axes = plt.subplots(
    2, 4,
    figsize=(22.5, 13.0),
    facecolor="white"
)

axes = axes.ravel()

for idx, (species, color) in enumerate(SPECIES):

    ax = axes[idx]

    sub = summary[
        summary["Species"].eq(species)
    ].sort_values("Period_Index")

    x = sub["Period_Index"].to_numpy(int)
    y = sub["Mean_ARGs_Per_Genome"].to_numpy(float)
    sem = sub["SEM"].to_numpy(float)
    n = sub["N"].to_numpy(int)

    good = (n >= MIN_N) & np.isfinite(y)
    sparse = (n > 0) & ~good

    ax.errorbar(
        x,
        np.where(good, y, np.nan),
        yerr=np.where(
            good,
            np.nan_to_num(sem, nan=0),
            np.nan
        ),
        color=color,
        marker="o",
        markersize=6.5,
        linewidth=2.2,
        capsize=3.5,
        elinewidth=1.4,
        markeredgecolor="white",
        markeredgewidth=0.8,
        zorder=4
    )

    if sparse.any():
        ax.scatter(
            x[sparse],
            y[sparse],
            marker="x",
            s=50,
            color=color,
            zorder=5
        )

    # Sample sizes placed below corresponding points.
    for xx, yy, nn in zip(x, y, n):
        if nn > 0 and np.isfinite(yy):
            ax.annotate(
                f"n={nn}",
                (xx, yy),
                xytext=(0, -15),
                textcoords="offset points",
                ha="center",
                va="top",
                fontsize=7.4,
                color="#42576B"
            )

    if pairwise.empty:
        eligible_pairs = pairwise
    else:
        eligible_pairs = pairwise[
            pairwise["Species"].eq(species)
            & pairwise["Significant_After_Correction"]
        ].sort_values(
            ["Holm_P_adjusted", "Index_2"]
        ).head(MAX_BRACKETS)

    finite = y[np.isfinite(y)]

    upper = max(
        2.0,
        float(np.max(finite)) if len(finite) else 2.0
    )

    if len(finite):
        sem_upper = y + np.nan_to_num(sem)
        upper = max(
            upper,
            float(np.nanmax(sem_upper))
        )

    step = max(0.7, upper * 0.14)
    bracket_base = upper + step * 0.60

    # Draw corrected significance brackets.
    for level, row in enumerate(
        eligible_pairs.itertuples(index=False)
    ):

        yy = bracket_base + level * step

        ax.plot(
            [
                row.Index_1,
                row.Index_1,
                row.Index_2,
                row.Index_2
            ],
            [
                yy - step * 0.18,
                yy,
                yy,
                yy - step * 0.18
            ],
            color="#27394C",
            linewidth=1.05,
            clip_on=False
        )

        ax.text(
            (row.Index_1 + row.Index_2) / 2,
            yy + step * 0.08,
            row.Stars,
            ha="center",
            va="bottom",
            fontsize=11,
            fontweight="bold",
            color="#243447"
        )

    ax.set_ylim(
        0,
        upper + step * (len(eligible_pairs) + 1.25)
    )

    ax.set_xlim(-0.45, 5.45)

    # Full years are printed on EVERY species graph.
    ax.set_xticks(range(6))

    ax.set_xticklabels(
        LABELS,
        rotation=37,
        ha="right",
        fontsize=8.7
    )

    ax.yaxis.set_major_locator(
        MaxNLocator(nbins=5)
    )

    ax.set_xlabel(
        "Five-year interval",
        fontsize=10,
        labelpad=7
    )

    ax.set_ylabel(
        "AMR genes per genome (mean ± SEM)",
        fontsize=9.5
    )

    ax.set_title(
        species.replace("Enterobacter", "E."),
        loc="left",
        fontstyle="italic",
        fontweight="bold",
        fontsize=13,
        color=color,
        pad=12
    )

    ax.grid(
        axis="y",
        linestyle=":",
        alpha=0.25
    )

    q = omnibus_q.get(species, np.nan)

    if np.isfinite(q):
        statistical_label = (
            f"Kruskal–Wallis q = {q:.3g}"
        )
    else:
        statistical_label = "Insufficient data"

    ax.text(
        0.98, 0.98,
        statistical_label,
        transform=ax.transAxes,
        ha="right",
        va="top",
        fontsize=8.5,
        color="#324458",
        bbox={
            "boxstyle": "round,pad=0.25",
            "facecolor": "white",
            "edgecolor": "none",
            "alpha": 0.92
        }
    )

# Eighth panel = explanation.
ax = axes[7]
ax.axis("off")

ax.text(
    0.05, 0.94,
    "Statistical interpretation",
    fontsize=15,
    fontweight="bold",
    color="#173352",
    transform=ax.transAxes
)

ax.text(
    0.05, 0.79,
    "Points: mean AMR genes/genome\n"
    "Error bars: standard error of mean\n"
    "n: genomes within each interval\n"
    "Kruskal–Wallis: omnibus test\n"
    "Omnibus q: BH correction across species\n"
    "Pairwise Mann–Whitney: Holm correction\n"
    "×: sparse interval (n < 5)\n"
    "2025–2026: partial two-year interval",
    transform=ax.transAxes,
    fontsize=10.3,
    color="#354A60",
    va="top",
    linespacing=1.65
)

ax.text(
    0.05, 0.15,
    "* adjusted p < 0.05\n"
    "** adjusted p < 0.01\n"
    "*** adjusted p < 0.001\n"
    "**** adjusted p < 0.0001",
    transform=ax.transAxes,
    fontsize=10,
    color="#354A60"
)

fig.suptitle(
    "Temporal changes in AMR gene burden across seven ECC species",
    fontsize=21,
    fontweight="bold",
    color="#163451",
    y=0.99
)

fig.text(
    0.5, 0.961,
    "Five-year intervals (2000–2024) and "
    "final partial interval (2025–2026)",
    ha="center",
    fontsize=11.5,
    color="#506278"
)

fig.text(
    0.5, 0.014,
    f"Year field: {year_col} ({year_kind}). "
    "Burden includes all annotated AMR gene names; "
    "intrinsic determinants are not excluded.",
    ha="center",
    fontsize=8.7,
    color="#607184"
)

fig.subplots_adjust(
    left=0.07,
    right=0.985,
    top=0.925,
    bottom=0.145,
    wspace=0.37,
    hspace=0.67
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 93,
        "subsampling": 1
    }
)

plt.close(fig)

print("[SAVED JPEG]", JPG)

# ============================================================
# EXCEL STATISTICS
# ============================================================

qc = pd.DataFrame([
    ("Master genomes", len(df)),
    ("Seven-species genomes",
     int(df["Species_Used"].isin(valid_species).sum())),
    ("Eligible for temporal figure", len(work)),
    ("Year field", year_col),
    ("Year field meaning", year_kind),
    ("Annotation failures",
     int(df["Annotation_Failed"].sum())),
    ("Minimum n for significance", MIN_N),
    ("Omnibus method",
     "Kruskal-Wallis with BH FDR across seven species"),
    ("Pairwise method",
     "Mann-Whitney U with Holm correction"),
    ("Gene burden",
     "Distinct semicolon-delimited AMR gene names"),
    ("Data limitation",
     "Includes intrinsic AMR-associated annotations"),
], columns=["Parameter", "Value"])

export = df[[
    "Assembly_Accession",
    "Species_Used",
    "Year_Used",
    "Period_Index",
    "AMR_Genes_Per_Genome",
    "Annotation_Failed",
    "Figure_Eligible",
    "AMR_Genes"
]].copy()

export["Period"] = export["Period_Index"].map(
    lambda i: LABELS[int(i)] if pd.notna(i) else ""
)

with pd.ExcelWriter(
    XLS,
    engine="openpyxl"
) as writer:

    export.to_excel(
        writer,
        sheet_name="ALL_9157_GENOMES",
        index=False
    )

    summary.to_excel(
        writer,
        sheet_name="SPECIES_PERIOD_MEAN_SEM",
        index=False
    )

    omnibus.to_excel(
        writer,
        sheet_name="KRUSKAL_WALLIS_FDR",
        index=False
    )

    pairwise.to_excel(
        writer,
        sheet_name="PAIRWISE_HOLM_TESTS",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = PatternFill(
                "solid", fgColor="173352"
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )

        for j, cell in enumerate(sheet[1], 1):
            sheet.column_dimensions[
                get_column_letter(j)
            ].width = min(
                42,
                max(14, len(str(cell.value)) + 4)
            )

print()
print("=" * 100)
print("FIGURE 6 FINISHED")
print("=" * 100)
print("Master genomes:", len(df))
print("Eligible genomes:", len(work))
print("Year field:", year_col)
print("Year type:", year_kind)
print("JPEG:", JPG)
print("Excel:", XLS)
print()
print("Kruskal-Wallis results:")
print(omnibus.to_string(index=False))
print("=" * 100)
