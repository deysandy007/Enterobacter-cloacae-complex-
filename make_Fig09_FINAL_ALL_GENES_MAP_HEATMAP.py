
#!/usr/bin/env python3

# ============================================================
# FIGURE 9 — IMP AND MCR GLOBAL EPIDEMIOLOGY
#
# EXACTLY TWO OUTPUT FIGURES
#
# FIG 9A:
# Global map with country-income backgrounds
# Circle = blaIMP, square = mcr
# Source composition shown within both markers
# Country positive genome counts annotated
#
# FIG 9B:
# ALL automatically identified blaIMP and mcr alleles
# Seven ECC species x Human/Animal/Environment
# Cell values = positive genome counts
# Cell colors = within-stratum prevalence
# Right bars = allele-positive genome totals
#
# Also creates:
# Excel workbook for phylogenetic genome selection
# Separate IMP, MCR, dual-positive accession TXT files
#
# NO HARDCODED ALLELE COUNT OR TOP-N SELECTION
# ============================================================

from pathlib import Path
from collections import Counter
import re
import math
import urllib.request

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.patches import (
    Circle, Rectangle, Wedge, Patch
)
from matplotlib.colors import (
    LinearSegmentedColormap, Normalize
)

import geopandas as gpd
import pycountry

from adjustText import adjust_text

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

INCOME_DB = (
    ROOT / "Fig_01A_Global_Geographical_Distribution"
    / "Fig_01A_Country_Income_Database.xlsx"
)

OUT = ROOT / "Fig_09_IMP_MCR_FINAL"
OUT.mkdir(parents=True, exist_ok=True)

MAP_JPG = OUT / "Fig09A_FINAL_IMP_MCR_World_Map_600DPI.jpg"

HEAT_JPG = OUT / "Fig09B_FINAL_ALL_IMP_MCR_Allele_Heatmap_600DPI.jpg"

EXCEL = OUT / "Fig09_FINAL_IMP_MCR_Phylogeny_Master.xlsx"

GEO = OUT / "ne_110m_admin_0_countries.geojson"

GEO_URL = (
    "https://raw.githubusercontent.com/"
    "nvkelso/natural-earth-vector/master/"
    "geojson/ne_110m_admin_0_countries.geojson"
)

DPI = 600


# ============================================================
# COLORS — DO NOT CHANGE EXISTING CATEGORY MEANINGS
# ============================================================

SOURCES = [
    "Human",
    "Animal",
    "Environment"
]

SOURCE_COLORS = {
    "Human": "#D55E00",         # Reddish orange
    "Animal": "#56B4E9",        # Blue
    "Environment": "#009E73"    # Green
}

INCOME_COLORS = {
    "H": "#E8B91E",             # Strong yellow/gold
    "M": "#7E59CA",             # Purple
    "L+LM": "#CE669A",          # Pink
    "Unclassified": "#D9DDE1"   # Neutral grey
}

INCOME_NAMES = {
    "H": "High income",
    "M": "Upper-middle income",
    "L+LM": "Low + lower-middle income",
    "Unclassified": "Unclassified"
}

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 10,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
    "axes.spines.top": False,
    "axes.spines.right": False
})


# ============================================================
# METADATA CLEANING
# ============================================================

def clean_text(x):
    if pd.isna(x):
        return ""
    return re.sub(r"\s+", " ", str(x).strip())


def clean_key(x):
    return clean_text(x).casefold()


def clean_source(x):
    v = clean_key(x)

    if v == "human":
        return "Human"
    if v == "animal":
        return "Animal"
    if v == "environment":
        return "Environment"

    return "Excluded"


def clean_income(x):
    v = re.sub(
        r"[\s_-]+", "",
        clean_text(x).upper()
    )

    if v in {
        "H", "HIGH", "HIGHINCOME",
        "HIGHINCOMECOUNTRIES"
    }:
        return "H"

    if v in {
        "M", "UM", "UMIC",
        "UPPERMIDDLE",
        "UPPERMIDDLEINCOME",
        "UPPERMIDDLEINCOMECOUNTRIES"
    }:
        return "M"

    if v in {
        "L+LM", "LM+L", "LLM",
        "LOW", "LOWINCOME",
        "LOWERMIDDLE", "LOWERMIDDLEINCOME",
        "LOWERMIDDLE+LOW",
        "LOW+LOWERMIDDLE",
        "LOWERMIDDLEANDLOW",
        "LOWERMIDDLEANDLOWINCOME"
    }:
        return "L+LM"

    return "Unclassified"


ISO_ALIASES = {
    "united states": "USA",
    "united states of america": "USA",
    "usa": "USA",
    "united kingdom": "GBR",
    "uk": "GBR",
    "russia": "RUS",
    "russian federation": "RUS",
    "south korea": "KOR",
    "republic of korea": "KOR",
    "north korea": "PRK",
    "taiwan": "TWN",
    "vietnam": "VNM",
    "viet nam": "VNM",
    "iran": "IRN",
    "czech republic": "CZE",
    "czechia": "CZE",
    "ivory coast": "CIV",
    "cote d'ivoire": "CIV",
    "democratic republic of the congo": "COD",
    "democratic republic of congo": "COD",
    "dr congo": "COD",
    "republic of the congo": "COG",
    "congo": "COG",
    "tanzania": "TZA",
    "bolivia": "BOL",
    "venezuela": "VEN",
    "laos": "LAO",
    "syria": "SYR",
    "turkey": "TUR",
    "türkiye": "TUR",
    "kosovo": "XKX",
    "hong kong": "HKG",
    "macao": "MAC",
}


def get_iso3(country):
    value = clean_text(country)
    k = value.casefold()

    if not k:
        return ""

    if k in ISO_ALIASES:
        return ISO_ALIASES[k]

    try:
        return pycountry.countries.lookup(value).alpha_3
    except LookupError:
        return ""


# ============================================================
# EXACT ALLELE MATCHING
# ============================================================

IMP_RE = re.compile(
    r"^bla[\s_-]*IMP[\s_-]*(\d+[A-Za-z]?)$",
    re.I
)

MCR_RE = re.compile(
    r"^mcr[\s_-]*(10|[1-9])"
    r"((?:[._-]\d+)*)$",
    re.I
)


def extract_alleles(value):
    imp = set()
    mcr = set()
    review = set()

    text = "" if pd.isna(value) else str(value)

    for raw in re.split(r"[;,]", text):
        gene = raw.strip().strip("'\"")

        if not gene:
            continue

        a = IMP_RE.fullmatch(gene)
        b = MCR_RE.fullmatch(gene)

        if a:
            imp.add(
                "blaIMP-" + a.group(1).upper()
            )
        elif b:
            suffix = b.group(2).replace("_", ".").replace("-", ".")
            mcr.add(
                "mcr-" + b.group(1) + suffix
            )
        elif re.match(r"^bla[\s_-]*IMP", gene, re.I):
            review.add(gene)
        elif re.match(r"^mcr[\s_-]*\d", gene, re.I):
            review.add(gene)

    return (
        sorted(imp),
        sorted(mcr),
        sorted(review)
    )


def short_species(species):
    return species.replace("Enterobacter ", "E. ")


def save_jpg(fig, path):
    fig.savefig(
        path,
        dpi=DPI,
        format="jpeg",
        facecolor="white",
        pil_kwargs={
            "quality": 94,
            "subsampling": 0
        }
    )
    plt.close(fig)
    print("[SAVED]", path)


def style_excel(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid", fgColor="173352"
            )
            cell.font = Font(
                bold=True, color="FFFFFF"
            )

        for i, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(i)
            ].width = min(
                46, max(16, len(str(cell.value)) + 3)
            )


# ============================================================
# LOAD MASTER AND INCOME REFERENCE
# ============================================================

print("=" * 100)
print("FIGURE 9A + 9B — AUTOMATIC ALL-ALLELE ANALYSIS")
print("=" * 100)

if not MASTER.is_file():
    raise FileNotFoundError(MASTER)

if not INCOME_DB.is_file():
    raise FileNotFoundError(INCOME_DB)

cols = [
    "Assembly_Accession",
    "Original_NCBI_Accession",
    "Genome_File",
    "Normalized_Species",
    "MLST_ST",
    "Country_Harmonized",
    "Final_Source_Category",
    "Collection_Year_Numeric",
    "AMRFinder_Status",
    "AMR_Genes"
]

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=cols,
    dtype=str,
    engine="openpyxl"
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9,157 genomes; found {len(df)}"
    )

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicated assembly accessions."
    )

df["Species"] = (
    df["Normalized_Species"]
    .fillna("").str.strip()
)

df["Source"] = (
    df["Final_Source_Category"]
    .map(clean_source)
)

df["Country"] = (
    df["Country_Harmonized"]
    .map(clean_text)
)

df["Country_Key"] = df["Country"].map(clean_key)
df["ISO3"] = df["Country"].map(get_iso3)

df["Year"] = pd.to_numeric(
    df["Collection_Year_Numeric"],
    errors="coerce"
)

income = pd.read_excel(
    INCOME_DB,
    sheet_name="COUNTRY_INCOME_COUNTS",
    dtype=str,
    engine="openpyxl"
)

for column in ["Country_Harmonized", "Income_Group"]:
    if column not in income:
        raise ValueError(
            "Missing income column: " + column
        )

income["Key"] = (
    income["Country_Harmonized"].map(clean_key)
)
income["Income"] = (
    income["Income_Group"].map(clean_income)
)

income = income[income.Key.ne("")].copy()

conflicts = (
    income.groupby("Key")["Income"].nunique()
)

if (conflicts > 1).any():
    raise ValueError(
        "Conflicting country-income classification: "
        + str(conflicts[conflicts > 1].index.tolist()[:10])
    )

income_lookup = (
    income.drop_duplicates("Key")
    .set_index("Key")["Income"]
)

df["Income"] = (
    df["Country_Key"]
    .map(income_lookup)
    .fillna("Unclassified")
)

eligible = df[df.Source.isin(SOURCES)].copy()

calls = eligible.AMR_Genes.map(extract_alleles)

eligible["IMP_List"] = calls.map(lambda x: x[0])
eligible["MCR_List"] = calls.map(lambda x: x[1])
eligible["Review_List"] = calls.map(lambda x: x[2])

eligible["IMP_Pos"] = eligible.IMP_List.map(bool)
eligible["MCR_Pos"] = eligible.MCR_List.map(bool)

eligible["IMP_Alleles"] = (
    eligible.IMP_List.map(lambda x: "; ".join(x))
)
eligible["MCR_Alleles"] = (
    eligible.MCR_List.map(lambda x: "; ".join(x))
)

eligible["Gene_Group"] = np.select(
    [
        eligible.IMP_Pos & eligible.MCR_Pos,
        eligible.IMP_Pos,
        eligible.MCR_Pos
    ],
    [
        "IMP + mcr",
        "IMP only",
        "mcr only"
    ],
    default="Neither"
)

positive = eligible[
    eligible.IMP_Pos | eligible.MCR_Pos
].copy()

if positive.empty:
    raise ValueError(
        "No IMP/mcr-positive genomes were identified. "
        "Check AMR_Genes formatting."
    )

print("Eligible Human/Animal/Environment:", len(eligible))
print("IMP positive:", int(eligible.IMP_Pos.sum()))
print("mcr positive:", int(eligible.MCR_Pos.sum()))
print(
    "IMP and mcr positive:",
    int((eligible.IMP_Pos & eligible.MCR_Pos).sum())
)


# ============================================================
# AUTOMATIC ALLELE TABLE — NO TOP-N FILTER
# ============================================================

long_rows = []

for row in eligible.itertuples(index=False):
    for allele in row.IMP_List:
        long_rows.append({
            "Assembly_Accession": row.Assembly_Accession,
            "Species": row.Species,
            "Source": row.Source,
            "Country": row.Country,
            "ISO3": row.ISO3,
            "Income": row.Income,
            "Family": "IMP",
            "Allele": allele
        })

    for allele in row.MCR_List:
        long_rows.append({
            "Assembly_Accession": row.Assembly_Accession,
            "Species": row.Species,
            "Source": row.Source,
            "Country": row.Country,
            "ISO3": row.ISO3,
            "Income": row.Income,
            "Family": "mcr",
            "Allele": allele
        })

long = pd.DataFrame(long_rows)

allele_summary = (
    long.groupby(["Family", "Allele"])
    .Assembly_Accession.nunique()
    .reset_index(name="Positive_Genomes")
    .sort_values(
        ["Family", "Positive_Genomes", "Allele"],
        ascending=[True, False, True]
    )
)

imp_names = (
    allele_summary[
        allele_summary.Family.eq("IMP")
    ].Allele.tolist()
)

mcr_names = (
    allele_summary[
        allele_summary.Family.eq("mcr")
    ].Allele.tolist()
)

all_names = imp_names + mcr_names

print("Distinct IMP alleles:", len(imp_names))
print("Distinct mcr alleles:", len(mcr_names))
print("Total heatmap rows:", len(all_names))

counts = (
    long.groupby(["Allele", "Species", "Source"])
    .Assembly_Accession.nunique()
)

denoms = (
    eligible.groupby(["Species", "Source"])
    .size()
)

columns = [
    (species, source)
    for species in SPECIES
    for source in SOURCES
]

NROWS = len(all_names)
NCOLS = len(columns)

matrix_n = np.zeros(
    (NROWS, NCOLS), dtype=int
)

matrix_pct = np.full(
    (NROWS, NCOLS), np.nan
)

matrix_rows = []

for i, allele in enumerate(all_names):
    for j, (sp, src) in enumerate(columns):
        n = int(
            counts.get((allele, sp, src), 0)
        )
        total = int(
            denoms.get((sp, src), 0)
        )

        matrix_n[i, j] = n

        if total:
            matrix_pct[i, j] = 100*n/total

        matrix_rows.append({
            "Allele": allele,
            "Species": sp,
            "Source": src,
            "Positive_Genomes": n,
            "Total_Genomes": total,
            "Prevalence_Percent": (
                100*n/total if total else np.nan
            )
        })

matrix_excel = pd.DataFrame(matrix_rows)

allele_totals = [
    int(
        long.loc[
            long.Allele.eq(allele),
            "Assembly_Accession"
        ].nunique()
    )
    for allele in all_names
]


# ============================================================
# COUNTRY-LEVEL GENE AND SOURCE COUNTS
# ============================================================

country_records = []

for (iso, country), group in positive.groupby(
    ["ISO3", "Country"]
):
    if not iso:
        continue

    for family, flag in [
        ("IMP", "IMP_Pos"),
        ("mcr", "MCR_Pos")
    ]:
        part = group[group[flag]]

        if part.empty:
            continue

        vals = (
            part.Source.value_counts()
            .reindex(SOURCES, fill_value=0)
        )

        country_records.append({
            "Country": country,
            "ISO3": iso,
            "Income": part.Income.mode().iloc[0],
            "Family": family,
            "Positive_Genomes": len(part),
            "Human": int(vals["Human"]),
            "Animal": int(vals["Animal"]),
            "Environment": int(vals["Environment"])
        })

country_markers = pd.DataFrame(country_records)


# ============================================================
# NATURAL EARTH WORLD MAP
# ============================================================

if not GEO.is_file():
    print("Downloading Natural Earth map geometry...")
    urllib.request.urlretrieve(GEO_URL, GEO)

world = gpd.read_file(GEO)

if "ADM0_A3" in world.columns:
    world["ISO3"] = world["ADM0_A3"].astype(str)
elif "ISO_A3" in world.columns:
    world["ISO3"] = world["ISO_A3"].astype(str)
else:
    raise ValueError(
        "Natural Earth country-code field not found."
    )

world["ISO3"] = world["ISO3"].replace({
    "KOS": "XKX"
})

world = world[
    world.ISO3.ne("ATA")
].copy()

points = world.geometry.representative_point()

world["X"] = points.x
world["Y"] = points.y

# Country-income backgrounds use the entire frozen
# reference table, not just positive-genome countries.

income_iso = income.copy()

income_iso["ISO3"] = (
    income_iso["Country_Harmonized"]
    .map(get_iso3)
)

income_iso = income_iso[
    income_iso.ISO3.ne("")
]

iso_conflicts = (
    income_iso.groupby("ISO3")
    .Income.nunique()
)

if (iso_conflicts > 1).any():
    raise ValueError(
        "Conflicting ISO-level income assignments: "
        + str(iso_conflicts[iso_conflicts > 1].index.tolist())
    )

world_income = (
    income_iso.drop_duplicates("ISO3")
    .set_index("ISO3")["Income"]
)

world["Income"] = (
    world.ISO3.map(world_income)
    .fillna("Unclassified")
)

locations = (
    world.drop_duplicates("ISO3")
    .set_index("ISO3")[["X", "Y"]]
)

unmapped = country_markers[
    ~country_markers.ISO3.isin(locations.index)
].copy()

mapped = country_markers[
    country_markers.ISO3.isin(locations.index)
].copy()

print("Country markers requiring location review:",
      len(unmapped))


# ============================================================
# FIGURE 9A — GLOBAL MAP
# ============================================================

def pie_circle(ax, x, y, radius, counts):
    total = sum(counts.values())
    if not total:
        return

    angle = 90

    for source in SOURCES:
        n = counts[source]

        if not n:
            continue

        sweep = 360 * n / total

        ax.add_patch(Wedge(
            (x, y), radius,
            angle, angle+sweep,
            facecolor=SOURCE_COLORS[source],
            edgecolor="white",
            linewidth=0.48,
            zorder=6
        ))

        angle += sweep

    ax.add_patch(Circle(
        (x, y), radius,
        facecolor="none",
        edgecolor="#25374A",
        linewidth=0.65,
        zorder=8
    ))


def stacked_square(ax, x, y, size, counts):
    total = sum(counts.values())
    if not total:
        return

    left = x-size/2
    cursor = left

    # Side-by-side source composition, clipped to square.
    for source in SOURCES:
        n = counts[source]
        if not n:
            continue

        width = size*n/total

        ax.add_patch(Rectangle(
            (cursor, y-size/2),
            width, size,
            facecolor=SOURCE_COLORS[source],
            edgecolor="none",
            zorder=6
        ))

        cursor += width

    ax.add_patch(Rectangle(
        (left, y-size/2),
        size, size,
        fill=False,
        edgecolor="#25374A",
        linewidth=0.75,
        zorder=8
    ))


fig = plt.figure(
    figsize=(22, 12.5),
    facecolor="white"
)

ax = fig.add_axes([
    0.025, 0.175, 0.95, 0.710
])

for category in [
    "Unclassified", "H", "M", "L+LM"
]:
    sub = world[
        world.Income.eq(category)
    ]

    if not sub.empty:
        sub.plot(
            ax=ax,
            color=INCOME_COLORS[category],
            edgecolor="white",
            linewidth=0.32,
            zorder=1
        )

ax.set_xlim(-172, 191)
ax.set_ylim(-60, 85)
ax.axis("off")
ax.set_aspect("auto")

maximum = (
    mapped.Positive_Genomes.max()
    if not mapped.empty else 1
)

country_labels = []
marker_output = []

for iso, group in mapped.groupby("ISO3"):

    x = float(locations.loc[iso, "X"])
    y = float(locations.loc[iso, "Y"])

    has_two = len(group) == 2

    shift = 2.2 if has_two else 0

    for row in group.itertuples(index=False):

        n = int(row.Positive_Genomes)

        radius = (
            0.75 + 4.5*math.sqrt(n/maximum)
        )

        xpos = x

        if has_two:
            xpos += (
                -shift if row.Family == "IMP"
                else shift
            )

        values = {
            source: int(getattr(row, source))
            for source in SOURCES
        }

        if row.Family == "IMP":
            pie_circle(
                ax, xpos, y, radius, values
            )
        else:
            stacked_square(
                ax, xpos, y, radius*1.75, values
            )

        marker_output.append({
            "Country": row.Country,
            "ISO3": iso,
            "Family": row.Family,
            "Positive_Genomes": n,
            "Marker_X": xpos,
            "Marker_Y": y
        })

# Label the largest positive countries, keeping all
# smaller-country counts in the Excel map table.
label_data = (
    positive.groupby(["Country", "ISO3"])
    .agg(
        IMP=("IMP_Pos", "sum"),
        MCR=("MCR_Pos", "sum")
    )
    .reset_index()
)

label_data["Any_Positive"] = (
    label_data["IMP"]
    + label_data["MCR"]
)

label_data = label_data[
    label_data.ISO3.isin(locations.index)
].nlargest(12, "Any_Positive")

texts = []

for row in label_data.itertuples(index=False):

    x = float(locations.loc[row.ISO3, "X"])
    y = float(locations.loc[row.ISO3, "Y"])

    label = (
        f"{row.Country}\n"
        f"IMP {int(row.IMP):,} | "
        f"mcr {int(row.MCR):,}"
    )

    t = ax.text(
        x, y+7.0,
        label,
        fontsize=8.2,
        fontweight="bold",
        ha="center",
        va="center",
        color="#152E43",
        bbox=dict(
            boxstyle="round,pad=0.18",
            facecolor="white",
            edgecolor="none",
            alpha=0.88
        ),
        zorder=12
    )

    texts.append(t)

if texts:
    adjust_text(
        texts,
        ax=ax,
        only_move={"text": "xy"},
        force_text=(0.6, 0.8),
        force_static=(0.4, 0.5),
        arrowprops=dict(
            arrowstyle="-",
            color="#51667A",
            lw=0.55
        )
    )

fig.text(
    0.5, .963,
    "Global distribution of blaIMP- and mcr-positive "
    "Enterobacter cloacae complex genomes",
    ha="center",
    fontsize=20,
    fontweight="bold",
    color="#18324C"
)

fig.text(
    0.5, .930,
    "Circle: blaIMP | Square: mcr | "
    "Human, animal and environmental genomes",
    ha="center",
    fontsize=11.5,
    color="#526B80"
)

legend_income = [
    Patch(
        facecolor=INCOME_COLORS[i],
        edgecolor="none",
        label=INCOME_NAMES[i]
    )
    for i in INCOME_COLORS
]

legend_source = [
    Patch(
        facecolor=SOURCE_COLORS[s],
        edgecolor="none",
        label=s
    )
    for s in SOURCES
]

fig.legend(
    handles=legend_income,
    loc="lower left",
    bbox_to_anchor=(0.04, .080),
    ncol=2,
    fontsize=9.5,
    frameon=False
)

fig.legend(
    handles=legend_source,
    loc="lower center",
    bbox_to_anchor=(.66, .099),
    ncol=3,
    fontsize=10,
    frameon=False
)

fig.text(
    .65, .062,
    "Marker size reflects positive genome count; "
    "colored segments represent the source distribution.",
    fontsize=9.3,
    color="#465D70",
    ha="center"
)

fig.text(
    .5, .029,
    "Countries without eligible reference metadata "
    "are shown in grey. Smaller-country counts are "
    "available in the associated Excel table.",
    ha="center",
    fontsize=8.8,
    color="#63778B"
)

save_jpg(fig, MAP_JPG)


# ============================================================
# FIGURE 9B — ALL IMP AND MCR ALLELES
# ============================================================

# Height automatically expands according to the number
# of alleles found; no variants are manually excluded.

FIG_H = max(
    11.0,
    4.5 + 0.44*NROWS
)

fig = plt.figure(
    figsize=(24.5, FIG_H),
    facecolor="white"
)

# Reserved areas:
# left: complete allele names
# top: species and source annotations
# right: allele-positive total bars
# bottom: colorbar

ax = fig.add_axes([
    .17, .155, .635, .67
])

axb = fig.add_axes([
    .825, .155, .125, .67
])

header_ax = fig.add_axes([
    .17, .832, .635, .092
])

header_ax.set_xlim(-.5, NCOLS-.5)
header_ax.set_ylim(0, 2)
header_ax.axis("off")

finite = matrix_pct[
    np.isfinite(matrix_pct)
]

vmax = (
    max(1.0, float(finite.max()))
    if finite.size else 1.0
)

cmap = LinearSegmentedColormap.from_list(
    "ECC_ALLELES",
    [
        "#FFFFFF",
        "#DDEFFC",
        "#93CAE9",
        "#3892CC",
        "#00579F",
        "#003B75"
    ]
)

cmap.set_bad("#E1E5E8")

im = ax.imshow(
    np.ma.masked_invalid(matrix_pct),
    aspect="auto",
    cmap=cmap,
    norm=Normalize(0, vmax)
)

ax.set_xlim(-.5, NCOLS-.5)
ax.set_ylim(NROWS-.5, -.5)

ax.set_xticks([])
ax.set_yticks(np.arange(NROWS))
ax.set_yticklabels(
    all_names,
    fontsize=10,
    fontweight="bold"
)

ax.tick_params(axis="y", length=0, pad=8)

# Show counts in every cell.
for i in range(NROWS):
    for j in range(NCOLS):

        v = matrix_pct[i, j]
        n = matrix_n[i, j]

        if not np.isfinite(v):
            label = "–"
        else:
            label = str(n)

        color = (
            "white"
            if np.isfinite(v) and v > .55*vmax
            else "#17364E"
        )

        ax.text(
            j, i, label,
            ha="center",
            va="center",
            fontsize=8.0,
            color=color
        )

for j in np.arange(-.5, NCOLS, 1):
    ax.axvline(
        j,
        color="white",
        linewidth=.65
    )

for i in np.arange(-.5, NROWS, 1):
    ax.axhline(
        i,
        color="white",
        linewidth=.65
    )

for j in range(3, NCOLS, 3):
    ax.axvline(
        j-.5,
        color="#234A65",
        linewidth=1.7
    )

if len(imp_names) and len(mcr_names):
    ax.axhline(
        len(imp_names)-.5,
        color="#B23C75",
        linewidth=2.7
    )

# Separate IMP and mcr row-family strips.
for i, allele in enumerate(all_names):

    family_color = (
        "#0072B2"
        if allele.startswith("blaIMP")
        else "#CC79A7"
    )

    ax.add_patch(Rectangle(
        (-.66, i-.47),
        .12, .94,
        color=family_color,
        clip_on=False
    ))

# Header, isolated from the heatmap data region.
for j, (sp, src) in enumerate(columns):

    header_ax.add_patch(Rectangle(
        (j-.5, .10),
        1, .68,
        facecolor=SOURCE_COLORS[src],
        edgecolor="white",
        linewidth=.55
    ))

    header_ax.text(
        j, .44,
        src[0],
        ha="center",
        va="center",
        fontsize=10,
        fontweight="bold",
        color="white"
    )

for i, sp in enumerate(SPECIES):

    center = 3*i + 1

    n_species = sum(
        int(denoms.get((sp, s), 0))
        for s in SOURCES
    )

    header_ax.text(
        center, 1.56,
        short_species(sp)
        + f"\n(n={n_species:,})",
        fontsize=10,
        fontweight="bold",
        ha="center",
        va="center",
        color=SPECIES_COLORS[sp]
    )

# Right-side horizontal total count bars.
ypos = np.arange(NROWS)

bars_color = [
    "#0072B2" if a.startswith("blaIMP")
    else "#CC79A7"
    for a in all_names
]

axb.barh(
    ypos,
    allele_totals,
    height=.69,
    color=bars_color
)

max_total = max(allele_totals)

axb.set_xlim(
    0, max_total*1.28
)
axb.set_ylim(NROWS-.5, -.5)

axb.set_yticks([])
axb.set_xlabel(
    "Positive genomes",
    fontsize=10
)

for i, n in enumerate(allele_totals):
    axb.text(
        n + max_total*.025,
        i, f"{n:,}",
        fontsize=9,
        va="center",
        fontweight="bold",
        color="#173148"
    )

axb.spines[
    ["top", "right", "left"]
].set_visible(False)

fig.text(
    .5, .970,
    "Distribution of blaIMP and mcr variants "
    "across the Enterobacter cloacae complex",
    ha="center",
    fontsize=19,
    fontweight="bold",
    color="#18324C"
)

fig.text(
    .5, .942,
    f"All detected alleles: {len(imp_names)} blaIMP "
    f"and {len(mcr_names)} mcr | "
    "Human, animal and environmental sources",
    ha="center",
    fontsize=11,
    color="#53697E"
)

cax = fig.add_axes([
    .33, .065, .34, .018
])

fig.colorbar(
    im, cax=cax,
    orientation="horizontal"
)

cax.set_xlabel(
    "Genomes carrying each allele within "
    "the species–source stratum (%)",
    fontsize=10
)

fig.text(
    .5, .022,
    "Cell numbers are observed genome counts. "
    "Grey = no eligible genomes; zero = not detected. "
    "Every detected allele is retained.",
    ha="center",
    fontsize=9,
    color="#647A8B"
)

save_jpg(fig, HEAT_JPG)


# ============================================================
# EXPORT PHYLOGENY-READY EXCEL
# ============================================================

phylo_cols = [
    "Assembly_Accession",
    "Original_NCBI_Accession",
    "Genome_File",
    "Species",
    "MLST_ST",
    "Country",
    "ISO3",
    "Income",
    "Source",
    "Year",
    "IMP_Alleles",
    "MCR_Alleles",
    "Gene_Group",
]

positive_export = positive[
    phylo_cols
].copy()

positive_export["FASTA_Path_Exists"] = (
    positive_export.Genome_File.fillna("")
    .map(lambda p: bool(str(p).strip()) and Path(str(p)).is_file())
)

review = eligible[
    eligible.Review_List.map(bool)
].copy()

review["Review_Genes"] = (
    review.Review_List.map(lambda x: "; ".join(x))
)

qc = pd.DataFrame([
    ["Original master genomes", len(df)],
    ["Eligible Human/Animal/Environment", len(eligible)],
    ["IMP-positive genomes", int(eligible.IMP_Pos.sum())],
    ["mcr-positive genomes", int(eligible.MCR_Pos.sum())],
    ["Dual-positive genomes", int(
        (eligible.IMP_Pos & eligible.MCR_Pos).sum()
    )],
    ["Either-positive genomes", len(positive)],
    ["Detected IMP alleles", len(imp_names)],
    ["Detected mcr alleles", len(mcr_names)],
    ["Heatmap rows", NROWS],
    ["Heatmap columns", NCOLS],
    ["Unmapped map-marker rows", len(unmapped)],
    ["Allele review records", len(review)],
    ["Gene-positive interpretation", "Genomic detection, not phenotype"],
    ["Map circle", "blaIMP"],
    ["Map square", "mcr"],
    ["Map source colors", str(SOURCE_COLORS)],
    ["Country-income colors", str(INCOME_COLORS)],
], columns=["Parameter", "Value"])

with pd.ExcelWriter(
    EXCEL,
    engine="openpyxl"
) as writer:

    qc.to_excel(
        writer,
        sheet_name="QC",
        index=False
    )

    allele_summary.to_excel(
        writer,
        sheet_name="ALL_ALLELE_COUNTS",
        index=False
    )

    matrix_excel.to_excel(
        writer,
        sheet_name="FULL_ALLELE_MATRIX",
        index=False
    )

    positive_export.to_excel(
        writer,
        sheet_name="PHYLOGENY_ALL_POSITIVE",
        index=False
    )

    positive_export[
        positive.IMP_Pos.values
    ].to_excel(
        writer,
        sheet_name="IMP_GENOMES",
        index=False
    )

    positive_export[
        positive.MCR_Pos.values
    ].to_excel(
        writer,
        sheet_name="MCR_GENOMES",
        index=False
    )

    positive_export[
        (positive.IMP_Pos & positive.MCR_Pos).values
    ].to_excel(
        writer,
        sheet_name="DUAL_POSITIVE",
        index=False
    )

    country_markers.to_excel(
        writer,
        sheet_name="COUNTRY_MARKERS",
        index=False
    )

    pd.DataFrame(marker_output).to_excel(
        writer,
        sheet_name="MAP_COORDINATES",
        index=False
    )

    unmapped.to_excel(
        writer,
        sheet_name="UNMAPPED_COUNTRIES",
        index=False
    )

    review[
        ["Assembly_Accession", "Species", "Country",
         "Source", "Review_Genes"]
    ].to_excel(
        writer,
        sheet_name="GENE_NAME_REVIEW",
        index=False
    )

    style_excel(writer)

# Standalone accession lists.
for name, part in {
    "IMP": positive[positive.IMP_Pos],
    "MCR": positive[positive.MCR_Pos],
    "DUAL": positive[
        positive.IMP_Pos & positive.MCR_Pos
    ],
    "EITHER": positive
}.items():

    path = OUT / f"{name}_PHYLOGENY_ACCESSIONS.txt"

    accessions = part.Assembly_Accession.astype(str).tolist()

    path.write_text(
        "\n".join(accessions) + ("\n" if accessions else ""),
        encoding="utf-8"
    )

print()
print("=" * 100)
print("FIGURE 9 COMPLETED")
print("=" * 100)
print("Eligible genomes:", len(eligible))
print("IMP-positive genomes:", int(eligible.IMP_Pos.sum()))
print("mcr-positive genomes:", int(eligible.MCR_Pos.sum()))
print("All IMP variants:", len(imp_names))
print("All mcr variants:", len(mcr_names))
print("Global map:", MAP_JPG)
print("Allele heatmap:", HEAT_JPG)
print("Excel:", EXCEL)
print("=" * 100)
