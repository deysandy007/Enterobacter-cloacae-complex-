
#!/usr/bin/env python3

# ================================================================
# FIGURE 8B — ALL-ST HIERARCHICAL MLST + INCOME TREEMAP
#
# One combined landscape figure, seven ECC species.
# EVERY valid ST gets a separate rectangle.
# NO "Other STs" grouping.
#
# Species block area = ST-assigned genome count
# ST rectangle area  = genomes assigned to that ST
# Background color   = species
# Internal strip     = income composition of that ST
# Dark border        = ST in >=2 classified income groups
#
# Outputs:
# - One 600-DPI JPEG
# - Complete Excel workbook with ST x income counts
#
# Uses existing 9,157-genome master and frozen Figure 1A
# country-income database.
# ================================================================

from pathlib import Path
import re
import gc
import math

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Patch
from matplotlib.colors import to_rgb

import squarify

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ================================================================
# PATHS
# ================================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

INCOME_FILE = (
    ROOT / "Fig_01A_Global_Geographical_Distribution"
    / "Fig_01A_Country_Income_Database.xlsx"
)

OUT = (
    ROOT / "Fig_08B_Combined_MLST_Treemap"
    / "INCOME_STRATIFIED_ALL_ST"
)
OUT.mkdir(parents=True, exist_ok=True)

JPG = OUT / "Fig08B_All_ST_Income_Combined_600DPI.jpg"
XLSX = OUT / "Fig08B_All_ST_Income_Complete_Data.xlsx"

DPI = 600

# Landscape. Keep the aspect ratio close to the previous
# combined figure but allow additional label space.
FIG_W = 24.0
FIG_H = 17.0

# ST tiles must exceed both size thresholds to receive
# an income strip. Tiny tiles still remain in the treemap.
MIN_STRIP_WIDTH_IN = 0.58
MIN_STRIP_HEIGHT_IN = 0.42

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

# Previous combined-treemap arrangement.
SPECIES_ROWS = [
    ["Enterobacter hormaechei"],
    [
        "Enterobacter cloacae",
        "Enterobacter roggenkampii",
        "Enterobacter asburiae",
    ],
    [
        "Enterobacter kobei",
        "Enterobacter ludwigii",
        "Enterobacter bugandensis",
    ],
]

INCOME_GROUPS = ["H", "M", "L+LM", "Unclassified"]

INCOME_COLORS = {
    "H": "#C98A00",
    "M": "#673AB7",
    "L+LM": "#C94F82",
    "Unclassified": "#929DA8",
}

INCOME_LABELS = {
    "H": "High income",
    "M": "Upper-middle income",
    "L+LM": "Lower-middle + low income",
    "Unclassified": "Income unclassified",
}

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 11,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

# ================================================================
# DATA UTILITIES
# ================================================================

def clean_st(value):
    if pd.isna(value):
        return ""

    s = str(value).strip().strip("'\"").strip()

    if s.upper().startswith("ST"):
        s = s[2:].strip()

    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]

    if not re.fullmatch(r"\d+", s):
        return ""

    number = int(s)
    return f"ST{number}" if number > 0 else ""


def country_key(value):
    if pd.isna(value):
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip()
    ).casefold()


def normalize_income(value):
    if pd.isna(value):
        return "Unclassified"

    s = str(value).strip().upper()
    s = re.sub(r"[\s_\-]+", "", s)

    mapping = {
        "H": "H",
        "HIGH": "H",
        "HIGHINCOME": "H",
        "HIGHINCOMECOUNTRIES": "H",

        "M": "M",
        "UM": "M",
        "UMIC": "M",
        "UPPERMIDDLE": "M",
        "UPPERMIDDLEINCOME": "M",
        "UPPERMIDDLEINCOMECOUNTRIES": "M",

        "L+LM": "L+LM",
        "LM+L": "L+LM",
        "LLM": "L+LM",
        "LMICLIC": "L+LM",
        "LOWERMIDDLE+LOW": "L+LM",
        "LOW+LOWERMIDDLE": "L+LM",
        "LOWERMIDDLEANDLOW": "L+LM",
        "LOWERMIDDLEANDLOWINCOME": "L+LM",

        # These are both included in the combined L+LM bin.
        "LOWERMIDDLE": "L+LM",
        "LOWERMIDDLEINCOME": "L+LM",
        "LOW": "L+LM",
        "LOWINCOME": "L+LM",
    }

    return mapping.get(s, "Unclassified")


def lighten(hex_color, amount):
    rgb = np.array(to_rgb(hex_color), dtype=float)

    return tuple(
        np.clip(
            rgb + (1.0 - rgb) * amount,
            0.0, 1.0
        )
    )


def text_color(face):
    r, g, b = to_rgb(face)

    luminosity = (
        0.2126 * r
        + 0.7152 * g
        + 0.0722 * b
    )

    return "white" if luminosity < 0.47 else "#173148"


def pcent(n, denominator):
    return (
        100.0 * n / denominator
        if denominator else np.nan
    )


def style_excel(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="173352"
            )
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

        for j, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(j)
            ].width = min(
                46,
                max(17, len(str(cell.value)) + 4)
            )


# ================================================================
# READ MASTER
# ================================================================

print("=" * 100)
print("FIGURE 8B — MLST ABUNDANCE + INCOME COMPOSITION")
print("=" * 100)

for file in [MASTER, INCOME_FILE]:
    if not file.is_file():
        raise FileNotFoundError(file)

master_columns = [
    "Assembly_Accession",
    "Normalized_Species",
    "MLST_Scheme",
    "MLST_ST",
    "Country_Harmonized",
]

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    engine="openpyxl",
    usecols=master_columns,
    dtype=str,
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9157 genomes; found {len(df)}"
    )

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicated assembly accessions"
    )

df["Species"] = (
    df["Normalized_Species"]
    .fillna("").str.strip()
)

df["Scheme"] = (
    df["MLST_Scheme"]
    .fillna("").str.strip().str.lower()
)

df["ST"] = df["MLST_ST"].map(clean_st)

df["Country_Key"] = (
    df["Country_Harmonized"]
    .map(country_key)
)

df["Resolved"] = df["Species"].isin(SPECIES)

df["Valid_ST"] = (
    df["ST"].ne("")
    & df["Scheme"].eq("ecloacae")
)

df["Eligible"] = (
    df["Resolved"] & df["Valid_ST"]
)

resolved = df[df["Resolved"]].copy()
typed = df[df["Eligible"]].copy()

if len(resolved) != 8882:
    raise ValueError(
        f"Expected 8882 resolved genomes; found {len(resolved)}"
    )

if typed.empty:
    raise ValueError("No valid ST-assigned genomes")

# ================================================================
# READ FROZEN INCOME MAPPING
# ================================================================

income = pd.read_excel(
    INCOME_FILE,
    sheet_name="COUNTRY_INCOME_COUNTS",
    engine="openpyxl",
    dtype=str,
)

for col in ["Country_Harmonized", "Income_Group"]:
    if col not in income.columns:
        raise ValueError(
            f"Income database missing column: {col}\n"
            f"Columns: {list(income.columns)}"
        )

income["Country_Key"] = (
    income["Country_Harmonized"]
    .map(country_key)
)

income["Income"] = (
    income["Income_Group"]
    .map(normalize_income)
)

# Allow repeated rows only if their classification agrees.
conflicts = (
    income[income["Country_Key"].ne("")]
    .groupby("Country_Key")["Income"]
    .nunique()
)

bad_keys = conflicts[conflicts > 1].index.tolist()

if bad_keys:
    raise ValueError(
        "Conflicting income categories for countries: "
        + ", ".join(bad_keys[:20])
    )

income_map = (
    income[income["Country_Key"].ne("")]
    .drop_duplicates("Country_Key")
    .set_index("Country_Key")["Income"]
)

df["Income"] = (
    df["Country_Key"]
    .map(income_map)
    .fillna("Unclassified")
)

resolved = df[df["Resolved"]].copy()
typed = df[df["Eligible"]].copy()

print("Master genomes:", len(df))
print("Resolved genomes:", len(resolved))
print("ST-assigned genomes:", len(typed))

print("\nIncome groups among typed genomes:")
print(typed["Income"].value_counts().to_string())

# ================================================================
# COMPLETE ST ABUNDANCE AND INCOME COUNTS
# ================================================================

species_all = (
    resolved.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

species_typed = (
    typed.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

abundance = (
    typed.groupby(["Species", "ST"])
    .agg(
        Genome_Count=("Assembly_Accession", "nunique")
    )
    .reset_index()
)

abundance["ST_Number"] = (
    abundance["ST"].str[2:].astype(int)
)

income_counts = (
    typed.groupby(["Species", "ST", "Income"])
    .size()
    .unstack("Income", fill_value=0)
    .reindex(columns=INCOME_GROUPS, fill_value=0)
    .reset_index()
)

income_counts = income_counts.rename(
    columns={
        "H": "N_H",
        "M": "N_M",
        "L+LM": "N_LLM",
        "Unclassified": "N_Unclassified",
    }
)

all_st = abundance.merge(
    income_counts,
    on=["Species", "ST"],
    how="left",
    validate="one_to_one",
)

all_st["Species_Total"] = (
    all_st["Species"].map(species_all)
)

all_st["Species_Typed"] = (
    all_st["Species"].map(species_typed)
)

all_st["Percent_Within_Species"] = (
    100.0 * all_st["Genome_Count"]
    / all_st["Species_Total"]
)

all_st["Percent_Among_Typed"] = (
    100.0 * all_st["Genome_Count"]
    / all_st["Species_Typed"]
)

for name, col in [
    ("H", "N_H"),
    ("M", "N_M"),
    ("LLM", "N_LLM"),
    ("Unclassified", "N_Unclassified"),
]:
    all_st[f"Income_{name}_Percent"] = (
        100.0 * all_st[col]
        / all_st["Genome_Count"]
    )

all_st["Income_Breadth"] = (
    (all_st[["N_H", "N_M", "N_LLM"]] > 0)
    .sum(axis=1)
)

all_st["Cross_Income_ST"] = (
    all_st["Income_Breadth"] >= 2
)

all_st = all_st.sort_values(
    ["Species", "Genome_Count", "ST_Number"],
    ascending=[True, False, True]
).reset_index(drop=True)

all_st["ST_Abundance_Rank"] = (
    all_st.groupby("Species")
    .cumcount() + 1
)

# Numerical validation
income_columns = [
    "N_H", "N_M", "N_LLM", "N_Unclassified"
]

if not (
    all_st[income_columns].sum(axis=1)
    == all_st["Genome_Count"]
).all():
    raise ValueError(
        "Income subcounts do not match ST genome counts."
    )

if int(all_st["Genome_Count"].sum()) != len(typed):
    raise ValueError(
        "Total ST counts do not match typed genomes."
    )

# ================================================================
# FIGURE CANVAS AND LABEL FUNCTION
# ================================================================

fig = plt.figure(
    figsize=(FIG_W, FIG_H),
    facecolor="white"
)

# Make the actual treemap large and retain a footer
# for the income legend.
AX_LEFT = 0.025
AX_BOTTOM = 0.135
AX_WIDTH = 0.950
AX_HEIGHT = 0.775

ax = fig.add_axes([
    AX_LEFT, AX_BOTTOM,
    AX_WIDTH, AX_HEIGHT
])

AX_W_IN = FIG_W * AX_WIDTH
AX_H_IN = FIG_H * AX_HEIGHT

CANVAS_W = 1000.0
CANVAS_H = 720.0

ax.set_xlim(0, CANVAS_W)
ax.set_ylim(0, CANVAS_H)
ax.invert_yaxis()
ax.set_aspect("auto")
ax.axis("off")

tile_rows = []
block_rows = []


def draw_st_label(
    patch, x, y, w, h, st, n, pct, face,
    strip_drawn=False
):
    """
    Labels are fitted and clipped inside their tiles.
    Returns a readable status for the Excel audit.
    """

    width_in = w / CANVAS_W * AX_W_IN
    height_in = h / CANVAS_H * AX_H_IN

    # Some vertical height is reserved for income strip.
    if strip_drawn:
        height_in *= 0.68

    if width_in >= 0.95 and height_in >= 0.67:
        label = f"{st}\nn={n:,}\n{pct:.1f}%"
        fs = min(13.0, width_in * 11, height_in * 13)

    elif width_in >= 0.63 and height_in >= 0.36:
        label = f"{st}\nn={n:,}"
        fs = min(10.3, width_in * 12, height_in * 15)

    elif width_in >= 0.41 and height_in >= 0.21:
        label = st
        fs = min(8.0, width_in * 12, height_in * 17)

    else:
        return "Unlabeled: small tile"

    if fs < 5.8:
        return "Unlabeled: small tile"

    label_y = (
        y + h*0.42
        if strip_drawn else y + h*0.50
    )

    artist = ax.text(
        x + w/2,
        label_y,
        label,
        ha="center",
        va="center",
        fontsize=fs,
        fontweight="bold",
        color=text_color(face),
        linespacing=1.14,
        zorder=10,
    )

    artist.set_clip_path(patch)

    return "Labeled"


def draw_income_strip(x, y, w, h, row, species):
    """
    Draw each income segment proportional to its actual
    share among genomes assigned to this individual ST.
    """

    width_in = w / CANVAS_W * AX_W_IN
    height_in = h / CANVAS_H * AX_H_IN

    if (
        width_in < MIN_STRIP_WIDTH_IN
        or height_in < MIN_STRIP_HEIGHT_IN
    ):
        return False

    # Narrow strip near the lower part of the tile.
    strip_height = min(
        h*0.20,
        12.0
    )

    strip_x = x + w*0.055
    strip_y = y + h*0.75

    strip_w = w*0.89

    counts = [
        int(row.N_H),
        int(row.N_M),
        int(row.N_LLM),
        int(row.N_Unclassified)
    ]

    cumulative_x = strip_x
    total = int(row.Genome_Count)

    for income, count in zip(INCOME_GROUPS, counts):

        if count <= 0:
            continue

        segment_w = strip_w * count / total

        ax.add_patch(
            Rectangle(
                (cumulative_x, strip_y),
                segment_w,
                strip_height,
                facecolor=INCOME_COLORS[income],
                edgecolor="none",
                zorder=11
            )
        )

        cumulative_x += segment_w

    # Thin outline makes the strip visible over pale tiles.
    ax.add_patch(
        Rectangle(
            (strip_x, strip_y),
            strip_w,
            strip_height,
            fill=False,
            edgecolor="#46576A",
            linewidth=0.35,
            zorder=12
        )
    )

    return True


# ================================================================
# HIERARCHICAL TREEMAP — ALL STs
# ================================================================

TOTAL_TYPED = len(typed)
current_y = 0.0

for species_group in SPECIES_ROWS:

    row_total = sum(
        int(species_typed[s])
        for s in species_group
    )

    if row_total == 0:
        continue

    row_height = (
        CANVAS_H * row_total / TOTAL_TYPED
    )

    current_x = 0.0

    for species in species_group:

        typed_n = int(species_typed[species])
        all_n = int(species_all[species])

        if typed_n == 0:
            continue

        block_w = CANVAS_W * typed_n / row_total
        block_h = row_height

        x, y = current_x, current_y

        base = SPECIES_COLORS[species]

        # Main species background and boundaries.
        ax.add_patch(
            Rectangle(
                (x+1, y+1),
                max(.1, block_w-2),
                max(.1, block_h-2),
                facecolor=base,
                edgecolor="white",
                linewidth=2.3,
                zorder=1
            )
        )

        header_h = min(
            30.0,
            block_h * 0.22
        )

        species_pct = pcent(
            typed_n, TOTAL_TYPED
        )

        # Use full name for large blocks, abbreviated
        # only where the smaller block cannot accommodate it.
        if block_w >= 340:
            species_title = (
                f"{species}  "
                f"(n={typed_n:,}; {species_pct:.1f}%)"
            )
            species_fs = 14.0
        else:
            species_title = (
                species.replace("Enterobacter ", "E. ")
                + f"  (n={typed_n:,})"
            )
            species_fs = (
                10.5 if block_w > 180 else 8.5
            )

        ax.text(
            x+9,
            y+header_h*.54,
            species_title,
            ha="left",
            va="center",
            color="white",
            fontsize=species_fs,
            fontstyle="italic",
            fontweight="bold",
            zorder=8
        )

        # Every ST occupies its own squarified tile.
        px = x + 6
        py = y + header_h + 2

        pw = max(.1, block_w-12)
        ph = max(.1, block_h-header_h-8)

        sp = all_st[
            all_st["Species"].eq(species)
        ].copy()

        sp = sp.sort_values(
            ["Genome_Count", "ST_Number"],
            ascending=[False, True]
        ).reset_index(drop=True)

        sizes = sp["Genome_Count"].astype(float).tolist()

        rectangles = squarify.squarify(
            squarify.normalize_sizes(
                sizes, pw, ph
            ),
            px, py, pw, ph
        )

        if len(rectangles) != len(sp):
            raise ValueError(
                f"Tile count mismatch for {species}"
            )

        for i, (r, row) in enumerate(
            zip(rectangles, sp.itertuples(index=False))
        ):

            rx, ry = float(r["x"]), float(r["y"])
            rw, rh = float(r["dx"]), float(r["dy"])

            # Species-specific tone remains the main fill.
            shade = (
                0.08
                + 0.67 * i/max(1, len(sp)-1)
            )

            face = lighten(base, shade)

            # Cross-income ST: thin dark boundary.
            # Single-income ST: subtle white boundary.
            if int(row.Income_Breadth) >= 2:
                border_color = "#1D3043"
                line_width = 0.9
            else:
                border_color = "white"
                line_width = 0.35

            # Reduce boundaries for extremely tiny tiles
            # so their fill remains visible.
            if rw < 1.3 or rh < 1.3:
                line_width = min(line_width, 0.18)

            patch = Rectangle(
                (rx, ry),
                rw, rh,
                facecolor=face,
                edgecolor=border_color,
                linewidth=line_width,
                zorder=3,
            )

            ax.add_patch(patch)

            strip_drawn = draw_income_strip(
                rx, ry, rw, rh, row, species
            )

            label_status = draw_st_label(
                patch,
                rx, ry, rw, rh,
                row.ST,
                int(row.Genome_Count),
                float(row.Percent_Within_Species),
                face,
                strip_drawn
            )

            tile_rows.append({
                "Species": species,
                "ST": row.ST,
                "Genome_Count": int(row.Genome_Count),
                "Percent_Within_Species": (
                    float(row.Percent_Within_Species)
                ),
                "N_H": int(row.N_H),
                "N_M": int(row.N_M),
                "N_LLM": int(row.N_LLM),
                "N_Unclassified": int(row.N_Unclassified),
                "Income_Breadth": int(row.Income_Breadth),
                "Cross_Income_ST": bool(row.Cross_Income_ST),
                "Income_Strip_Drawn": strip_drawn,
                "Label_Status": label_status,
                "Tile_X": rx,
                "Tile_Y": ry,
                "Tile_Width": rw,
                "Tile_Height": rh,
            })

        block_rows.append({
            "Species": species,
            "Species_All_Genomes": all_n,
            "Species_Typed_Genomes": typed_n,
            "Unassigned_ST_Genomes": all_n-typed_n,
            "Unique_STs": len(sp),
            "Percent_Of_Typed_Total": species_pct,
            "Block_X": x,
            "Block_Y": y,
            "Block_Width": block_w,
            "Block_Height": block_h,
        })

        current_x += block_w

    current_y += row_height

# ================================================================
# FIGURE TITLES AND LEGENDS
# ================================================================

fig.text(
    .5, .970,
    "Combined MLST abundance and country-income composition",
    ha="center",
    fontsize=22,
    fontweight="bold",
    color="#173352"
)

fig.text(
    .5, .943,
    (
        f"Enterobacter cloacae complex | "
        f"{TOTAL_TYPED:,} ST-assigned genomes"
    ),
    ha="center",
    fontsize=12.4,
    color="#566B7E"
)

income_handles = [
    Patch(
        facecolor=INCOME_COLORS[group],
        edgecolor="none",
        label=INCOME_LABELS[group],
    )
    for group in INCOME_GROUPS
]

income_handles.append(
    Patch(
        facecolor="white",
        edgecolor="#1D3043",
        linewidth=1.5,
        label="ST detected in ≥2 income groups",
    )
)

fig.legend(
    handles=income_handles,
    loc="lower center",
    bbox_to_anchor=(.5, .092),
    ncol=5,
    frameon=False,
    fontsize=10.2,
    columnspacing=1.5,
    handlelength=1.5,
)

fig.text(
    .5, .063,
    "Species block area and individual ST tile area "
    "are proportional to genome counts. "
    "All assigned STs are shown separately.",
    ha="center",
    fontsize=10,
    color="#40576C"
)

fig.text(
    .5, .044,
    "Income strips show the within-ST proportion of genomes "
    "from each income category. Tiny ST tiles are retained "
    "without overcrowding.",
    ha="center",
    fontsize=9.2,
    color="#586E80"
)

fig.text(
    .5, .025,
    "Unassigned STs are excluded from rectangle areas; "
    "income-unclassified genomes are retained and shown in grey "
    "where a strip is drawn.",
    ha="center",
    fontsize=8.6,
    color="#65788B"
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 94,
        "subsampling": 1
    }
)

plt.close(fig)
gc.collect()

# ================================================================
# VERIFY AND EXPORT COMPLETE STATISTICS
# ================================================================

tiles = pd.DataFrame(tile_rows)
blocks = pd.DataFrame(block_rows)

if len(tiles) != len(all_st):
    raise ValueError(
        "Some STs were omitted from the figure."
    )

if tiles.duplicated(["Species", "ST"]).any():
    raise ValueError(
        "Duplicate ST tiles detected."
    )

if int(tiles["Genome_Count"].sum()) != TOTAL_TYPED:
    raise ValueError(
        "Tile counts do not equal typed genomes."
    )

# Species-by-income totals, based only on typed genomes.
species_income = (
    typed.groupby(["Species", "Income"])
    .size()
    .unstack("Income", fill_value=0)
    .reindex(
        index=SPECIES,
        columns=INCOME_GROUPS,
        fill_value=0
    )
    .reset_index()
)

# STs shared across multiple income levels.
cross_income = all_st[
    all_st["Cross_Income_ST"]
].copy()

qc = pd.DataFrame([
    ["Total master genomes", len(df)],
    ["Resolved seven-species genomes", len(resolved)],
    ["ST-assigned genomes", len(typed)],
    ["Unassigned ST genomes", len(resolved)-len(typed)],
    ["Species-ST combinations", len(all_st)],
    ["Actual ST rectangles drawn", len(tiles)],
    ["Grouped or Other ST rectangles", 0],
    ["Income database", str(INCOME_FILE)],
    ["Income classified typed genomes",
     int(typed["Income"].ne("Unclassified").sum())],
    ["Income unclassified typed genomes",
     int(typed["Income"].eq("Unclassified").sum())],
    ["STs in >=2 classified income groups",
     len(cross_income)],
    ["Income strips drawn",
     int(tiles["Income_Strip_Drawn"].sum())],
    ["ST tiles too small for strip",
     int((~tiles["Income_Strip_Drawn"]).sum())],
    ["ST tile area", "Genome count"],
    ["Species block area", "ST-assigned genome count"],
    ["ST percent denominator", "All genomes in species"],
    ["Income strip denominator", "All assigned genomes in ST"],
    ["Interpretation", "MLST sharing across income groups, not transmission"],
    ["Plot format", "Single landscape JPEG, 600 DPI"],
], columns=["Parameter", "Value"])

# Export genome-level source columns and QC flags.
genome_export = df[[
    "Assembly_Accession",
    "Species",
    "ST",
    "Scheme",
    "Country_Harmonized",
    "Income",
    "Resolved",
    "Valid_ST",
    "Eligible",
]].copy()

with pd.ExcelWriter(
    XLSX,
    engine="openpyxl"
) as writer:

    genome_export.to_excel(
        writer,
        sheet_name="ALL_9157_GENOMES",
        index=False
    )

    all_st.to_excel(
        writer,
        sheet_name="ALL_ST_INCOME_COUNTS",
        index=False
    )

    tiles.to_excel(
        writer,
        sheet_name="EVERY_ST_TILE",
        index=False
    )

    cross_income.to_excel(
        writer,
        sheet_name="CROSS_INCOME_STS",
        index=False
    )

    species_income.to_excel(
        writer,
        sheet_name="SPECIES_INCOME_SUMMARY",
        index=False
    )

    blocks.to_excel(
        writer,
        sheet_name="SPECIES_BLOCKS",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    style_excel(writer)

print()
print("=" * 100)
print("FIGURE 8B — COMPLETED")
print("=" * 100)
print("Master genomes:", len(df))
print("Resolved genomes:", len(resolved))
print("ST-assigned genomes:", TOTAL_TYPED)
print("Individual ST rectangles:", len(tiles))
print("Cross-income STs:", len(cross_income))
print("Income strips drawn:", int(tiles["Income_Strip_Drawn"].sum()))
print("Grouped STs: ZERO")
print("JPEG:", JPG)
print("Excel:", XLSX)
print("=" * 100)
