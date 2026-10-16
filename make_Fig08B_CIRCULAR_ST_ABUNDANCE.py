
#!/usr/bin/env python3

# ============================================================
# FIGURE 8B — CIRCULAR MLST SEQUENCE-TYPE ABUNDANCE
#
# Seven separate circular bar plots for seven ECC species.
#
# Bar height  = number of genomes assigned to each ST.
# Bar color   = relative genome abundance.
# Outer label = ST name and genome count.
#
# Top 25 STs displayed per species.
# ALL STs exported to Excel.
#
# No source, host, income or environmental classification.
#
# Input:
# ECC_FINAL_2000_2026_MASTER.xlsx
#
# Outputs:
# 7 publication-style JPEGs at 600 DPI
# Complete ST abundance Excel workbook
# ============================================================

from pathlib import Path
import re
import gc

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib import cm, colors
from matplotlib.patches import Patch

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT
    / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

OUT = ROOT / "Fig_08B_Circular_MLST_Abundance"

OUT.mkdir(parents=True, exist_ok=True)

FIGDIR = OUT / "01_SEVEN_SPECIES_CIRCULAR_PLOTS"
FIGDIR.mkdir(exist_ok=True)

EXCEL = OUT / "Fig08B_All_Species_ST_Abundance.xlsx"

# ============================================================
# SETTINGS
# ============================================================

DPI = 600
TOP_N = 25

# Change this to 30 or 35 if you want more STs displayed.
# All STs are retained in Excel regardless.

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

SPECIES_COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 11,
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

# ============================================================
# FUNCTIONS
# ============================================================

def normalize_st(value):
    """
    Accept only valid positive numerical ST identifiers.
    Examples:
      113 -> ST113
      113.0 -> ST113
      ST113 -> ST113

    Exclude unassigned values such as:
      '-, -, NA, missing, failed
    """

    if pd.isna(value):
        return ""

    value = (
        str(value)
        .strip()
        .strip("'\"")
        .strip()
    )

    if value.upper().startswith("ST"):
        value = value[2:]

    if re.fullmatch(r"\d+\.0+", value):
        value = value.split(".")[0]

    if not re.fullmatch(r"\d+", value):
        return ""

    number = int(value)

    if number <= 0:
        return ""

    return f"ST{number}"


def format_integer(value):
    return f"{int(value):,}"


def save_jpeg(fig, path):
    fig.savefig(
        path,
        dpi=DPI,
        format="jpeg",
        facecolor="white",
        pil_kwargs={
            "quality": 94,
            "subsampling": 1
        }
    )

    plt.close(fig)
    gc.collect()

    print(
        "[SAVED]",
        path.name,
        f"{path.stat().st_size / 1024**2:.1f} MB"
    )


def excel_style(writer):
    for ws in writer.book.worksheets:

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid",
                fgColor="173352"
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )

        for j, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(j)
            ].width = min(
                45,
                max(17, len(str(cell.value)) + 3)
            )


# ============================================================
# LOAD DATA
# ============================================================

print("=" * 95)
print("FIGURE 8B — CIRCULAR ST ABUNDANCE")
print("=" * 95)

if not MASTER.is_file():
    raise FileNotFoundError(MASTER)

needed = [
    "Assembly_Accession",
    "Normalized_Species",
    "MLST_Scheme",
    "MLST_ST",
    "MLST_Status"
]

df = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=needed,
    engine="openpyxl",
    dtype=str
)

if len(df) != 9157:
    raise ValueError(
        f"Expected 9157 genomes, found {len(df)}"
    )

if (
    df["Assembly_Accession"].isna().any()
    or df["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicated assembly accession"
    )

df["Species"] = (
    df["Normalized_Species"]
    .fillna("")
    .str.strip()
)

df["Scheme"] = (
    df["MLST_Scheme"]
    .fillna("")
    .str.strip()
    .str.lower()
)

df["ST"] = df["MLST_ST"].map(normalize_st)

df["Is_Resolved_Species"] = (
    df["Species"].isin(SPECIES)
)

df["Valid_ST"] = df["ST"].ne("")

df["Valid_Scheme"] = (
    df["Scheme"].eq("ecloacae")
)

df["Figure_Eligible"] = (
    df["Is_Resolved_Species"]
    & df["Valid_ST"]
    & df["Valid_Scheme"]
)

valid = df[
    df["Figure_Eligible"]
].copy()

print("Total genomes:", len(df))
print(
    "Resolved-species genomes:",
    int(df["Is_Resolved_Species"].sum())
)
print("Valid typed genomes:", len(valid))
print(
    "Invalid/unassigned ST genomes:",
    int((~df["Valid_ST"]).sum())
)

# ============================================================
# CALCULATE ST ABUNDANCE
# ============================================================

all_abundance = (
    valid
    .groupby(["Species", "ST"], as_index=False)
    .agg(
        Genome_Count=(
            "Assembly_Accession",
            "nunique"
        )
    )
)

species_totals = (
    df[df["Is_Resolved_Species"]]
    .groupby("Species")
    .size()
    .to_dict()
)

typed_totals = (
    valid
    .groupby("Species")
    .size()
    .to_dict()
)

all_abundance["Species_Total_Genomes"] = (
    all_abundance["Species"]
    .map(species_totals)
)

all_abundance["Typed_Genomes"] = (
    all_abundance["Species"]
    .map(typed_totals)
)

all_abundance["Percent_of_All_Species_Genomes"] = (
    100.0
    * all_abundance["Genome_Count"]
    / all_abundance["Species_Total_Genomes"]
)

all_abundance["Percent_of_Typed_Genomes"] = (
    100.0
    * all_abundance["Genome_Count"]
    / all_abundance["Typed_Genomes"]
)

all_abundance = (
    all_abundance
    .sort_values(
        [
            "Species",
            "Genome_Count",
            "ST"
        ],
        ascending=[
            True,
            False,
            True
        ]
    )
    .reset_index(drop=True)
)

all_abundance["Abundance_Rank"] = (
    all_abundance
    .groupby("Species")
    .cumcount() + 1
)

all_abundance["Displayed_In_Figure"] = (
    all_abundance["Abundance_Rank"] <= TOP_N
)

summary_rows = []

# ============================================================
# PLOT ONE CIRCULAR ABUNDANCE FIGURE PER SPECIES
# ============================================================

for species_index, species in enumerate(SPECIES, 1):

    sp = all_abundance[
        all_abundance["Species"].eq(species)
    ].copy()

    total_n = int(
        species_totals.get(species, 0)
    )

    typed_n = int(
        typed_totals.get(species, 0)
    )

    unassigned_n = total_n - typed_n

    unique_sts = len(sp)

    top = sp.head(TOP_N).copy()

    plotted_n = int(top["Genome_Count"].sum())

    remaining_n = typed_n - plotted_n

    summary_rows.append({
        "Species": species,
        "All_Genomes": total_n,
        "Valid_Typed_Genomes": typed_n,
        "Unassigned_ST_Genomes": unassigned_n,
        "Unique_STs": unique_sts,
        "STs_Displayed": len(top),
        "Genomes_In_Top_STs": plotted_n,
        "Genomes_In_Remaining_STs": remaining_n,
        "Top_ST_Coverage_Percent": (
            100 * plotted_n / typed_n
            if typed_n else np.nan
        )
    })

    if top.empty:
        print(
            "[SKIP]",
            species,
            "no numeric ST"
        )
        continue

    # Arrange STs by descending abundance.
    st_names = top["ST"].tolist()
    counts = top["Genome_Count"].to_numpy(float)

    n = len(st_names)

    # Circular angles.
    angles = np.linspace(
        0,
        2*np.pi,
        n,
        endpoint=False
    )

    width = 2*np.pi/n * 0.81

    # Square-root scaling prevents one dominant ST
    # from compressing every other radial bar.
    maximum = max(1.0, counts.max())

    bar_heights = (
        0.18
        + 0.57 * np.sqrt(counts/maximum)
    )

    bar_bottom = 0.36

    # Bar color encodes the actual count
    # through a continuous colormap.
    normalizer = colors.Normalize(
        vmin=0,
        vmax=maximum
    )

    cmap = plt.get_cmap("YlOrRd")

    bar_colors = [
        cmap(normalizer(v))
        for v in counts
    ]

    fig = plt.figure(
        figsize=(15.0, 15.0),
        facecolor="white"
    )

    ax = fig.add_axes(
        [0.12, 0.16, 0.76, 0.70],
        projection="polar"
    )

    ax.set_theta_offset(np.pi/2)
    ax.set_theta_direction(-1)

    # Clean radial chart: no polar grid clutter.
    ax.grid(False)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.spines["polar"].set_visible(False)

    ax.set_ylim(0, 1.42)

    bars = ax.bar(
        angles,
        bar_heights,
        width=width,
        bottom=bar_bottom,
        align="center",
        color=bar_colors,
        edgecolor="white",
        linewidth=1.5,
        alpha=0.98,
        zorder=3
    )

    # Fine guide ring at bar origin.
    ring = np.linspace(
        0,
        2*np.pi,
        500
    )

    ax.plot(
        ring,
        np.full_like(ring, bar_bottom),
        color="#B9C6D0",
        linewidth=0.9,
        alpha=0.7
    )

    # Labels outside each radial bar.
    for angle, height, st, count in zip(
        angles,
        bar_heights,
        st_names,
        counts
    ):

        angle_deg = np.degrees(angle)

        # Tangential text orientation.
        rotation = 90 - angle_deg

        if 90 < angle_deg < 270:
            rotation += 180

        radius = (
            bar_bottom
            + height
            + 0.11
        )

        ax.text(
            angle,
            radius,
            f"{st}\n(n={int(count):,})",
            rotation=rotation,
            rotation_mode="anchor",
            ha="center",
            va="center",
            fontsize=10.2,
            linespacing=1.2,
            color="#24384A",
            fontweight="medium",
            clip_on=False
        )

    # White central area for summary values.
    ax.scatter(
        [0],
        [0],
        s=10,
        color="white",
        zorder=5
    )

    ax.text(
        0,
        0,
        (
            f"{total_n:,} genomes\n"
            f"{unique_sts:,} unique STs\n"
            f"{typed_n:,} assigned\n"
            f"{unassigned_n:,} unassigned"
        ),
        ha="center",
        va="center",
        fontsize=13,
        fontweight="bold",
        color="#26394D",
        linespacing=1.75,
        zorder=10
    )

    # Only the full species name at the top.
    fig.suptitle(
        species,
        fontsize=23,
        fontweight="bold",
        fontstyle="italic",
        color=SPECIES_COLORS[species],
        y=0.965
    )

    # Quantitative color scale.
    cax = fig.add_axes([
        0.32,
        0.090,
        0.36,
        0.018
    ])

    scalar = cm.ScalarMappable(
        norm=normalizer,
        cmap=cmap
    )

    scalar.set_array([])

    cbar = fig.colorbar(
        scalar,
        cax=cax,
        orientation="horizontal"
    )

    cbar.set_label(
        "Genome count per sequence type",
        fontsize=11,
        labelpad=8
    )

    cbar.ax.tick_params(
        labelsize=9
    )

    fig.text(
        0.5,
        0.040,
        (
            f"Top {n} STs shown by abundance; "
            f"{remaining_n:,} typed genomes belong "
            "to STs outside the displayed set."
        ),
        ha="center",
        fontsize=10,
        color="#52677B"
    )

    fig.text(
        0.5,
        0.022,
        (
            "Radial bar height uses square-root scaling "
            "for legibility; color represents the actual "
            "genome count. Full abundance table in Excel."
        ),
        ha="center",
        fontsize=9,
        color="#687C8E"
    )

    outfile = FIGDIR / (
        f"Fig08B_{species_index:02d}_"
        f"{species.split()[-1]}_"
        "Circular_ST_Abundance_600DPI.jpg"
    )

    save_jpeg(fig, outfile)

# ============================================================
# EXTRA TABLE — ABUNDANCE BY SPECIES
# ============================================================

summary = pd.DataFrame(summary_rows)

# Complete species x ST count matrix
st_matrix = all_abundance.pivot_table(
    index="ST",
    columns="Species",
    values="Genome_Count",
    aggfunc="sum",
    fill_value=0
)

st_matrix = st_matrix.reindex(
    columns=SPECIES,
    fill_value=0
)

st_matrix["Total_Genomes_Across_Species"] = (
    st_matrix.sum(axis=1)
)

st_matrix = (
    st_matrix
    .sort_values(
        "Total_Genomes_Across_Species",
        ascending=False
    )
    .reset_index()
)

# Important: same ST number in different species
# is NOT treated as the same transmitted lineage.

qc = pd.DataFrame([
    ["Master genomes", len(df)],
    [
        "Seven resolved species genomes",
        int(df["Is_Resolved_Species"].sum())
    ],
    ["Valid ST-assigned genomes", len(valid)],
    [
        "Invalid or unassigned ST genomes",
        int((~df["Valid_ST"]).sum())
    ],
    ["MLST scheme", "ecloacae"],
    ["Plotted STs per species", TOP_N],
    [
        "Abundance metric",
        "Number of genomes assigned to a species-specific ST"
    ],
    [
        "Bar scaling",
        "Square root of genome count, displayed as radial height"
    ],
    [
        "Color metric",
        "Genome count, normalized to each species' maximum"
    ],
    [
        "ST matching rule",
        "ST labels analyzed separately within each species"
    ],
    [
        "Data limitation",
        "Only valid positive numerical MLST_ST labels were used"
    ]
], columns=["Parameter", "Value"])

# ============================================================
# EXPORT FULL EXCEL
# ============================================================

with pd.ExcelWriter(
    EXCEL,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="ALL_9157_MLST_QC",
        index=False
    )

    valid.to_excel(
        writer,
        sheet_name="VALID_ST_GENOMES",
        index=False
    )

    all_abundance.to_excel(
        writer,
        sheet_name="ALL_ST_ABUNDANCE",
        index=False
    )

    st_matrix.to_excel(
        writer,
        sheet_name="SPECIES_ST_MATRIX",
        index=False
    )

    summary.to_excel(
        writer,
        sheet_name="SPECIES_SUMMARY",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    excel_style(writer)

print()
print("=" * 95)
print("FIGURE 8B — COMPLETED")
print("=" * 95)

print("Master genomes:", len(df))
print("Valid typed genomes:", len(valid))
print("Species figures:", len(SPECIES))
print("Figures:", FIGDIR)
print("Excel:", EXCEL)

print("\nSPECIES SUMMARY")
print(summary.to_string(index=False))

print("=" * 95)
