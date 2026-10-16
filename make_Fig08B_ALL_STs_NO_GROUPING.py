
#!/usr/bin/env python3

# ============================================================
# FIGURE 8B — COMPLETE ECC MLST TREEMAP
#
# EVERY VALID ST IS A SEPARATE RECTANGLE.
# NO "OTHER STs" AND NO AGGREGATION.
#
# One large landscape figure, seven species.
# Species area proportional to ST-assigned genome count.
# ST tile area proportional to its genome count.
#
# Tiny tiles are retained even when labels cannot fit.
# All labels and counts exported to Excel.
#
# Outputs: 600-DPI JPEG and XLSX only.
# ============================================================

from pathlib import Path
import re
import gc

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.colors import to_rgb

import squarify

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\Enterobacter\ECC_NPJ_AMR_FIGURES")

MASTER = (
    ROOT / "00_FINAL_COUNTRY_REPORTED_MASTER"
    / "ECC_FINAL_2000_2026_MASTER.xlsx"
)

OUT = ROOT / "Fig_08B_Combined_MLST_Treemap"
OUT.mkdir(parents=True, exist_ok=True)

JPG = OUT / "Fig08B_ALL_STs_Complete_Combined_Treemap_600DPI.jpg"
XLSX = OUT / "Fig08B_ALL_STs_Complete_Abundance.xlsx"

DPI = 600
FIG_WIDTH = 25
FIG_HEIGHT = 18

SPECIES = [
    "Enterobacter hormaechei",
    "Enterobacter cloacae",
    "Enterobacter roggenkampii",
    "Enterobacter asburiae",
    "Enterobacter kobei",
    "Enterobacter ludwigii",
    "Enterobacter bugandensis",
]

COLORS = {
    "Enterobacter hormaechei": "#0072B2",
    "Enterobacter cloacae": "#D55E00",
    "Enterobacter roggenkampii": "#009E73",
    "Enterobacter asburiae": "#CC79A7",
    "Enterobacter kobei": "#7851A9",
    "Enterobacter ludwigii": "#A6761D",
    "Enterobacter bugandensis": "#56B4E9",
}

ROWS = [
    ["Enterobacter hormaechei"],
    [
        "Enterobacter cloacae",
        "Enterobacter roggenkampii",
        "Enterobacter asburiae",
    ],
    [
        "Enterobacter kobei",
        "Enterobacter ludwigii",
        "Enterobacter bugandensis",
    ],
]

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "figure.facecolor": "white",
    "savefig.facecolor": "white",
})

# ============================================================
# CLEAN DATA
# ============================================================

def clean_st(value):
    if pd.isna(value):
        return ""

    s = str(value).strip().strip("'\"").strip()

    if s.upper().startswith("ST"):
        s = s[2:].strip()

    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]

    if re.fullmatch(r"[0-9]+", s) is None:
        return ""

    n = int(s)
    return f"ST{n}" if n > 0 else ""


def lighten(color, amount):
    rgb = np.array(to_rgb(color))
    return tuple(
        np.clip(rgb + (1 - rgb) * amount, 0, 1)
    )


def text_color(color):
    r, g, b = to_rgb(color)

    lum = (
        0.2126*r +
        0.7152*g +
        0.0722*b
    )

    return "white" if lum < 0.48 else "#163044"


def style_workbook(writer):
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = PatternFill(
                "solid", fgColor="173352"
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )

        for index, cell in enumerate(ws[1], 1):
            ws.column_dimensions[
                get_column_letter(index)
            ].width = min(
                46,
                max(16, len(str(cell.value)) + 3)
            )


def fitted_label(ax, patch, x, y, w, h,
                 st, count, percentage, face,
                 axes_width_inches, axes_height_inches):

    # Convert tile dimensions into physical units.
    tile_width_in = w / 1000.0 * axes_width_inches
    tile_height_in = h / 720.0 * axes_height_inches

    if tile_width_in >= 1.00 and tile_height_in >= 0.68:
        label = (
            f"{st}\n"
            f"n={count:,}\n"
            f"{percentage:.2f}%"
        )
        fs = min(
            12.5,
            tile_width_in * 72 / max(5.0, len(st) * 0.68),
            tile_height_in * 21.0
        )

    elif tile_width_in >= 0.62 and tile_height_in >= 0.38:
        label = f"{st}\nn={count:,}"
        fs = min(
            9.5,
            tile_width_in * 72 / max(5.0, len(st) * 0.68),
            tile_height_in * 22.0
        )

    elif tile_width_in >= 0.42 and tile_height_in >= 0.22:
        label = st
        fs = min(
            7.8,
            tile_width_in * 72 / max(5.0, len(st) * 0.68),
            tile_height_in * 31.0
        )

    else:
        return "Too_small_to_label"

    if fs < 5.7:
        return "Too_small_to_label"

    text = ax.text(
        x + w/2,
        y + h/2,
        label,
        ha="center",
        va="center",
        color=text_color(face),
        fontweight="bold",
        fontsize=fs,
        linespacing=1.12,
        zorder=5
    )

    text.set_clip_path(patch)

    return "Labeled"


# ============================================================
# READ FINAL MASTER
# ============================================================

print("=" * 100)
print("FIGURE 8B — ALL SEQUENCE TYPES, NO GROUPING")
print("=" * 100)

if not MASTER.is_file():
    raise FileNotFoundError(MASTER)

data = pd.read_excel(
    MASTER,
    sheet_name="FINAL_MASTER_2000_2026",
    usecols=[
        "Assembly_Accession",
        "Normalized_Species",
        "MLST_Scheme",
        "MLST_ST",
        "MLST_Status",
    ],
    engine="openpyxl",
    dtype=str
)

if len(data) != 9157:
    raise ValueError(
        f"Expected 9157 genomes; found {len(data)}"
    )

if (
    data["Assembly_Accession"].isna().any()
    or data["Assembly_Accession"].duplicated().any()
):
    raise ValueError(
        "Missing or duplicate assembly accessions."
    )

data["Species"] = (
    data["Normalized_Species"]
    .fillna("")
    .str.strip()
)

data["Scheme"] = (
    data["MLST_Scheme"]
    .fillna("")
    .str.strip()
    .str.lower()
)

data["ST"] = data["MLST_ST"].map(clean_st)

data["Resolved"] = data["Species"].isin(SPECIES)

data["ST_Assigned"] = (
    data["Resolved"]
    & data["Scheme"].eq("ecloacae")
    & data["ST"].ne("")
)

resolved = data[data["Resolved"]].copy()
typed = data[data["ST_Assigned"]].copy()

if len(resolved) != 8882:
    raise ValueError(
        f"Expected 8882 resolved genomes, found {len(resolved)}"
    )

if typed.empty:
    raise ValueError("No valid ST assignments.")

species_all = (
    resolved.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

species_typed = (
    typed.groupby("Species")
    .size()
    .reindex(SPECIES, fill_value=0)
)

abundance = (
    typed.groupby(["Species", "ST"])
    .agg(Genome_Count=("Assembly_Accession", "nunique"))
    .reset_index()
)

abundance["ST_Number"] = (
    abundance["ST"].str[2:].astype(int)
)

abundance = abundance.sort_values(
    ["Species", "Genome_Count", "ST_Number"],
    ascending=[True, False, True]
).reset_index(drop=True)

abundance["Rank"] = (
    abundance.groupby("Species").cumcount() + 1
)

abundance["Species_Total_Genomes"] = (
    abundance["Species"].map(species_all)
)

abundance["Species_Typed_Genomes"] = (
    abundance["Species"].map(species_typed)
)

abundance["Percent_of_Species_Total"] = (
    100.0 * abundance["Genome_Count"]
    / abundance["Species_Total_Genomes"]
)

abundance["Percent_of_ST_Assigned_Genomes"] = (
    100.0 * abundance["Genome_Count"]
    / abundance["Species_Typed_Genomes"]
)

# Verify all assigned genomes are represented.
if int(abundance["Genome_Count"].sum()) != len(typed):
    raise ValueError(
        "ST abundance counts do not sum to assigned genomes."
    )

print("Full master:", len(data))
print("Resolved seven-species genomes:", len(resolved))
print("Valid ST-assigned genomes:", len(typed))
print("Distinct species-ST combinations:", len(abundance))
print("Unassigned ST genomes:", len(resolved)-len(typed))

# ============================================================
# HIERARCHICAL SPECIES BLOCKS
# ============================================================

fig = plt.figure(
    figsize=(FIG_WIDTH, FIG_HEIGHT),
    facecolor="white"
)

ax = fig.add_axes([
    0.028, 0.092,
    0.944, 0.817
])

AXES_W_IN = FIG_WIDTH * 0.944
AXES_H_IN = FIG_HEIGHT * 0.817

ax.set_xlim(0, 1000)
ax.set_ylim(0, 720)
ax.invert_yaxis()
ax.axis("off")
ax.set_aspect("auto")

TOTAL_TYPED = len(typed)
current_y = 0.0

tile_records = []
block_records = []
summary_records = []

for group in ROWS:

    row_total = sum(
        int(species_typed[sp])
        for sp in group
    )

    row_height = 720.0 * row_total / TOTAL_TYPED
    current_x = 0.0

    for species in group:

        n_typed = int(species_typed[species])
        n_all = int(species_all[species])

        width = 1000.0 * n_typed / row_total
        height = row_height

        x = current_x
        y = current_y
        color = COLORS[species]

        if n_typed == 0:
            current_x += width
            continue

        # Outer species block.
        ax.add_patch(
            Rectangle(
                (x + 1.4, y + 1.4),
                width - 2.8,
                height - 2.8,
                facecolor=color,
                edgecolor="white",
                linewidth=2.4,
                zorder=1
            )
        )

        header = min(31.0, height * 0.24)

        species_pct = (
            100.0 * n_typed / TOTAL_TYPED
        )

        short_name = species.replace(
            "Enterobacter ", "E. "
        )

        if width > 370:
            title = (
                f"{species} "
                f"(n={n_typed:,}; {species_pct:.1f}%)"
            )
            title_fs = 15
        else:
            title = (
                f"{short_name} (n={n_typed:,})"
            )
            title_fs = 10.2 if width >= 200 else 8.2

        ax.text(
            x + 9,
            y + header * 0.57,
            title,
            fontsize=title_fs,
            fontweight="bold",
            fontstyle="italic",
            color="white",
            va="center",
            zorder=7
        )

        # Region for ALL individual STs.
        px = x + 6
        py = y + header + 2
        pw = max(0.1, width - 12)
        ph = max(0.1, height - header - 8)

        subset = abundance[
            abundance["Species"].eq(species)
        ].copy()

        # Squarify works on sorted positive areas.
        subset = subset.sort_values(
            ["Genome_Count", "ST_Number"],
            ascending=[False, True]
        ).reset_index(drop=True)

        sizes = subset["Genome_Count"].astype(float).tolist()

        rects = squarify.squarify(
            squarify.normalize_sizes(sizes, pw, ph),
            px, py, pw, ph
        )

        if len(rects) != len(subset):
            raise ValueError(
                f"Tile count mismatch for {species}"
            )

        # Each ST has its own rectangle.
        for idx, (rect, row) in enumerate(
            zip(rects, subset.itertuples(index=False))
        ):

            rx = float(rect["x"])
            ry = float(rect["y"])
            rw = float(rect["dx"])
            rh = float(rect["dy"])

            # Darker colors for abundant STs.
            # Light shades for rarer STs.
            shade = (
                0.08 + 0.76 * idx / max(1, len(subset)-1)
            )

            face = lighten(color, shade)

            # Do not inset or remove small rectangles:
            # every ST must remain visible in the layout.
            patch = Rectangle(
                (rx, ry),
                rw, rh,
                facecolor=face,
                edgecolor="white",
                linewidth=0.20 if rw < 2 or rh < 2 else 0.60,
                zorder=3
            )

            ax.add_patch(patch)

            status = fitted_label(
                ax, patch,
                rx, ry, rw, rh,
                row.ST,
                int(row.Genome_Count),
                float(row.Percent_of_Species_Total),
                face,
                AXES_W_IN,
                AXES_H_IN
            )

            tile_records.append({
                "Species": species,
                "ST": row.ST,
                "Genome_Count": int(row.Genome_Count),
                "Percent_of_Species_Total":
                    float(row.Percent_of_Species_Total),
                "Percent_of_Assigned":
                    float(row.Percent_of_ST_Assigned_Genomes),
                "Rank": int(row.Rank),
                "Tile_X": rx,
                "Tile_Y": ry,
                "Tile_Width": rw,
                "Tile_Height": rh,
                "Label_Status": status,
            })

        block_records.append({
            "Species": species,
            "Block_X": x,
            "Block_Y": y,
            "Block_Width": width,
            "Block_Height": height,
            "ST_Assigned_Genomes": n_typed,
            "Percent_of_Assigned_Total": species_pct
        })

        summary_records.append({
            "Species": species,
            "All_Genomes": n_all,
            "ST_Assigned_Genomes": n_typed,
            "Unassigned_ST_Genomes": n_all - n_typed,
            "Unique_STs": len(subset),
            "Percent_of_All_Resolved_Genomes":
                100.0 * n_all / len(resolved),
            "Percent_of_All_Assigned_Genomes": species_pct
        })

        current_x += width

    current_y += row_height

# ============================================================
# FIGURE ANNOTATIONS
# ============================================================

fig.text(
    0.5, 0.968,
    "Combined hierarchical MLST abundance treemap",
    ha="center",
    fontsize=24,
    fontweight="bold",
    color="#173352"
)

fig.text(
    0.5, 0.939,
    (
        "Enterobacter cloacae complex"
        f"  |  {TOTAL_TYPED:,} ST-assigned genomes"
    ),
    ha="center",
    fontsize=13,
    color="#4C6277"
)

fig.text(
    0.5, 0.068,
    "Each individual rectangle represents ONE MLST sequence type. "
    "No sequence types are combined.",
    ha="center",
    fontsize=11,
    color="#34495E"
)

fig.text(
    0.5, 0.047,
    "Species and ST areas are proportional to ST-assigned "
    "genome counts. Percentages within tiles use all genomes "
    "of the corresponding species.",
    ha="center",
    fontsize=9.9,
    color="#53697D"
)

fig.text(
    0.5, 0.028,
    "Rare STs remain as individual small rectangles; "
    "complete ST identifiers, counts and prevalence "
    "are provided in Excel.",
    ha="center",
    fontsize=9.4,
    color="#65798A"
)

fig.savefig(
    JPG,
    dpi=DPI,
    format="jpeg",
    facecolor="white",
    pil_kwargs={
        "quality": 94,
        "subsampling": 1
    }
)

plt.close(fig)
gc.collect()

# ============================================================
# EXPORT COMPLETE SOURCE DATA AND QC
# ============================================================

tiles = pd.DataFrame(tile_records)
blocks = pd.DataFrame(block_records)
summary = pd.DataFrame(summary_records)

if len(tiles) != len(abundance):
    raise ValueError(
        "Not every ST received a tile."
    )

if int(tiles["Genome_Count"].sum()) != len(typed):
    raise ValueError(
        "Tile genome counts do not match typed genomes."
    )

qc = pd.DataFrame([
    ["Master genomes", len(data)],
    ["Seven resolved-species genomes", len(resolved)],
    ["ST-assigned genomes", len(typed)],
    ["Unassigned ST genomes", len(resolved)-len(typed)],
    ["Distinct species-ST combinations", len(abundance)],
    ["Individual ST tiles drawn", len(tiles)],
    ["Grouped ST tiles", 0],
    ["MLST scheme", "ecloacae"],
    ["ST area metric", "Actual ST-assigned genome count"],
    ["Species area metric", "ST-assigned genomes per species"],
    ["ST prevalence denominator", "All genomes in species"],
    ["MLST interpretation", "Sequence type, not serotype"],
    ["Figure resolution", "600 DPI JPEG"],
    ["Figure dimensions", f"{FIG_WIDTH} x {FIG_HEIGHT} inches"],
    ["Excluded from tiles", "Unresolved species and invalid ST assignments"],
    ["Tiny ST labeling", "Exact labels in Excel when tile is too small"],
], columns=["Parameter", "Value"])

with pd.ExcelWriter(
    XLSX, engine="openpyxl"
) as writer:

    data.to_excel(
        writer,
        sheet_name="ALL_9157_MLST_QC",
        index=False
    )

    abundance.to_excel(
        writer,
        sheet_name="ALL_ST_ABUNDANCE",
        index=False
    )

    tiles.to_excel(
        writer,
        sheet_name="EVERY_PLOTTED_ST_TILE",
        index=False
    )

    blocks.to_excel(
        writer,
        sheet_name="SPECIES_LAYOUT",
        index=False
    )

    summary.to_excel(
        writer,
        sheet_name="SEVEN_SPECIES_SUMMARY",
        index=False
    )

    qc.to_excel(
        writer,
        sheet_name="QC_METHODS",
        index=False
    )

    style_workbook(writer)

print()
print("=" * 100)
print("COMPLETE ALL-ST TREEMAP FINISHED")
print("=" * 100)
print("Resolved genomes:", len(resolved))
print("ST-assigned genomes:", len(typed))
print("Individual ST tiles:", len(tiles))
print("Grouped tiles: ZERO")
print("JPEG:", JPG)
print("Excel:", XLSX)
print()
print(summary.to_string(index=False))
print("=" * 100)
